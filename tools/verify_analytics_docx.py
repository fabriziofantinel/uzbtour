from pathlib import Path
from zipfile import ZipFile, BadZipFile
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
files = [p for p in (ROOT / "docs").rglob("*.docx") if "templates" not in p.parts]
marker = "Aggiornamento Analytics per l'agenzia"
failures = []
results = []

for path in files:
    try:
        with ZipFile(path) as package:
            bad = package.testzip()
            if bad:
                failures.append(f"{path.name}: elemento OOXML corrotto {bad}")
        document = Document(path)
        marker_count = sum(1 for paragraph in document.paragraphs if paragraph.text.strip() == marker)
        all_text = "\n".join(
            [paragraph.text for paragraph in document.paragraphs]
            + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
        )
        if marker_count != 1:
            failures.append(f"{path.name}: marker Analytics={marker_count}")
        expected = (
            "Evento Analytics" if "Logico" in path.name else
            "ops.product_analytics_events" if "Fisico" in path.name else
            "Componenti e responsabilita" if "Architettura" in path.name else
            "Casi d'uso Analytics" if "Casi_Uso" in path.name else
            "Funzionalita Analytics"
        )
        if expected not in all_text:
            failures.append(f"{path.name}: contenuto Analytics incompleto")
        empty_tables = sum(1 for table in document.tables if not table.rows or not table.columns)
        if empty_tables:
            failures.append(f"{path.name}: {empty_tables} tabelle vuote")
        results.append({
            "document": str(path.relative_to(ROOT)),
            "paragraphs": len(document.paragraphs),
            "tables": len(document.tables),
            "analytics_marker": marker_count,
        })
    except (BadZipFile, Exception) as error:
        failures.append(f"{path.name}: {error}")

if failures:
    raise SystemExit("\n".join(failures))

for result in results:
    print(result)
print({"status": "passed", "documents": len(results)})
