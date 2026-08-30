from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
DOCS = [p for p in (ROOT / "docs").rglob("*.docx") if "templates" not in p.parts]
MARKER = "Aggiornamento Analytics per l'agenzia"


def remove_existing_appendix(document: Document) -> None:
    body = document._element.body
    marker = next((p for p in document.paragraphs if p.text.strip() == MARKER), None)
    if marker is None:
        return
    current = marker._element
    while current is not None:
        following = current.getnext()
        if current.tag != qn("w:sectPr"):
            body.remove(current)
        current = following


def keep_with_next(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    p_pr.append(keep)


def heading(document: Document, text: str, level: int = 1):
    paragraph = document.add_heading(text, level=level)
    keep_with_next(paragraph)
    return paragraph


def bullet(document: Document, text: str) -> None:
    document.add_paragraph(text, style="List Bullet")


def table(document: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    result = document.add_table(rows=1, cols=len(headers))
    result.style = "Table Grid"
    result.autofit = False
    header = result.rows[0]
    header._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for index, value in enumerate(headers):
        header.cells[index].text = value
        header.cells[index].width = Inches(widths[index])
        header.cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in header.cells[index].paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(9)
    for values in rows:
        cells = result.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
            cells[index].width = Inches(widths[index])
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cells[index].paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
    document.add_paragraph()


def common_intro(document: Document) -> None:
    document.add_page_break()
    heading(document, MARKER, 1)
    document.add_paragraph(
        "Questa appendice recepisce la funzionalita Analytics introdotta nel pannello agenzia. "
        "La soluzione misura l'adozione e l'utilizzo del servizio senza mostrare all'agenzia "
        "il dettaglio personale delle azioni del singolo viaggiatore."
    )


def add_architecture(document: Document) -> None:
    heading(document, "Componenti e responsabilita", 2)
    table(document, ["Componente", "Responsabilita", "Tecnologia"], [
        ["Client viaggiatore", "Emette eventi di sessione, consultazione programma, apertura elenco documenti e download.", "React 19, fetch keepalive, sessionStorage"],
        ["API Analytics", "Valida identita, payload, partenza, gruppo e giornata prima della registrazione.", "Next.js App Router, Zod"],
        ["Persistenza eventi", "Conserva eventi append-only, idempotenti e isolati per tenant.", "Neon PostgreSQL, RLS, SECURITY DEFINER"],
        ["Servizio aggregazione", "Calcola indicatori per agenzia, periodo, partenza e tappa.", "SQL tipizzato, query aggregate"],
        ["Dashboard agenzia", "Espone funnel, indicatori, confronto partenze e feedback operativo.", "React Server Component, CSS white-label"],
    ], [1.25, 3.55, 1.7])
    heading(document, "Flusso end-to-end", 2)
    for item in [
        "Il viaggiatore autenticato apre una sezione dell'esperienza di viaggio.",
        "Il client genera session_id e client_operation_id UUID e invia l'evento all'API.",
        "L'API ricava l'attore dalla sessione autenticata; la stored function verifica appartenenza a partenza e gruppo.",
        "Neon applica deduplicazione semantica, limite orario, vincoli referenziali e RLS forzata.",
        "La dashboard interroga soltanto dati aggregati entro il perimetro agency_id dell'utente agenzia.",
    ]:
        bullet(document, item)
    heading(document, "Sicurezza, resilienza e privacy", 2)
    for item in [
        "agency_id e il campo leading degli indici operativi e delle policy RLS.",
        "Il client non puo selezionare l'agenzia; lo scope deriva dalle relazioni partenza-gruppo-viaggiatore.",
        "Gli eventi duplicati nella stessa sessione non alterano i conteggi e il rate limit limita replay automatizzati.",
        "Gli errori Analytics non interrompono l'esperienza del viaggiatore; il client ritenta gli errori HTTP o di rete.",
        "La dashboard mostra conteggi e medie, non una cronologia nominativa delle azioni individuali.",
    ]:
        bullet(document, item)


def add_logical(document: Document) -> None:
    heading(document, "Estensione del modello concettuale", 2)
    table(document, ["Entita", "Scopo", "Relazioni principali"], [
        ["Evento Analytics", "Fatto atomico prodotto dall'utilizzo dell'app viaggiatore.", "Agenzia, Utente, Viaggiatore, Partenza, Gruppo, Giornata"],
        ["Sessione Analytics", "Contesto client temporaneo usato per deduplicare consultazioni equivalenti.", "Utente e insieme di Eventi Analytics"],
        ["Indicatore Agenzia", "Vista aggregata calcolata, non persistita come dato master.", "Periodo, Agenzia e facoltativamente Partenza"],
        ["Feedback Tappa", "Valutazione gia presente, aggregata per giornata e attivita effettiva.", "Partenza, Giornata, Tappa, Viaggiatore"],
    ], [1.35, 2.75, 2.4])
    heading(document, "Semantica degli indicatori", 2)
    table(document, ["Indicatore", "Definizione funzionale"], [
        ["Invitati", "Viaggiatori associati a una partenza nel perimetro selezionato."],
        ["Attivati", "Viaggiatori con identita digitale attiva."],
        ["Attivi", "Viaggiatori distinti con almeno un evento Analytics nel periodo."],
        ["Consultazioni programma", "Eventi programme_view validi e deduplicati."],
        ["Uso documenti", "Viaggiatori distinti con almeno un document_download."],
        ["Assistenza", "Messaggi inviati dai viaggiatori nella chat operativa."],
        ["Engagement", "Tentativi quiz e sfide, oltre alle partecipazioni ai contest, completati nel periodo."],
        ["Feedback per tappa", "Media, numerosita e ordinamento delle valutazioni per attivita effettiva."],
    ], [1.75, 4.75])
    heading(document, "Regole di business", 2)
    for item in [
        "Ogni evento appartiene a una sola agenzia, partenza e gruppo.",
        "La giornata e opzionale per gli eventi generali e obbligatoriamente coerente con la partenza quando valorizzata.",
        "La selezione di una partenza restringe tutti gli indicatori della pagina.",
        "Le metriche storiche di consultazione decorrono dall'attivazione del tracking; feedback, chat ed engagement usano anche dati applicativi gia esistenti.",
    ]:
        bullet(document, item)


def add_physical(document: Document) -> None:
    heading(document, "Oggetto fisico ops.product_analytics_events", 2)
    table(document, ["Campo", "Tipo", "Vincolo / significato"], [
        ["id", "UUID", "PK, default uuidv7()"],
        ["agency_id", "UUID", "FK iam.agencies; chiave tenant obbligatoria"],
        ["actor_user_id", "UUID", "FK iam.users; autore autenticato"],
        ["traveler_id", "UUID nullable", "FK travel.traveler_profiles"],
        ["departure_id / party_id", "UUID", "FK composita travel.travel_parties"],
        ["departure_day_id", "UUID nullable", "FK composita travel.departure_days"],
        ["event_name", "varchar(40)", "CHECK su traveler_session, programme_view, document_list_view, document_download"],
        ["session_id", "UUID", "Sessione client"],
        ["client_operation_id", "UUID", "Idempotenza della singola operazione"],
        ["properties", "JSONB", "Oggetto con soli attributi contestuali non sensibili"],
        ["occurred_at / created_at", "timestamptz", "Tempo evento e persistenza in UTC"],
    ], [1.55, 1.3, 3.65])
    heading(document, "Indici e vincoli operativi", 2)
    for item in [
        "product_analytics_tenant_period_idx (agency_id, occurred_at DESC, event_name, departure_id).",
        "product_analytics_departure_actor_idx (agency_id, departure_id, actor_user_id, occurred_at DESC).",
        "Indice univoco semantico per attore, sessione, evento, partenza, giornata e documentId.",
        "Vincolo univoco (actor_user_id, client_operation_id) per retry idempotenti.",
        "ENABLE e FORCE RLS con policy basata su current_setting('app.agency_id').",
    ]:
        bullet(document, item)
    heading(document, "API di scrittura e autorizzazioni", 2)
    document.add_paragraph(
        "app.record_product_analytics_event_v3 e una funzione SECURITY DEFINER con search_path bloccato. "
        "Risolve l'identita legacy, verifica l'appartenenza attiva del viaggiatore al gruppo, controlla la giornata, "
        "applica un limite di 120 eventi per attore/ora e restituisce l'identificatore gia esistente in caso di duplicato. "
        "Il ruolo smf_app dispone soltanto di EXECUTE sulla funzione e SELECT sulla tabella protetta da RLS."
    )


def add_functional(document: Document) -> None:
    heading(document, "Funzionalita Analytics agenzia", 2)
    table(document, ["Area", "Funzione offerta", "Valore operativo"], [
        ["Adozione", "Funnel invitati, attivati e utenti attivi con percentuali.", "Individua gruppi che necessitano accompagnamento."],
        ["Programma", "Conteggio consultazioni per periodo e partenza.", "Misura se il programma viene realmente usato."],
        ["Documenti", "Numero di viaggiatori distinti che scaricano file.", "Evidenzia voucher o documenti poco consultati."],
        ["Assistenza", "Conteggio richieste inviate dai viaggiatori.", "Segnala carico operativo e partenze critiche."],
        ["Engagement", "Azioni completate su quiz, sfide e contest.", "Misura partecipazione alle attivita."],
        ["Feedback", "Media e numero risposte per tappa; valori bassi in evidenza.", "Permette interventi puntuali su servizi e visite."],
    ], [1.15, 3.15, 2.2])
    heading(document, "Esperienza utente agenzia", 2)
    for item in [
        "La voce Analytics e disponibile nella navigazione del pannello agenzia per responsabile e agente autorizzato.",
        "Il filtro periodo offre 7, 30 e 90 giorni; il filtro partenza accetta una partenza dell'agenzia o tutte.",
        "La testata e i controlli rispettano logo e colore dell'agenzia con contrasto WCAG AA.",
        "La tabella partenze e scorrevole su schermi stretti; le schede KPI diventano a colonna su smartphone.",
        "Gli stati senza dati sono espliciti e non vengono rappresentati come valore zero ambiguo.",
    ]:
        bullet(document, item)
    heading(document, "Criteri di accettazione", 2)
    for item in [
        "Un utente agenzia non puo visualizzare indicatori di un'altra agenzia modificando query string o URL.",
        "Il cambio di periodo o partenza aggiorna coerentemente tutti i blocchi della pagina.",
        "Una risposta API Analytics fallita non blocca programma, documenti o navigazione del viaggiatore.",
        "Due invii equivalenti nella stessa sessione producono un solo evento conteggiabile.",
        "Il feedback mostra giornata, tappa, media e numero di risposte, ordinando prima le valutazioni peggiori nella giornata.",
    ]:
        bullet(document, item)


def add_use_cases(document: Document) -> None:
    heading(document, "Casi d'uso Analytics agenzia", 2)
    cases = [
        ("UC-ANA-001", "Aprire la dashboard Analytics", "Responsabile o agente autenticato", "La pagina mostra KPI e dati della sola agenzia attiva."),
        ("UC-ANA-002", "Filtrare per periodo", "Dashboard aperta", "7, 30 o 90 giorni aggiornano tutti gli indicatori."),
        ("UC-ANA-003", "Filtrare per partenza", "Almeno una partenza disponibile", "Sono incluse soltanto metriche della partenza selezionata."),
        ("UC-ANA-004", "Verificare il funnel di adozione", "Viaggiatori invitati", "Invitati, attivati, attivi e percentuali sono coerenti."),
        ("UC-ANA-005", "Registrare una consultazione programma", "Viaggiatore autenticato", "La prima apertura della giornata nella sessione incrementa programme_view una sola volta."),
        ("UC-ANA-006", "Registrare l'uso documenti", "Documento visibile al gruppo", "Elenco e download sono registrati; il download contribuisce agli utenti documenti."),
        ("UC-ANA-007", "Misurare richieste di assistenza", "Chat operativa disponibile", "I messaggi del viaggiatore compaiono nel conteggio della partenza."),
        ("UC-ANA-008", "Misurare engagement", "Quiz, sfide o contest disponibili", "Le azioni completate entrano nel totale senza contaminare altri gruppi."),
        ("UC-ANA-009", "Analizzare feedback per tappa", "Feedback inviati", "Sono mostrati giornata, tappa, media, risposte e segnalazione sotto 3/5."),
        ("UC-ANA-010", "Gestire assenza di dati", "Periodo senza eventi", "La pagina usa stati vuoti e trattini per medie non disponibili."),
        ("UC-ANA-011", "Verificare isolamento tenant", "Due agenzie con dati", "Nessun filtro o identificatore espone dati cross-tenant."),
        ("UC-ANA-012", "Verificare resilienza tracking", "API temporaneamente non disponibile", "L'app viaggiatore resta utilizzabile e l'evento puo essere ritentato."),
        ("UC-ANA-013", "Verificare idempotenza e anti-replay", "Evento gia inviato", "Retry o replay semantico non aumentano il conteggio; il limite orario blocca abuso."),
        ("UC-ANA-014", "Verificare accessibilita e responsive", "Colore agenzia chiaro o scuro", "Contrasto AA, focus, lettura mobile e target interattivi restano conformi."),
    ]
    table(document, ["ID", "Scenario", "Precondizione", "Risultato atteso"], [list(row) for row in cases], [0.85, 1.6, 1.7, 2.35])
    heading(document, "Dati di prova minimi", 2)
    for item in [
        "Due agenzie, ciascuna con almeno una partenza, un gruppo e due viaggiatori attivi.",
        "Una partenza con piu giornate, documenti di gruppo, chat, quiz, sfide, contest e feedback su almeno due tappe.",
        "Eventi duplicati con stesso client_operation_id e con UUID differenti ma stessa semantica di sessione.",
        "Periodi con dati e senza dati, valori feedback sopra e sotto la soglia 3/5, colori agenzia chiari e scuri.",
    ]:
        bullet(document, item)


def update(path: Path) -> None:
    document = Document(path)
    remove_existing_appendix(document)
    common_intro(document)
    name = path.name.lower()
    if "architettura" in name:
        add_architecture(document)
    elif "logico" in name:
        add_logical(document)
    elif "fisico" in name:
        add_physical(document)
    elif "casi_uso" in name:
        add_use_cases(document)
    else:
        add_functional(document)
    document.save(path)


for target in DOCS:
    update(target)
    print(target.relative_to(ROOT))
