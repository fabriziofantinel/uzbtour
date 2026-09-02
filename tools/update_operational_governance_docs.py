from pathlib import Path
from shutil import copy2
from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"

FILES = [
    (DOCS / "architecture/SMF_Travel_Architettura_Soluzione_v1.1.docx", DOCS / "architecture/SMF_Travel_Architettura_Soluzione_v1.2.docx", "architecture"),
    (DOCS / "data-model/SMF_Travel_Modello_Logico_Dati_v1.3.docx", DOCS / "data-model/SMF_Travel_Modello_Logico_Dati_v1.4.docx", "logical"),
    (DOCS / "data-model/SMF_Travel_Modello_Fisico_Dati_v1.3.docx", DOCS / "data-model/SMF_Travel_Modello_Fisico_Dati_v1.4.docx", "physical"),
    (DOCS / "functional/SMF_Travel_Catalogo_Funzionale_v1.0.docx", DOCS / "functional/SMF_Travel_Catalogo_Funzionale_v1.1.docx", "functional"),
    (DOCS / "testing/SMF_Travel_Casi_Uso_Completi_v1.1.docx", DOCS / "testing/SMF_Travel_Casi_Uso_Completi_v1.2.docx", "usecases"),
    (DOCS / "testing/SMF_Travel_Rapporto_Completo_Test_v1.3_2026-09-02.docx", DOCS / "testing/SMF_Travel_Rapporto_Completo_Test_v1.4_2026-09-03.docx", "tests"),
]

DECISIONS = [
    "Le segnalazioni personali contengono esclusivamente indicazioni operative essenziali per l'assistenza durante il viaggio.",
    "Le segnalazioni sono leggibili solo dal responsabile dell'agenzia e dai Tour Leader assegnati alla partenza.",
    "Il trattamento richiede consenso esplicito; il contenuto viene eliminato automaticamente entro 30 giorni dal rientro.",
    "È vietata l'archiviazione di passaporti, copie di passaporto, diagnosi e documenti sanitari.",
    "Il Tour Leader è un ruolo operativo associato alla singola partenza, senza privilegi generali sull'agenzia.",
    "Il Tour Leader gestisce comunicazioni, presenze, emergenze, documenti e modifiche operative del programma.",
    "Non viene raccolta o memorizzata la posizione del viaggiatore.",
    "SOS resta una pagina semplice di contatti e istruzioni essenziali.",
    "Il capogruppo maggiorenne può gestire dati e consensi operativi dei minori del proprio gruppo.",
    "Le chat sono separate in tre ambiti: viaggio, gruppo e conversazione individuale con il viaggiatore.",
]

def add_title(doc, title):
    doc.add_page_break()
    doc.add_heading(title, level=1)
    p = doc.add_paragraph("Aggiornamento: 3 settembre 2026")
    p.runs[0].italic = True

def add_decisions(doc):
    doc.add_heading("Decisioni funzionali e di governance", level=2)
    for item in DECISIONS:
        doc.add_paragraph(item, style="List Bullet")

def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, value in enumerate(headers):
        table.rows[0].cells[index].text = value
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value

def update(doc, kind):
    add_title(doc, "Addendum - Operatività della partenza e Tour Leader")
    add_decisions(doc)
    if kind == "architecture":
        doc.add_heading("Architettura applicativa", level=2)
        add_table(doc, ["Componente", "Responsabilità", "Controllo"], [
            ("IAM e assegnazioni", "Assegna il ruolo Tour Leader a una partenza", "Least privilege e audit"),
            ("Operational control", "Presenze, segnalazioni e operazioni sul campo", "Autorizzazione per departure_id"),
            ("Operational chat", "Canali trip, group e traveler", "Isolamento di tenant, partenza, gruppo e soggetto"),
            ("Privacy retention", "Consenso e cancellazione automatica", "Scadenza a rientro + 30 giorni"),
        ])
        doc.add_paragraph("Le API non espongono dati sensibili tramite query dirette: le letture passano da funzioni SECURITY DEFINER con verifica dell'attore e della partenza. Le tabelle applicano RLS e indici con agency_id iniziale.")
    elif kind == "logical":
        doc.add_heading("Entità e relazioni aggiunte", level=2)
        add_table(doc, ["Entità", "Relazioni principali", "Regole"], [
            ("DepartureStaffAssignment", "Agency, Departure, User", "Ruolo tour_leader; assegnazione revocabile"),
            ("DepartureAttendance", "DepartureDay, Party, Traveler, User", "Un esito per viaggiatore e giornata"),
            ("TravelerOperationalAlert", "Traveler, Party, Departure, Consent", "Solo contenuto essenziale; retention 30 giorni"),
            ("OperationalMessage", "Departure e scope opzionale Party/Traveler", "Ambiti trip, group, traveler mutuamente esclusivi"),
        ])
    elif kind == "physical":
        doc.add_heading("Oggetti fisici introdotti dalla migrazione 135", level=2)
        add_table(doc, ["Schema.oggetto", "Chiave/indice", "Protezione"], [
            ("travel.departure_staff_assignments", "departure_id, user_id, role", "RLS; scrittura owner"),
            ("journey.departure_attendance", "departure_day_id, traveler_id", "RLS; owner/Tour Leader"),
            ("privacy.traveler_operational_alerts", "departure_id, traveler_id", "RLS; lettura owner/Tour Leader"),
            ("journey.operational_messages", "departure + scope + party/traveler + operation id", "Scope authorization e idempotenza"),
        ])
        doc.add_paragraph("Funzioni principali: is_departure_operator_v3, assign_tour_leader_v3, record_departure_attendance_v3, save_operational_alert_v3, purge_expired_operational_alerts_v3, list_operational_alerts_v3 e funzioni chat scoped.")
    elif kind == "functional":
        doc.add_heading("Nuove capacità funzionali", level=2)
        add_table(doc, ["Codice", "Capacità", "Attori"], [
            ("OPS-073", "Assegnazione Tour Leader alla partenza", "Responsabile agenzia"),
            ("OPS-074", "Registro presenze giornaliero", "Responsabile, Tour Leader"),
            ("OPS-075", "Segnalazioni operative essenziali con consenso", "Viaggiatore, capogruppo per minore"),
            ("OPS-076", "Consultazione protetta delle segnalazioni", "Responsabile, Tour Leader"),
            ("OPS-077", "Chat viaggio, gruppo e individuale", "Staff operativo, viaggiatori"),
            ("OPS-078", "Cancellazione automatica post-viaggio", "Sistema"),
        ])
    elif kind == "usecases":
        doc.add_heading("Casi d'uso aggiunti", level=2)
        rows = [
            ("UC-108", "Il responsabile assegna un Tour Leader a una sola partenza", "Il ruolo non abilita altre partenze o funzioni amministrative"),
            ("UC-109", "Il Tour Leader registra presenze per giornata", "Stato persistito e auditabile"),
            ("UC-110", "Un adulto registra una segnalazione con consenso", "Visibile solo a responsabile e Tour Leader"),
            ("UC-111", "Il capogruppo registra la segnalazione di un minore", "Ammesso solo per minore dello stesso gruppo"),
            ("UC-112", "Un utente tenta di inserire dati senza consenso", "Richiesta respinta"),
            ("UC-113", "Scade la retention post-rientro", "Testo cancellato e record marcato expired"),
            ("UC-114", "Chat di viaggio", "Messaggi visibili a tutti i partecipanti della partenza"),
            ("UC-115", "Chat di gruppo", "Messaggi invisibili agli altri gruppi"),
            ("UC-116", "Chat individuale", "Messaggi visibili solo al soggetto e allo staff autorizzato"),
            ("UC-117", "SOS senza localizzazione", "Contatti disponibili; nessuna richiesta geolocation"),
            ("UC-118", "Upload di passaporto o documento sanitario", "Funzione non prevista e contenuto non ammesso"),
        ]
        add_table(doc, ["ID", "Scenario", "Risultato atteso"], rows)
    elif kind == "tests":
        doc.add_heading("Stato verifica incremento operativo", level=2)
        add_table(doc, ["Area", "Esito", "Evidenza"], [
            ("Migrazione 135", "SUPERATO", "Dry-run e applicazione con gate completi"),
            ("Catalogo Neon", "SUPERATO", "85 tabelle, 68 RLS, 0 indici/vincoli invalidi"),
            ("Compilazione TypeScript", "SUPERATO", "tsc --noEmit"),
            ("Test unitari", "SUPERATO", "3 file, 6 test"),
            ("Casi UC-108..UC-118", "DA ESEGUIRE E2E", "Richiedono account e gruppi rappresentativi"),
        ])

for source, target, kind in FILES:
    copy2(source, target)
    document = Document(target)
    update(document, kind)
    document.save(target)
    print(target)
