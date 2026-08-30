from pathlib import Path
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

path=Path("docs/functional/SMF_Travel_Catalogo_Funzionale_v1.0.docx")
doc=Document(path)
if any(p.text.strip()=="11. Governance commerciale, variazioni e onboarding" for p in doc.paragraphs):
    print("already_updated");raise SystemExit(0)

doc.add_heading("11. Governance commerciale, variazioni e onboarding",level=1)
doc.add_paragraph("Questa sezione rende verificabili e commercialmente non ambigue le misure Analytics, le comunicazioni operative, i contenuti sensibili sul Paese e il primo accesso dei viaggiatori.")
doc.add_heading("11.1 Dizionario KPI versionato",level=2)
rows=[
 ("Tasso di attivazione","Account attivati / inviti personali consegnati x 100","Inviti annullati, duplicati tecnici o non consegnati","Eventi immediati; aggregato entro 15 minuti; consolidamento giornaliero"),
 ("Adozione per partenza","Viaggiatori con almeno una sessione significativa / viaggiatori abilitati alla partenza x 100","Impersonificazioni e membership rimosse","Entro 15 minuti; consolidamento giornaliero"),
 ("Consultazione programma","Viaggiatori distinti che aprono almeno una giornata / viaggiatori abilitati x 100","Refresh duplicati nella stessa sessione","Entro 15 minuti"),
 ("Utilizzo documenti","Viaggiatori distinti con apertura riuscita di almeno un documento / viaggiatori abilitati x 100","Tentativi falliti e download tecnici duplicati","Entro 15 minuti"),
 ("Apertura notifiche","Notifiche aperte / notifiche consegnate dal provider x 100","Invii falliti, subscription scadute","Entro 15 minuti"),
 ("Engagement","Viaggiatori con almeno una sfida completata, un ricordo inviato o una spesa registrata / viaggiatori abilitati x 100","Bozze e azioni annullate","Entro 15 minuti"),
]
t=doc.add_table(rows=1,cols=4);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
for i,v in enumerate(("KPI","Formula","Esclusioni","Aggiornamento")):t.rows[0].cells[i].text=v
for row in rows:
 c=t.add_row().cells
 for i,v in enumerate(row):c[i].text=v
for row in t.rows:
 row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
for cell in t.rows[0].cells:
 tcPr=cell._tc.get_or_add_tcPr();shd=OxmlElement("w:shd");shd.set(qn("w:fill"),"12313B");tcPr.append(shd)
 for run in cell.paragraphs[0].runs:run.font.bold=True
doc.add_paragraph("Ogni definizione possiede codice e versione. Una modifica della formula crea una nuova versione con data di efficacia; i dati storici conservano la definizione utilizzata al momento del calcolo.")

doc.add_heading("11.2 Registro variazioni operativo",level=2)
for text in [
 "Ogni modifica pubblicata al programma genera una voce immutabile con partenza, giornata o tappa, valore precedente, nuovo valore, autore, data, severita e testo comprensibile.",
 "Il viaggiatore visualizza gli aggiornamenti non letti, apre la giornata interessata e usa Ho letto. La presa visione è registrata individualmente e non equivale ad accettazione contrattuale.",
 "Le comunicazioni urgenti possono essere associate a push; consegna, apertura e presa visione rimangono eventi distinti e verificabili.",
]:doc.add_paragraph(text,style="List Bullet")

doc.add_heading("11.3 Governance dei contenuti sensibili",level=2)
doc.add_paragraph("Salute, sicurezza, documenti di ingresso, emergenze e ambasciata non sono considerati affidabili per il solo fatto di essere stati generati dall'AI.")
for text in [
 "Ogni contenuto conserva fonte, URL, data di acquisizione, data di verifica, scadenza, stato di revisione, revisore e disclaimer.",
 "Gli stati ammessi sono Da verificare, Approvato, Scaduto e Rifiutato. Solo l'approvazione umana rende il contenuto verificato.",
 "Le sezioni sensibili scadono dopo 7 giorni; le informazioni culturali e operative non sensibili dopo 90 giorni, salvo regole più restrittive.",
 "Le fonti preferenziali sono Viaggiare Sicuri, Ministero degli Affari Esteri, Ambasciate italiane, Ministero della Salute, OMS e autorità ufficiali locali.",
]:doc.add_paragraph(text,style="List Bullet")

doc.add_heading("11.4 Onboarding personale a basso attrito",level=2)
doc.add_paragraph("L'identità rimane basata su username univoco; l'email può essere condivisa da più componenti della famiglia. Ogni invito è personale, monouso, revocabile e collegato a un solo utente.")
for text in [
 "Accesso rapido: il viaggiatore apre il link personale, verifica lo username e accede senza creare subito una password.",
 "Accesso tradizionale: nello stesso flusso può scegliere una password e continuare a usare username e password.",
 "Il consumo di un invito non modifica né invalida gli inviti degli altri utenti che condividono la stessa email.",
 "La passkey resta un'evoluzione opzionale subordinata alla configurazione WebAuthn del dominio Cognito; il magic link è il percorso passwordless attualmente supportato.",
 "Dopo il primo accesso l'app propone l'installazione PWA con istruzioni dedicate a iOS e Android.",
]:doc.add_paragraph(text,style="List Bullet")

doc.save(path)
print(path)
