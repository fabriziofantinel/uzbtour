from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "architecture" / "SMF_Travel_Architettura_Soluzione_v1.0.docx"
OUTPUT = ROOT / "docs" / "architecture" / "SMF_Travel_Architettura_Soluzione_v1.1.docx"


REPLACEMENTS = [
    ("1.0", "1.1"),
    ("Candidate for Architecture Approval", "Consolidated As-Built - Ready for Architecture Approval"),
    ("Candidate", "Consolidated As-Built"),
    ("Neon Auth 0.5 beta", "Amazon Cognito Lite dietro adapter proprietario"),
    ("Neon PostgreSQL + Neon Auth", "Neon PostgreSQL + Amazon Cognito Lite"),
    ("serverless driver e Neon Auth SDK", "serverless driver Neon e adapter Cognito"),
    ("login Neon Auth e attivazione invito", "login Cognito username/password e attivazione invito"),
    ("cookie Neon Auth", "cookie di sessione Cognito"),
    ("endpoint Neon Auth", "endpoint Cognito"),
    ("Neon Auth", "Amazon Cognito Lite"),
    ("Neon Auth, cookie HttpOnly, RBAC applicativo; SDK Auth beta", "Cognito Lite, cookie HttpOnly, RBAC applicativo e adapter auth-provider"),
    ("SDK Neon Auth beta", "Dipendenza dal provider di autenticazione"),
    ("Pin versione, test upgrade, piano sostituzione", "Adapter proprietario, test di regressione e piano di sostituzione provider"),
    ("DLQ con almeno un messaggio.", "DLQ bonificata; allarme dedicato in stato OK."),
    ("DLQ con 1 messaggio", "DLQ vuota e allarme in stato OK"),
    ("Allarmi senza destinatario SNS evidenziato", "Cinque allarmi CloudWatch con destinatario SNS confermato"),
    ("notifica SNS non evidenziata nello stack", "topic SNS e sottoscrizione e-mail confermata"),
    ("SNS opzionale", "SNS operativo"),
    ("allineare a Nova 2 Lite o leggere da unico source", "Nova 2 Lite allineato tra IaC, runtime ed esempio ambiente"),
    ("P0 prima del sign-off operativo: chiudere DLQ, attivare destinatario allarmi, allineare modello Bedrock configurato.", "P0 chiusi: DLQ bonificata, destinatario SNS confermato e modello Bedrock allineato a Nova 2 Lite."),
    ("P2 per crescita: branch Neon per PR, metriche business/SLO, rate limiting, WAF, budget e cost allocation tag.", "Branch Neon per PR e gate di consolidamento sono operativi. WAF e rate limiting sono rinviati fino all'adozione di un dominio personalizzato; restano evolutive metriche business/SLO e cost allocation."),
]


def all_paragraphs(doc):
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in doc.sections:
        yield from section.header.paragraphs
        yield from section.footer.paragraphs


def replace_paragraph_text(paragraph, old, new):
    if old not in paragraph.text:
        return False
    text = paragraph.text.replace(old, new)
    for run in paragraph.runs:
        run.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = text
    else:
        paragraph.add_run(text)
    return True


def add_evidence_appendix(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    doc.add_paragraph("Appendice G - Evidenze finali di consolidamento", style="Title")
    doc.add_paragraph(
        "Stato verificato al 30 agosto 2026. Le evidenze seguenti sostituiscono le finding "
        "operative aperte riportate nelle sezioni storiche del documento."
    )

    rows = [
        ("Database Neon V3", "67 tabelle; 54 con RLS; 0 vincoli non validati; 0 indici invalidi; 0 tabelle tenant senza indice agency_id leading."),
        ("Riconciliazione", "Backfill core e operational applicati: 0 utenti mancanti, 0 movimenti di cassa mancanti, 0 mapping duplicati."),
        ("CI/CD", "Branch Neon effimero per pull request, migrazioni, smoke RLS cross-tenant e cleanup automatico; quality gate e Vercel verdi."),
        ("Identita", "Amazon Cognito Lite con accesso username/password; e-mail usata per inviti e reset; adapter lib/auth/auth-provider.ts."),
        ("AWS asincrono", "SQS Standard con MessageGroupId=agency_id, Lambda ARM64, DLQ bonificata e cinque allarmi CloudWatch in stato OK."),
        ("Notifiche", "Topic SNS operativo e sottoscrizione e-mail confermata."),
        ("AI", "Amazon Bedrock Converse, Nova 2 Lite, Tool Use forzato, validazione Zod e human-in-the-loop prima della pubblicazione."),
        ("Storage", "Cloudflare R2 privato in giurisdizione UE; Bucket Lock smf-travel-retention-30d attivo sul prefisso agencies/."),
        ("Edge security", "WAF e rate limiting Cloudflare rinviati: l'account non contiene una zona DNS e il dominio Vercel condiviso non e proteggibile direttamente."),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Ambito"
    table.rows[0].cells[1].text = "Evidenza consolidata"
    for left, right in rows:
        cells = table.add_row().cells
        cells[0].text = left
        cells[1].text = right

    doc.add_heading("Decisione aggiornata", level=2)
    doc.add_paragraph(
        "Architettura consolidata per la fase corrente. Non risultano finding P0 aperte. "
        "Il disaster recovery periodico resta un controllo operativo ricorrente; WAF e rate limiting "
        "sono un rischio accettato fino all'introduzione di un dominio personalizzato."
    )


def main():
    doc = Document(SOURCE)
    doc.core_properties.title = "SMF Travel - Architettura di Soluzione v1.1"
    doc.core_properties.subject = "Architettura as-built consolidata"
    doc.core_properties.comments = "Aggiornata dopo consolidamento Neon, AWS, SNS e Cloudflare R2"

    for paragraph in all_paragraphs(doc):
        for old, new in REPLACEMENTS:
            replace_paragraph_text(paragraph, old, new)

    # Il paragrafo descrittivo completo dell'identita richiede una sostituzione semantica.
    for paragraph in all_paragraphs(doc):
        if paragraph.text.startswith("Neon Auth gestisce identit"):
            text = (
                "Amazon Cognito Lite gestisce autenticazione username/password, inviti e reset. "
                "L'e-mail e un recapito e non costituisce l'identificativo di accesso. Il server "
                "risolve il subject Cognito nelle identita applicative V3; lib/auth/auth-provider.ts "
                "isola il dominio dal provider e consente una sostituzione controllata."
            )
            for run in paragraph.runs:
                run.text = ""
            (paragraph.runs[0] if paragraph.runs else paragraph.add_run()).text = text
        elif "La DLQ conteneva 1 messaggio" in paragraph.text:
            text = "La DLQ e stata bonificata; i cinque allarmi CloudWatch risultano in stato OK e mantengono actions enabled."
            for run in paragraph.runs:
                run.text = ""
            (paragraph.runs[0] if paragraph.runs else paragraph.add_run()).text = text
        elif "Il topic SNS e la subscription email sono condizionali" in paragraph.text:
            text = "Il topic SNS e operativo e la sottoscrizione e-mail ai.fabrizio.fantinel@gmail.com risulta confermata."
            for run in paragraph.runs:
                run.text = ""
            (paragraph.runs[0] if paragraph.runs else paragraph.add_run()).text = text
        elif "Deployment Vercel Production READY" in paragraph.text:
            text = (
                "Deployment Vercel Production operativo; stack CloudFormation AWS UPDATE_COMPLETE; "
                "Lambda Active; event source SQS Enabled; coda primaria e DLQ senza backlog. "
                "I cinque allarmi CloudWatch risultano OK con notifica SNS confermata."
            )
            for run in paragraph.runs:
                run.text = ""
            (paragraph.runs[0] if paragraph.runs else paragraph.add_run()).text = text

    add_evidence_appendix(doc)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
