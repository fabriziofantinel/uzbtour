from __future__ import annotations

from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "functional" / "SMF_Travel_Catalogo_Funzionale_v1.0.docx"

NAVY = "12313B"
TEAL = "137F7B"
TEAL_LIGHT = "E8F4F2"
GOLD = "C46A3A"
INK = "162D35"
MUTED = "60747A"
LIGHT = "F4F1EA"
PALE = "F8FAF9"
WHITE = "FFFFFF"
RED = "A33A32"
GREEN = "287A55"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_run(run, size=11, bold=False, italic=False, color=INK, font="Aptos"):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = rgb(color)


def shade(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: Iterable[float], indent_dxa=120):
    widths = list(widths)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(round(w * 1440) for w in widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(round(width * 1440)))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx])
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            tc_w.set(qn("w:w"), str(round(widths[idx] * 1440)))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_field(paragraph, instruction: str):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run(run, size=9, color=MUTED)


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.78)
    section.left_margin = Inches(0.86)
    section.right_margin = Inches(0.86)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in (
        ("Heading 1", 16, TEAL, 18, 10),
        ("Heading 2", 13, TEAL, 14, 7),
        ("Heading 3", 12, NAVY, 10, 5),
    ):
        style = doc.styles[style_name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Aptos"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    header = section.header
    p = header.paragraphs[0]
    p.text = "SMF TRAVEL  |  CATALOGO FUNZIONALE"
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(3)
    set_run(p.runs[0], size=8.5, bold=True, color=TEAL)
    p_pr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), "C9D9D6")
    borders.append(bottom)
    p_pr.append(borders)

    footer = section.footer
    t = footer.add_table(rows=1, cols=2, width=Inches(6.78))
    set_table_geometry(t, [5.65, 1.13], indent_dxa=0)
    t.rows[0].cells[0].paragraphs[0].text = "Versione 1.0 · 30 agosto 2026 · Documento funzionale as-built"
    t.rows[0].cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    t.rows[0].cells[1].paragraphs[0].add_run("Pag. ")
    add_field(t.rows[0].cells[1].paragraphs[0], "PAGE")
    for cell in t.rows[0].cells:
        for run in cell.paragraphs[0].runs:
            set_run(run, size=8, color=MUTED)


def add_para(doc, text, *, size=11, bold=False, italic=False, color=INK, align=None, before=0, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if align is not None:
        p.alignment = align
    set_run(p.add_run(text), size=size, bold=bold, italic=italic, color=color)
    return p


def add_bullets(doc, items: Iterable[str]):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_label_detail(doc, rows, widths=(1.62, 5.16), fill=PALE):
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in rows:
        cells = table.add_row().cells
        shade(cells[0], fill)
        set_run(cells[0].paragraphs[0].add_run(label), size=9.3, bold=True, color=TEAL)
        set_run(cells[1].paragraphs[0].add_run(value), size=9.5, color=INK)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_matrix(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        shade(table.rows[0].cells[idx], TEAL_LIGHT)
        set_run(table.rows[0].cells[idx].paragraphs[0].add_run(header), size=9.2, bold=True, color=NAVY)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_run(cells[idx].paragraphs[0].add_run(str(value)), size=8.9, color=INK)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_function(doc, f):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    set_run(p.add_run(f["id"]), size=9, bold=True, color=WHITE)
    shade_paragraph(p, TEAL)
    p.add_run("  ")
    set_run(p.add_run(f["title"]), size=12, bold=True, color=NAVY)
    add_label_detail(doc, [
        ("Scopo", f["purpose"]),
        ("Attori", f["actors"]),
        ("Prerequisiti", f["pre"]),
        ("Flusso principale", f["flow"]),
        ("Regole e varianti", f["rules"]),
        ("Esiti e dati", f["outputs"]),
        ("Copertura test", f["tests"]),
    ])


def shade_paragraph(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


FUNCTIONS = [
    {"id":"IAM-01","title":"Accesso con username e password","purpose":"Consentire a ogni persona un account individuale anche quando più persone condividono la stessa e-mail.","actors":"Superuser, responsabile, agente, viaggiatore.","pre":"Account attivo, username univoco, password Cognito e agenzia non sospesa.","flow":"L’utente inserisce username e password; il server autentica su Amazon Cognito, risolve l’identità applicativa in Neon e apre una sessione protetta mediante cookie HttpOnly.","rules":"L’e-mail non identifica l’account. Utente disabilitato, invito non accettato o agenzia sospesa bloccano l’accesso; i token non sono esposti a React.","outputs":"Sessione associata a persona, ruolo, agenzia, partenze e gruppi autorizzati; instradamento all’area corretta.","tests":"Credenziali valide/errate; username inesistente; account invitato/disabilitato; agenzia sospesa; rinnovo e scadenza sessione; due username sulla stessa e-mail."},
    {"id":"IAM-02","title":"Invito personale e attivazione account","purpose":"Creare credenziali distinte per responsabili, agenti e viaggiatori.","actors":"Superuser o utente agenzia autorizzato; invitato.","pre":"Username disponibile, e-mail valida e soggetto censito nel perimetro corretto.","flow":"Il sistema crea l’utente in stato invitato, genera un token monouso con scadenza, invia il link via e-mail e, all’accettazione, fa scegliere la password e attiva solo l’utente collegato al token.","rules":"Token indipendenti anche per la stessa casella; validità 14 giorni; un invito non attiva gli altri; token usato/scaduto non è riutilizzabile.","outputs":"Utente Cognito, identità collegata, stato attivo, membership coerente e sessione iniziale.","tests":"Tre inviti stessa e-mail; accettazione selettiva; token scaduto/usato/alterato; password non conforme; reinvio e-mail fallito."},
    {"id":"IAM-03","title":"Recupero password per username","purpose":"Reimpostare la password dello specifico account senza rendere l’e-mail un identificatore.","actors":"Qualsiasi utente censito.","pre":"Username noto; e-mail verificata sul relativo account Cognito.","flow":"L’utente indica lo username; riceve il codice/link sulla e-mail del relativo account; conferma codice e nuova password.","rules":"La risposta iniziale è generica per non rivelare l’esistenza dell’account; account differenti con stessa e-mail mantengono reset separati.","outputs":"Password aggiornata per un solo username.","tests":"Username valido/inesistente; codice errato/scaduto/riusato; password debole; stessa e-mail con più username."},
    {"id":"IAM-04","title":"Logout, rinnovo e blocco delle sessioni","purpose":"Garantire continuità controllata e revoca effettiva dell’accesso.","actors":"Tutti gli utenti; sistema.","pre":"Sessione esistente.","flow":"Il middleware rinnova i token quando possibile; il logout cancella i cookie; a ogni richiesta l’app ricontrolla identità e stato tenant.","rules":"La sospensione dell’agenzia o la disabilitazione dell’utente blocca anche una sessione già emessa alla richiesta successiva.","outputs":"Sessione rinnovata, chiusa o rifiutata con messaggio coerente.","tests":"Refresh valido/scaduto; logout; cookie alterato; cambio stato utente/agenzia durante la sessione."},
    {"id":"IAM-05","title":"Controllo disponibilità username","purpose":"Prevenire duplicati nell’intera applicazione.","actors":"Superuser e utenti agenzia durante la creazione di responsabili, agenti e viaggiatori.","pre":"Form di inserimento aperto.","flow":"L’interfaccia normalizza lo username e interroga il controllo globale prima del salvataggio.","rules":"Unicità globale e confronto normalizzato; il server ricontrolla in modo autoritativo per gestire richieste concorrenti.","outputs":"Esito disponibile/non disponibile e blocco del salvataggio se occupato.","tests":"Maiuscole/minuscole, spazi, caratteri non ammessi, duplicato simultaneo, username di utente rimosso."},
    {"id":"IAM-06","title":"Login come utente per assistenza","purpose":"Consentire verifica e supporto senza conoscere la password dell’utente.","actors":"Superuser; responsabile/agente per i viaggiatori consentiti.","pre":"Operatore autenticato e autorizzato; utente target attivo.","flow":"L’operatore seleziona un’identità, entra nell’esperienza target con banner di impersonazione e può terminare la modalità tornando al proprio profilo.","rules":"Il superuser può operare nel perimetro piattaforma; l’agente soltanto come viaggiatore dei propri viaggi; ogni accesso conserva identità reale e target per audit.","outputs":"Sessione delegata limitata e reversibile.","tests":"Target fuori agenzia/viaggio, utente sospeso, stop impersonazione, tentativo di concatenare impersonazioni, audit."},

    {"id":"ADM-01","title":"Cruscotto superuser","purpose":"Offrire una vista complessiva su agenzie, viaggi, utenti e stato della piattaforma.","actors":"Superuser.","pre":"Accesso con ruolo piattaforma.","flow":"La pagina aggrega conteggi e collegamenti alle anagrafiche agenzie e utenti.","rules":"Dati cross-tenant visibili solo al ruolo piattaforma; errori di lettura non devono esporre dettagli tecnici.","outputs":"Indicatori e accesso alle funzioni amministrative.","tests":"Conteggi con zero/molti record; accesso negato agli altri ruoli; agenzie sospese/in cancellazione."},
    {"id":"ADM-02","title":"Creazione agenzia e primo responsabile","purpose":"Attivare un nuovo tenant con la propria anagrafica e il proprio amministratore.","actors":"Superuser.","pre":"Dati obbligatori, username responsabile disponibile ed e-mail valida.","flow":"Il superuser inserisce dati agenzia e responsabile; l’app crea tenant, utente, ruolo owner e invito personale.","rules":"Un solo responsabile attivo; creazione atomica; l’eventuale errore di invio e-mail è segnalato senza creare duplicati.","outputs":"Agenzia attiva, responsabile invitato e tracciamento dell’operazione.","tests":"Campi mancanti/invalidi, username duplicato, e-mail condivisa, errore invio invito, doppio submit."},
    {"id":"ADM-03","title":"Modifica dati agenzia e contatti responsabile","purpose":"Mantenere aggiornata l’anagrafica fiscale, commerciale e di contatto.","actors":"Superuser.","pre":"Agenzia esistente.","flow":"L’operatore modifica nome, ragione sociale, partita IVA, codice fiscale, sede, città, CAP, provincia, paese, PEC, SDI, telefono, e-mail, sito e contatti del responsabile.","rules":"Username, nome e cognome del responsabile non sono modificabili in questo flusso; validazioni specifiche per formato e lunghezza.","outputs":"Anagrafica aggiornata e audit.","tests":"Campi lunghi, formati fiscali/contatto, valori vuoti, salvataggio concorrente, responsive layout."},
    {"id":"ADM-04","title":"Branding white-label dell’agenzia","purpose":"Applicare logo e colore dell’agenzia al pannello agenzia e all’esperienza viaggiatore.","actors":"Superuser.","pre":"Agenzia esistente; file immagine supportato.","flow":"L’operatore seleziona colore e carica logo locale; il browser invia il file direttamente allo storage privato mediante URL firmata; l’app registra il riferimento e aggiorna il tema.","rules":"PNG/JPG/WebP, massimo 2 MB; contrasto e leggibilità prevalgono sul colore scelto; testi e icone restano scuri; rimozione logo con fallback SMF.","outputs":"Identità visiva riutilizzata nelle pagine di agenzia e viaggio.","tests":"Formati/dimensioni, upload interrotto, colore estremo, logo rimosso, cache, tenant differenti."},
    {"id":"ADM-05","title":"Sospensione e riattivazione agenzia","purpose":"Bloccare immediatamente l’operatività senza cancellare dati.","actors":"Superuser.","pre":"Agenzia attiva o sospesa.","flow":"Il superuser conferma la variazione di stato; il sistema blocca o riabilita responsabili, agenti e viaggiatori al controllo successivo.","rules":"Operazione reversibile; i dati restano conservati; lo stato non dipende dalla sessione client.","outputs":"Tenant sospeso/attivo e audit.","tests":"Sessioni già aperte, riattivazione, accessi API diretti, import in corso."},
    {"id":"ADM-06","title":"Sostituzione del responsabile","purpose":"Trasferire la responsabilità dell’agenzia a una nuova persona.","actors":"Superuser.","pre":"Nuovo username disponibile; dati del nuovo responsabile completi.","flow":"Il sistema crea il nuovo utente, invia il link di attivazione, assegna il ruolo owner e rimuove il precedente responsabile dall’agenzia.","rules":"Non possono esistere due responsabili attivi; il passaggio è transazionale; l’ex responsabile non conserva accesso tenant.","outputs":"Nuovo owner invitato/attivo, vecchio owner rimosso e audit.","tests":"Invito fallito, username duplicato, concorrenza, vecchia sessione, stesso indirizzo e-mail."},
    {"id":"ADM-07","title":"Cancellazione asincrona dell’agenzia","purpose":"Rimuovere in sicurezza un tenant e i dati collegati evitando una transazione massiva.","actors":"Superuser; worker operativo.","pre":"Conferma esplicita e agenzia identificata.","flow":"La richiesta crea un job di cancellazione; il processo procede a tappe su accessi, viaggi, gruppi, dati operativi e riferimenti media, con possibilità di ripresa.","rules":"Nessuna DELETE sincrona a cascata sulla radice; oggetti protetti da retention sono purgati solo quando consentito; job idempotente.","outputs":"Stato in cancellazione, avanzamento e completamento verificabile.","tests":"Agenzia con molti dati, job interrotto/ripreso, retention, template condivisi, accessi durante il processo."},

    {"id":"AGY-01","title":"Cruscotto agenzia e catalogo viaggi","purpose":"Gestire rapidamente il portafoglio viaggi da desktop o tablet.","actors":"Responsabile e agente autorizzato.","pre":"Membership agenzia attiva.","flow":"La pagina mostra indicatori, elenco o schede, ricerca, filtri di stato, ordinamento e azioni contestuali sul viaggio.","rules":"Sono visibili solo i viaggi dell’agenzia; stati e azioni cambiano in base al ciclo di vita; l’elenco compatto privilegia una riga.","outputs":"Vista aggiornata e accesso a import, revisione, programma e gruppi.","tests":"Zero/molti viaggi, testi lunghi, mobile/tablet, filtri combinati, ordinamenti, stati in transizione."},
    {"id":"AGY-02","title":"Gestione agenti dell’agenzia","purpose":"Creare, invitare, elencare e rimuovere gli operatori.","actors":"Responsabile/admin agenzia.","pre":"Agenzia attiva; username disponibile.","flow":"L’utente apre la pagina Agenti, inserisce nome, username, e-mail e telefono; l’app crea membership e invito. Può successivamente rimuovere l’agente.","rules":"Il superuser non inserisce agenti da questo flusso; rimozione revoca il perimetro agenzia e non cancella dati storici prodotti.","outputs":"Elenco agenti, stato invito e audit.","tests":"Duplicati, e-mail condivisa, invito non consegnato, rimozione con sessione aperta, ultimo amministratore."},
    {"id":"AGY-03","title":"Chat operativa agenzia-gruppo","purpose":"Mantenere uno storico tracciabile delle comunicazioni operative.","actors":"Utenti agenzia autorizzati e viaggiatori del gruppo.","pre":"Partenza e gruppo attivi.","flow":"Le parti leggono la conversazione, inviano messaggi fino a 2.000 caratteri e aggiornano automaticamente lo storico.","rules":"Conversazione isolata per partenza e gruppo; invio idempotente mediante client_operation_id; polling periodico.","outputs":"Messaggio con mittente, ruolo e timestamp; eventuale notifica ai destinatari.","tests":"Accesso cross-group, doppio invio, testo vuoto/lungo, rete intermittente, ordinamento e aggiornamento ogni 15 secondi."},

    {"id":"TRP-01","title":"Creazione viaggio da preventivo","purpose":"Avviare il viaggio digitale partendo dal documento accettato dall’agenzia.","actors":"Responsabile o agente.","pre":"Titolo/periodo e file PDF, DOC o DOCX valido; agenzia attiva.","flow":"L’utente crea il viaggio, carica il preventivo originale nello storage e avvia l’interpretazione.","rules":"File massimo 20 MB; il preventivo deve descrivere il programma senza inventare elementi assenti; upload diretto con URL firmata.","outputs":"Viaggio in preparazione, documento originale, record import e job asincrono.","tests":"Formati validi/invalidi, file vuoto/grande, upload interrotto, doppio submit, date incoerenti."},
    {"id":"TRP-02","title":"Elaborazione asincrona OCR e AI","purpose":"Estrarre e strutturare automaticamente programma, località, servizi e dati commerciali.","actors":"Sistema SQS/Lambda/Textract/Bedrock.","pre":"Import registrato e documento accessibile.","flow":"Il job viene accodato per agenzia; il worker legge il file, segmenta/comprime se necessario, applica OCR quando serve, invoca Bedrock con output tipizzato e salva bozza, consumi e stato.","rules":"Fair sharing tramite agency_id; retry idempotenti e DLQ; massimo cinque blocchi Bedrock da 4,5 MB; fallback Textract per scansioni pesanti.","outputs":"Bozza normalizzata o errore recuperabile con chiavi di correlazione.","tests":"PDF testuale/scansione, DOCX con immagini, file grande, timeout, retry, DLQ, noisy neighbor, output AI non conforme."},
    {"id":"TRP-03","title":"Interpretazione del programma e dati commerciali","purpose":"Convertire il testo libero nel modello previsto dall’app.","actors":"Sistema AI; agente revisore.","pre":"Documento elaborabile.","flow":"L’AI estrae titolo, destinazione, date, sintesi, giornate, descrizioni estese, attività ordinate, trasferimenti, voli, treni, pasti, incontri, visite, pernottamenti e condizioni commerciali presenti.","rules":"Gli orari non sono obbligatori e sono conservati soprattutto per trasporti/prenotazioni; un giorno può avere più pernottamenti; ciò che non compare non è considerato incluso.","outputs":"TravelProgrammeDraft tipizzato con indicatori di incertezza.","tests":"Due città nello stesso giorno, due hotel, valle/regione vs città, pasto assente, testo ambiguo, date parziali."},
    {"id":"TRP-04","title":"Revisione umana della bozza","purpose":"Consentire all’agente di correggere l’interpretazione prima della pubblicazione.","actors":"Responsabile o agente.","pre":"Import nello stato da revisionare.","flow":"L’agente percorre le giornate, modifica descrizioni e attività, riordina o elimina voci, aggiunge/rimuove pernottamenti, completa orari previsti e valida entità dubbie.","rules":"Non si aggiungono nuove giornate dal revisore; pubblicazione bloccata se restano campi/validazioni obbligatorie; link di ricerca aiuta la verifica esterna.","outputs":"Bozza revisionata e salvabile più volte.","tests":"Modifiche parziali, navigazione tra giorni, due pernottamenti identici intenzionali, campi mancanti, refresh, concorrenza."},
    {"id":"TRP-05","title":"Salvataggio e ripresa della revisione","purpose":"Evitare perdita del lavoro dell’agente.","actors":"Agente revisore.","pre":"Bozza disponibile.","flow":"L’agente salva; l’app valida il payload e aggiorna la bozza senza pubblicarla; al successivo accesso riprende lo stato salvato.","rules":"Errori di validazione sono leggibili e non sovrascrivono l’ultima versione valida.","outputs":"Bozza persistita, stato e timestamp aggiornati.","tests":"Salvataggio valido/errato, perdita rete, doppio click, ricarica, conflitto di versione."},
    {"id":"TRP-06","title":"Pubblicazione controllata del programma","purpose":"Materializzare il viaggio solo dopo conferma esplicita dell’agente.","actors":"Agente; stored procedure di publish; worker di arricchimento.","pre":"Bozza completa e confermata; gate di pubblicazione superati.","flow":"L’app salva l’ultima revisione e avvia una transazione che crea versione programma, partenza, date, giornate, attività, pernottamenti e collegamenti alle anagrafiche; quindi avvia gli arricchimenti.","rules":"Human-in-the-loop obbligatorio; vincoli aggregati sono verificati in procedure transazionali; retry non duplica entità.","outputs":"Viaggio pubblicato/validato con partenza operativa e contenuti in generazione.","tests":"Publish riuscito, gate fallito, timeout dopo commit, retry, doppio click, tenant context, rollback."},
    {"id":"TRP-07","title":"Cataloghi condivisi di paesi, città, siti e hotel","purpose":"Riutilizzare entità geografiche e strutture senza duplicazioni incoerenti.","actors":"Sistema AI e catalogo; agente in revisione.","pre":"Programma in pubblicazione.","flow":"Il sistema normalizza nomi, distingue paese/città/regione/sito/hotel, corregge denominazioni verificabili e collega gli elementi alle giornate e attività.","rules":"La città predominante del giorno deriva dalla maggioranza delle visite, non solo dal pernottamento; una regione non viene forzata a città; coordinate e nomi sono riusabili.","outputs":"Riferimenti catalogo e mappa coerenti.","tests":"Omonimi, Fergana città/valle, hotel con refuso, giorno multicittà, coordinate mancanti, catalogo già esistente."},
    {"id":"TRP-08","title":"Generazione informazioni utili per paese","purpose":"Fornire contenuti pratici coerenti con la destinazione reale.","actors":"Bedrock e materializzatore contenuti.","pre":"Paese di destinazione identificato.","flow":"Il sistema genera fuso rispetto all’Italia, valuta/cambio, emergenze, ambasciata italiana, salute, documenti e sicurezza, abbigliamento, usi locali, mobilità, tradizioni e quadro demografico/politico/sociale.","rules":"Contenuti isolati per viaggio/paese, massimo sei curiosità; nessuna contaminazione con quiz o altri paesi; i due orologi sono calcolati a runtime.","outputs":"Schede informazioni utili e contatti telefonabili.","tests":"Uzbekistan/Vietnam/Belgio, sezioni mancanti, cambio non disponibile, timezone con ora legale, rigenerazione selettiva."},
    {"id":"TRP-09","title":"Generazione frasario locale","purpose":"Preparare frasi pratiche nella lingua pertinente alla destinazione.","actors":"Bedrock e viaggiatore.","pre":"Paese e lingue di riferimento noti.","flow":"L’AI crea categorie di frasi con italiano, lingua locale, pronuncia semplificata e contesto d’uso.","rules":"La lingua deve essere dinamica per paese; evitare che la lingua di un viaggio precedente contamini il nuovo contenuto.","outputs":"Frasario consultabile dal menu Informazioni/Altro.","tests":"Paese multilingue, accenti e caratteri non latini, pronuncia vuota, viaggio rigenerato."},
    {"id":"TRP-10","title":"Generazione quiz, missioni, bingo, giochi e contest","purpose":"Creare gamification coerente con programma, città, siti e tema del viaggio.","actors":"Bedrock, materializzatore e agente.","pre":"Giornate, località e siti pubblicati.","flow":"Il sistema produce dieci quiz per giorno, cinque missioni giornaliere, cartella bingo, giochi semplici e tema/criteri dei contest fotografici.","rules":"Domande e sfide devono riferirsi a elementi realmente visitati; niente duplicati o risposte ovvie nel testo; fallback deterministico se la generazione non raggiunge i minimi.","outputs":"Attività approvate e associate a template/giornate.","tests":"Giorno senza siti, siti multipli, contenuti duplicati, quantità esatta, lingua, rigenerazione e tempi di elaborazione."},
    {"id":"TRP-11","title":"Download preventivo originale e normalizzato","purpose":"Conservare la fonte e produrre un documento completo nel formato standard dell’app.","actors":"Utente agenzia.","pre":"Import esistente; per il normalizzato, revisione/pubblicazione completata.","flow":"Dalla pagina programma l’utente scarica il file originale oppure il DOCX ricostruito con anagrafica viaggio, programma, servizi e dati commerciali.","rules":"Il documento normalizzato segue il template agenzia e non contiene solo il programma; nomi file e autorizzazioni sono tenant-scoped.","outputs":"Download del preventivo originale e del preventivo normalizzato.","tests":"File mancante, formati diversi, dati lunghi, più pernottamenti, autorizzazione cross-tenant, resa Word/PDF."},
    {"id":"TRP-12","title":"Modifica del programma pubblicato","purpose":"Aggiornare il programma operativo senza ripetere la pubblicazione AI.","actors":"Utente agenzia autorizzato.","pre":"Partenza pubblicata e programma accessibile.","flow":"L’utente apre Programma, modifica descrizione, città, attività e pernottamenti, aggiunge/rimuove/riordina elementi e salva una giornata alla volta.","rules":"Non è disponibile Ripubblica; le modifiche valgono sul programma collegato; date e giornata devono appartenere alla partenza.","outputs":"Programma operativo aggiornato e visibile ai viaggiatori.","tests":"Salva giornata, ID giornata errato, riordino, cancellazione, due hotel, ticket collegati, concorrenza."},
    {"id":"TRP-13","title":"Gestione variazioni operative","purpose":"Conservare traccia di cancellazioni e riprogrammazioni senza perdere audit.","actors":"Agenzia e sistema.","pre":"Attività operativa esistente.","flow":"L’operatore può aggiornare lo stato o spostare un’attività mantenendo informazioni originali e nuova collocazione.","rules":"Le variazioni non cancellano fisicamente la storia; notifiche e programma devono riflettere lo stato effettivo.","outputs":"Attività cancellata/riprogrammata e potenziale avviso push.","tests":"Cambio giorno, cancellazione, ripristino, attività notturna, utenti offline."},
    {"id":"TRP-14","title":"Eliminazione viaggio","purpose":"Rimuovere un viaggio e i relativi asset nel perimetro corretto.","actors":"Utente agenzia autorizzato.","pre":"Conferma esplicita e viaggio appartenente all’agenzia.","flow":"L’app identifica template/partenze/import e rimuove o marca per rimozione record e file collegati.","rules":"Nessuna cancellazione cross-tenant; gli oggetti soggetti a retention seguono il purge differito; risposta idempotente se già eliminato.","outputs":"Viaggio non più visibile e conteggio asset trattati.","tests":"Viaggio con gruppi/documenti, già eliminato, ID non trovato, retention, richiesta ripetuta."},

    {"id":"GRP-01","title":"Creazione ed eliminazione gruppi","purpose":"Suddividere i partecipanti in nuclei operativi con dati privati separati.","actors":"Utente agenzia.","pre":"Viaggio con partenza pubblicata.","flow":"L’operatore crea un gruppo con nome/codice; può eliminarlo solo quando è vuoto.","rules":"Il termine funzionale è gruppo; ogni gruppo appartiene a una sola partenza; l’eliminazione non deve lasciare riferimenti orfani.","outputs":"Gruppo attivo oppure eliminato.","tests":"Nome duplicato/vuoto, gruppo con membri, partenza errata, concorrenza."},
    {"id":"GRP-02","title":"Inserimento e invito viaggiatore","purpose":"Associare una persona al gruppo e abilitarne l’accesso.","actors":"Utente agenzia.","pre":"Gruppo esistente; username disponibile.","flow":"L’operatore inserisce nome, username, e-mail, telefono e data di nascita; l’app crea profilo, membership e invito individuale.","rules":"La data di nascita determina adulto/minore dipendente; ogni token è personale; la stessa e-mail può servire più familiari.","outputs":"Viaggiatore invitato/attivo e collegato alla partenza/gruppo.","tests":"Minore, e-mail condivisa, username duplicato, invito fallito, membro già presente, dati opzionali."},
    {"id":"GRP-03","title":"Rimozione viaggiatore","purpose":"Revocare la partecipazione di una persona a uno specifico viaggio.","actors":"Utente agenzia.","pre":"Membership esistente.","flow":"L’operatore conferma; il sistema rimuove/disattiva la membership e aggiorna il gruppo.","rules":"Non deve cancellare impropriamente l’identità usata in altri viaggi; gestire il capogruppo prima della rimozione.","outputs":"Viaggiatore non più autorizzato sul viaggio.","tests":"Capogruppo, utente in più viaggi, dati finanziari/foto esistenti, doppia rimozione, sessione aperta."},
    {"id":"GRP-04","title":"Assegnazione e cessione capogruppo","purpose":"Garantire un referente adulto per sblocchi e operazioni di gruppo.","actors":"Agenzia; capogruppo corrente.","pre":"Almeno un membro adulto attivo.","flow":"L’agenzia assegna un adulto; il capogruppo può selezionare un altro adulto del proprio gruppo e cedergli la funzione.","rules":"Esattamente un organizer attivo per gruppo; un minore non può diventare capogruppo; aggiornamento transazionale e auditato.","outputs":"Nuovo capogruppo e precedente membro ordinario.","tests":"Minore, membro esterno, stesso leader, rimozione leader, due richieste concorrenti."},
    {"id":"GRP-05","title":"Partecipazione individuale ai giochi di viaggio","purpose":"Permettere a ogni viaggiatore di scegliere se concorrere anche contro gli altri gruppi.","actors":"Agenzia e singolo viaggiatore dal profilo Sfide.","pre":"Membership attiva.","flow":"Il flag viene impostato per il singolo viaggiatore e può essere modificato; i risultati restano sempre nel gruppo e, se abilitato, entrano anche nella classifica di viaggio.","rules":"Non è una proprietà del gruppo; la scelta influenza quiz, contest e classifiche senza duplicare punteggi.","outputs":"Preferenza personale e ambiti di ranking aggiornati.","tests":"Attiva/disattiva, più membri nello stesso gruppo, cambio dopo una gara, accesso non proprietario."},
    {"id":"GRP-06","title":"Privacy e consenso immagini per minori","purpose":"Formalizzare la base autorizzativa per foto, contest e ricordi dei minori.","actors":"Agenzia e responsabile adulto.","pre":"Viaggiatore classificato come dependent_minor.","flow":"L’operatore registra consenso concesso, negato o revocato; il sistema salva decisione e audit.","rules":"Assenza/negazione/revoca impedisce i flussi media che richiedono consenso; il minore non ha autonomia decisionale.","outputs":"Stato consenso verificabile e applicato alle funzioni foto.","tests":"Concesso/negato/revocato, modifica durante contest, adulto, tentativo upload senza consenso."},

    {"id":"DOC-01","title":"Documenti di viaggio per giornata e gruppo","purpose":"Distribuire voucher, biglietti e materiali soltanto ai destinatari corretti.","actors":"Utente agenzia; viaggiatori del gruppo.","pre":"Partenza, giornata e gruppo esistenti.","flow":"L’agenzia seleziona giornata e gruppo, inserisce titolo/descrizione, allega il file; i membri autorizzati lo vedono nella sezione Documenti.","rules":"Documento obbligatoriamente associato a giornata e gruppo; file privato in R2; nessuna visibilità agli altri gruppi.","outputs":"Metadati in Neon e asset privato scaricabile.","tests":"Gruppo/giorno errato, file grande/formato, cross-group, viaggio offline, descrizione lunga."},
    {"id":"DOC-02","title":"Download e consultazione documenti","purpose":"Rendere disponibili i documenti prima e durante il viaggio.","actors":"Agenzia e viaggiatore autorizzato.","pre":"Documento pronto e non cancellato.","flow":"L’utente apre la sezione Documenti, filtra implicitamente sul proprio viaggio/gruppo e scarica tramite endpoint autorizzato/URL firmata.","rules":"Link temporanei; metadati e contenuti entrano nel pacchetto offline quando selezionati; niente URL pubbliche permanenti.","outputs":"Visualizzazione o download del file.","tests":"Link scaduto, documento cancellato, offline già cache/non cache, cross-tenant/group, nome file speciale."},
    {"id":"DOC-03","title":"Cancellazione documenti da parte dell’agenzia","purpose":"Ritirare un documento non più valido.","actors":"Utente agenzia autorizzato.","pre":"Documento appartenente al viaggio dell’agenzia.","flow":"L’utente conferma la cancellazione; l’app applica soft delete ai metadati e revoca la distribuzione.","rules":"L’oggetto fisico può restare durante retention/legal hold; non deve essere più firmabile o visibile.","outputs":"Documento rimosso dalle liste e audit.","tests":"Documento in cache, link già firmato, cross-tenant, doppia cancellazione, retention."},
    {"id":"DOC-04","title":"Biglietti collegati alle attività","purpose":"Associare titoli di viaggio o voucher alla specifica tappa.","actors":"Utente agenzia.","pre":"Attività e giornata pubblicate.","flow":"Dal programma l’operatore carica un file sulla singola attività; il viaggiatore lo ritrova nel contesto previsto.","rules":"Attività, giorno e partenza devono essere coerenti; storage privato.","outputs":"Ticket/voucher collegato all’item operativo.","tests":"Item errato, sostituzione, download, modifica/cancellazione dell’attività."},

    {"id":"EXP-01","title":"Home e navigazione mobile del viaggio","purpose":"Offrire accesso rapido alle funzioni essenziali da smartphone.","actors":"Viaggiatore.","pre":"Accesso a una partenza e gruppo attivi.","flow":"La pagina applica logo/colore agenzia e propone Mappa, Programma, Documenti, Spese, Sfide e menu Altro con Ricordi, Informazioni utili, Frasario, SOS e Chat.","rules":"Touch target almeno 48 px, safe area iOS, testi/icone scuri, stato non espresso solo dal colore.","outputs":"Esperienza coerente con il viaggio e l’agenzia.","tests":"Schermi piccoli, zoom testo, landscape, tema molto chiaro/scuro, navigazione tastiera e lettore schermo."},
    {"id":"EXP-02","title":"Programma giornaliero ordinato","purpose":"Mostrare una scaletta chiara senza imporre orari inesistenti.","actors":"Viaggiatore.","pre":"Programma pubblicato.","flow":"L’utente seleziona/scorre i giorni, legge la descrizione sempre aperta e le attività nell’ordine previsto; orari appaiono solo dove presenti.","rules":"Nessuna sezione ridondante “dove sei”; hotel/visite non duplicati nei dettagli; swipe orizzontale e pull-to-refresh supportano l’uso mobile.","outputs":"Giornata corrente e sequenza operativa.","tests":"Giorni senza attività, più città/hotel, descrizione lunga, fuso/date, swipe, offline e aggiornamento programma."},
    {"id":"EXP-03","title":"Mappa operativa del giorno","purpose":"Localizzare visite, siti e pernottamenti pertinenti.","actors":"Viaggiatore.","pre":"Coordinate disponibili.","flow":"La mappa mostra i punti del giorno, consente selezione e apre indicazioni esterne; le tile già viste vengono conservate temporaneamente offline.","rules":"I punti devono derivare dal programma del viaggiatore; nessun tracciamento persistente della posizione; coordinate mancanti mostrano uno stato vuoto.","outputs":"Mappa e link di navigazione.","tests":"Punti multipli, coordinate errate/mancanti, offline, permesso posizione negato, itinerario di altro viaggio."},
    {"id":"EXP-04","title":"Informazioni utili e orologi","purpose":"Raccogliere indicazioni pratiche specifiche del paese.","actors":"Viaggiatore.","pre":"Contenuti paese materializzati.","flow":"La pagina mostra ora italiana e locale, cambio euro-valuta locale e nove sezioni informative con eventuali numeri cliccabili.","rules":"Paese, valuta e fuso derivano dalla partenza; nessun pulsante “Approfondisci”; fallback esplicito se una sezione è in aggiornamento.","outputs":"Schede utili consultabili anche offline se già sincronizzate.","tests":"Ora legale, viaggio multicountry, contenuti mancanti, valuta senza decimali, contaminazione da altro viaggio."},
    {"id":"EXP-05","title":"Frasario di viaggio","purpose":"Aiutare nelle interazioni quotidiane nella lingua locale.","actors":"Viaggiatore.","pre":"Frasi generate per la destinazione.","flow":"L’utente consulta categorie, traduzione italiana, frase locale e pronuncia.","rules":"Contenuto dinamico per paese/lingua; caratteri Unicode corretti; nessuna frase di altra destinazione.","outputs":"Frasi organizzate e disponibili offline dopo cache.","tests":"Lingue non latine, paese multilingue, contenuto vuoto, cambio partenza."},
    {"id":"EXP-06","title":"Meteo e contesto del giorno","purpose":"Fornire indicazioni meteorologiche utili alla giornata.","actors":"Viaggiatore; servizio meteo.","pre":"Località/data disponibili e rete per il primo caricamento.","flow":"L’app richiede il meteo per la località e presenta il dato nel contesto della giornata.","rules":"Il fallimento del provider non blocca il programma; dati marcati temporalmente e cacheabili.","outputs":"Previsione o fallback non invasivo.","tests":"Località omonima, provider indisponibile, offline, date lontane, unità metriche."},
    {"id":"EXP-07","title":"SOS e duty of care","purpose":"Rendere immediati i contatti e le istruzioni in emergenza.","actors":"Viaggiatore.","pre":"Numero emergenze/contenuti paese disponibili.","flow":"L’utente può chiamare il numero locale, aprire la chat agenzia e seguire una sequenza di azioni consigliate.","rules":"La posizione non è condivisa né conservata automaticamente; fallback 112 se manca un numero specifico.","outputs":"Chiamata/chat avviata e istruzioni visibili.","tests":"Numero mancante/non numerico, dispositivo senza telefonia, offline, paese differente."},

    {"id":"FIN-01","title":"Registrazione spesa di gruppo","purpose":"Tenere la cassa e ripartire i costi tra viaggiatori selezionati.","actors":"Viaggiatore del gruppo.","pre":"Gruppo attivo; importo, valuta e almeno un beneficiario.","flow":"L’utente inserisce descrizione, importo, valuta, giorno e partecipanti alla spesa; il sistema registra pagatore, cambio e quote.","rules":"Importi in unità minori/decimali controllati; isolamento per gruppo; client_operation_id evita duplicati; quote uguali tra selezionati nel flusso base.","outputs":"Spesa, shares, equivalente EUR e saldo aggiornato.","tests":"EUR/valuta locale, più partecipanti, arrotondamenti, importo zero/negativo, doppio invio, cross-group."},
    {"id":"FIN-02","title":"Totale speso in euro e totali per valuta","purpose":"Mostrare un totale comparabile indipendentemente dalla valuta di spesa.","actors":"Viaggiatore.","pre":"Spese e tassi disponibili.","flow":"L’app somma le spese in EUR e converte quelle locali usando il cambio applicato; mostra anche i totali originali per valuta.","rules":"Una spesa da 1 EUR più l’equivalente di 1 EUR in valuta locale produce 2 EUR; prelievi e cambi non sono automaticamente spese.","outputs":"Totale EUR, totale valuta locale e tasso mostrato.","tests":"Tassi diversi nel tempo, nessun tasso, arrotondamento, valuta del viaggio corretta, movimenti cancellati."},
    {"id":"FIN-03","title":"Prelievi e cambi valuta","purpose":"Tracciare disponibilità di contanti e tasso effettivamente applicato.","actors":"Viaggiatore del gruppo.","pre":"Giorno e valuta locale del viaggio.","flow":"L’utente registra prelievo o cambio con importi EUR/locali; l’app calcola e mette in evidenza il cambio applicato.","rules":"Nessuna voce commissione obbligatoria; valuta deriva dalla destinazione e non da viaggi precedenti; movimenti separati dalle spese.","outputs":"Movimento di cassa con autore, giorno e tasso.","tests":"EUR/locali invertiti, zero, valuta Vietnam/Uzbekistan, cambio implicito, doppio submit, offline."},
    {"id":"FIN-04","title":"Eliminazione spese, prelievi e cambi","purpose":"Correggere inserimenti errati.","actors":"Viaggiatore autorizzato del gruppo.","pre":"Movimento esistente e visibile.","flow":"L’utente seleziona Elimina; l’app verifica lo scope e aggiorna totali e saldi.","rules":"Impossibile eliminare movimenti di un altro gruppo; operazione idempotente e auditabile.","outputs":"Movimento rimosso e riepiloghi ricalcolati.","tests":"ID errato, doppio click, autore differente stesso gruppo, cross-group, offline."},
    {"id":"FIN-05","title":"Pareggio conti del gruppo","purpose":"Indicare chi deve ricevere o versare per bilanciare le spese.","actors":"Viaggiatori del gruppo.","pre":"Spese con pagatore e quote.","flow":"Il sistema calcola per ogni viaggiatore pagato meno quota dovuta; valori positivi indicano credito e negativi debito.","rules":"Calcolo in EUR, arrotondamenti coerenti e somma finale prossima a zero; solo membri del gruppo.","outputs":"Saldo individuale consultabile.","tests":"Un solo membro, pagatore escluso dalle quote, più valute, centesimi residui, membro rimosso."},
    {"id":"FIN-06","title":"Inserimenti finanziari offline e sincronizzazione","purpose":"Consentire operazioni sul campo senza connessione.","actors":"Viaggiatore; service worker.","pre":"PWA inizializzata e IndexedDB disponibile.","flow":"In assenza rete la mutazione viene accodata con UUIDv7; al ritorno online l’app/background sync invia in ordine, ritenta gli errori temporanei e mostra lo stato.","rules":"409 idempotente è considerato già acquisito; errori client definitivi sono rimossi e segnalati; nessun duplicato dopo retry.","outputs":"Coda pending/completed/failed e dato sincronizzato.","tests":"Offline prima/durante invio, chiusura app, retry 5xx/429, 4xx, duplicato, più operazioni ordinate."},

    {"id":"GAM-01","title":"Missioni giornaliere","purpose":"Stimolare esplorazione e partecipazione con cinque attività al giorno.","actors":"Viaggiatore; AI di validazione foto.","pre":"Giornata sbloccata e missioni generate.","flow":"Il viaggiatore sceglie una missione, allega una foto e una nota facoltativa; l’AI confronta immagine, tema e descrizione.","rules":"Esattamente cinque missioni; 10 punti se approvata; la foto non è valutata solo dal nome file; contenuti coerenti con il giorno.","outputs":"Tentativo queued/approved/rejected, motivazione e punteggio.","tests":"Foto pertinente/non pertinente, immagine non leggibile, upload grande, missione giorno errato, consenso minore."},
    {"id":"GAM-02","title":"Massimo due tentativi per missione","purpose":"Offrire una seconda possibilità senza tentativi illimitati.","actors":"Viaggiatore; AI.","pre":"Missione non ancora approvata.","flow":"Se la prima foto è rifiutata, l’utente può caricarne una seconda; se anche questa fallisce la missione si chiude senza punti.","rules":"Dopo approvazione o secondo rifiuto non sono ammessi altri upload; richieste concorrenti sono serializzate.","outputs":"attempt_number, attempts_remaining e stato finale.","tests":"Approvazione al primo/secondo, due rifiuti, terzo tentativo via API, doppio upload contemporaneo."},
    {"id":"GAM-03","title":"Bingo in stile tombola napoletana","purpose":"Trasformare la scoperta del viaggio in una cartella fotografica.","actors":"Viaggiatore; AI.","pre":"Cartella generata e viaggio attivo.","flow":"L’utente fotografa gli elementi della griglia 3x9 con 15 caselle; ogni casella usa lo stesso ciclo di validazione a due tentativi.","rules":"Layout 3 righe x 9 colonne; cinque caselle per riga; contenuti osservabili e pertinenti alla destinazione.","outputs":"Caselle approvate, progressione e combinazioni.","tests":"Geometria cartella, duplicati, validazione foto, due tentativi, viaggio diverso."},
    {"id":"GAM-04","title":"Punteggi ambo, terno, quaterna, cinquina e tombola","purpose":"Premiare combinazioni completate sulla stessa riga e l’intera cartella.","actors":"Sistema classifiche.","pre":"Caselle bingo approvate.","flow":"Il motore ricalcola la miglior combinazione per riga e la tombola quando tutte le caselle sono complete.","rules":"Ambo 5, terno 10, quaterna 20, cinquina 30 per riga; tombola 50 aggiuntivi; niente doppio conteggio improprio.","outputs":"Punteggio bingo e ranking aggiornati.","tests":"Ordine approvazioni, più righe, revoca, completa cartella, concorrenza."},
    {"id":"GAM-05","title":"Quiz giornalieri","purpose":"Valorizzare conoscenza di città e monumenti visitati.","actors":"Viaggiatore; capogruppo per lo sblocco.","pre":"Dieci domande approvate per la giornata e finestra temporale raggiunta.","flow":"L’utente seleziona risposte, invia e riceve punteggio; il sistema registra tentativo e risposte.","rules":"Dieci quiz al giorno; domande coerenti con siti/città del giorno; rilascio calcolato nel fuso della partenza e payload non anticipato offline.","outputs":"Risposte, score e classifica.","tests":"Quantità, correttezza, unlock timezone/ora alterata device, offline prima/dopo unlock, doppio invio."},
    {"id":"GAM-06","title":"Giochi semplici a tema","purpose":"Offrire attività leggere legate al viaggio.","actors":"Viaggiatore.","pre":"Giochi pubblicati.","flow":"L’utente svolge tre giochi facili, tra cui immagine da ricomporre e giochi di parole/ordine o domande visive, quindi verifica la soluzione.","rules":"Contenuti comprensibili da smartphone, coerenti con destinazione e senza schede vuote; un gioco riusa l’immagine dell’app originale.","outputs":"Esito, punti e ranking.","tests":"Tre giochi presenti, soluzione corretta/errata, reset, touch, immagini offline, accessibilità."},
    {"id":"GAM-07","title":"Contest fotografico: selezione e conferma","purpose":"Consentire a ogni viaggiatore di proporre fino a due foto prima della valutazione.","actors":"Viaggiatore.","pre":"Contest aperto e consenso media valido.","flow":"L’utente carica fino a due foto, può sostituirle o eliminarle finché non preme Conferma; dopo la conferma le proposte sono bloccate.","rules":"Massimo due slot; conferma esplicita; file privati; una sola submission confermata per contest/viaggiatore.","outputs":"Due candidature o meno in stato submitted.","tests":"Una/due foto, sostituzione, eliminazione, terza foto, doppia conferma, consenso revocato."},
    {"id":"GAM-08","title":"Valutazione AI e scelta della foto migliore","purpose":"Attribuire un punteggio trasparente e pubblicare una sola immagine.","actors":"Bedrock; sistema contest.","pre":"Submission confermata.","flow":"L’AI valuta ogni foto su composizione 0-25, qualità tecnica 0-20, storytelling 0-25, originalità 0-15 e aderenza al tema 0-15; il sistema seleziona il totale migliore.","rules":"Output strutturato e validato; pareggio risolto deterministicamente; le foto non selezionate restano private/non entrano in classifica.","outputs":"Score breakdown, nota AI, foto selected e stato scored/failed.","tests":"Una/due foto, pareggio, output fuori range, timeout/retry, foto non pertinente, errore modello."},
    {"id":"GAM-09","title":"Chiusura contest e classifiche temporali","purpose":"Congelare i risultati alle 06:00 del giorno successivo.","actors":"Scheduler EventBridge/Lambda; sistema ranking.","pre":"Contest del giorno esistente.","flow":"Un job periodico individua contest scaduti nel fuso della partenza, completa eventuali valutazioni e chiude il contest.","rules":"Chiusura idempotente; dopo close non sono ammesse modifiche; la miglior foto entra nella classifica gruppo e, se opt-in, viaggio.","outputs":"Contest closed, posizioni e vincitori.","tests":"Fusi/ora legale, job ripetuto, scoring pendente, invio alle 05:59/06:00, viaggio multicountry."},
    {"id":"GAM-10","title":"Classifiche e profilo Sfide","purpose":"Mostrare progressione personale, per categoria, gruppo e viaggio.","actors":"Viaggiatore.","pre":"Attività e punteggi disponibili.","flow":"Il profilo mostra totale, breakdown quiz/giochi/missioni/bingo/foto, badge e classifiche generale/categoria; l’utente gestisce anche l’opt-in ai giochi di viaggio.","rules":"Dati sempre limitati agli ambiti consentiti; i punteggi fotografici entrano solo dopo validazione/chiusura prevista.","outputs":"Ranking ordinati, badge e posizione utente.","tests":"Parità, zero punti, opt-in/out, gruppo vs viaggio, rimozione membro, aggiornamento concorrente."},

    {"id":"MEM-01","title":"Ricordi e diario di viaggio","purpose":"Raccogliere note, foto e momenti personali legati alle giornate.","actors":"Viaggiatore del gruppo.","pre":"Viaggio accessibile.","flow":"L’utente aggiunge contenuti alla giornata, consulta la raccolta ed elimina gli elementi consentiti.","rules":"Ricordi isolati per gruppo/utente secondo visibilità; file privati; supporto upload da fotocamera.","outputs":"Entry di diario e media associati.","tests":"Upload/eliminazione, giorno errato, cross-group, offline, consenso minore, metadati immagine."},
    {"id":"MEM-02","title":"Album finale PDF e condivisione","purpose":"Trasformare viaggio, tappe e ricordi in un prodotto condivisibile.","actors":"Viaggiatore.","pre":"Dati del viaggio e ricordi disponibili.","flow":"L’app genera un PDF del diario; su dispositivi compatibili usa la condivisione nativa, altrimenti consente download.","rules":"Includere solo contenuti autorizzati del gruppo; gestire assenza foto e testi lunghi.","outputs":"File PDF scaricato o condiviso.","tests":"Album vuoto/ricco, caratteri internazionali, immagini grandi, mobile share non disponibile, privacy."},
    {"id":"MEM-03","title":"Feedback sul programma","purpose":"Raccogliere valutazioni sulle giornate/attività anche con rete instabile.","actors":"Viaggiatore.","pre":"Giornata o attività visibile.","flow":"L’utente assegna una valutazione e invia; l’app salva subito online oppure accoda offline.","rules":"Operazione idempotente; valore entro scala prevista; scope sulla propria partenza/gruppo.","outputs":"Feedback persistito e stato sincronizzazione.","tests":"Valori limite, modifica, offline/retry, item non autorizzato, doppio invio."},

    {"id":"PWA-01","title":"Installazione PWA","purpose":"Usare SMF Travel come app standalone senza store.","actors":"Viaggiatore con browser compatibile.","pre":"HTTPS, manifest e service worker validi.","flow":"Android/desktop intercetta beforeinstallprompt e mostra banner; iOS Safari riceve istruzioni Condividi > Aggiungi alla schermata Home.","rules":"display standalone, portrait, start_url /viaggio, icone maskable e scorciatoie Programma/Spese/Documenti.","outputs":"Icona Home e avvio a tutto schermo.","tests":"Chrome Android, Safari iOS, prompt accettato/rifiutato, già installata, icone e start URL."},
    {"id":"PWA-02","title":"Consultazione offline","purpose":"Mantenere programma, documenti e informazioni essenziali in modalità aereo.","actors":"Service worker e viaggiatore.","pre":"Contenuti visitati o pacchetto offline scaricato.","flow":"App shell usa cache-first/stale-while-revalidate; dati viaggio network-first con fallback; documenti/media e tile già visualizzati sono recuperati dalla cache entro limiti.","rules":"Quote e scadenze evitano saturazione; contenuto sensibile resta nel contesto browser del dispositivo; le cache si aggiornano quando torna rete.","outputs":"Pagine e asset disponibili offline o messaggio esplicito se mai scaricati.","tests":"Primo accesso offline, cache completa/parziale, aggiornamento server, quota storage, logout su device condiviso."},
    {"id":"PWA-03","title":"Aggiornamento della PWA","purpose":"Applicare nuove versioni senza lasciare l’utente su codice incoerente.","actors":"Service worker e viaggiatore.","pre":"Nuovo worker installato.","flow":"L’app mostra un avviso discreto; l’utente seleziona Ricarica, il worker attiva skip waiting e la pagina si aggiorna.","rules":"Nessun reload continuo; preservare mutazioni offline pendenti.","outputs":"Versione client aggiornata.","tests":"Update con coda offline, più tab, dismiss, installazione iniziale, errore worker."},
    {"id":"PWA-04","title":"Notifiche push","purpose":"Richiamare l’utente per eventi rilevanti senza dipendere da chat esterne.","actors":"Viaggiatore; Web Push; scheduler.","pre":"Permesso browser concesso e subscription VAPID registrata.","flow":"L’utente attiva gli avvisi; il server salva endpoint e chiavi; il service worker mostra notifiche in background e apre la sezione corretta al tap.","rules":"Permesso richiesto contestualmente; subscription revocabile; endpoint scaduti vengono disattivati.","outputs":"Subscription e notifica ricevuta.","tests":"Concesso/negato, endpoint 404/410, più device, logout, tap, iOS PWA installata."},
    {"id":"PWA-05","title":"Promemoria programmati","purpose":"Inviare avvisi per partenza, quiz ed eventi operativi.","actors":"Scheduler e viaggiatore.","pre":"Subscription attiva e partenza pertinente.","flow":"Il processo pianificato individua partenze del giorno successivo, quiz da sbloccare e altri eventi; crea claim idempotenti e invia ai soli membri interessati.","rules":"Timezone della partenza; nessun invio duplicato; fallimenti registrati e ritentabili.","outputs":"Run sent/failed e notifiche consegnate.","tests":"Fuso, più device, job ripetuto, subscription revocata, utente rimosso, agenzia sospesa."},
    {"id":"PWA-06","title":"Accessibilità e preferenze inclusive","purpose":"Rendere l’app utilizzabile da persone con diverse capacità e familiarità digitale.","actors":"Tutti gli utenti.","pre":"Browser moderno.","flow":"La sezione Accessibilità e gli stili globali supportano navigazione tastiera, focus visibile, contrasto, zoom, target tattili e feedback non basati solo sul colore.","rules":"Obiettivo WCAG 2.2 AA; branding agenzia corretto quando non garantisce contrasto.","outputs":"Esperienza leggibile e operabile.","tests":"Tastiera, screen reader, zoom 200/400%, contrasto, reduced motion, messaggi errore, touch 48 px."},
]


def add_cover(doc):
    add_para(doc, "CATALOGO FUNZIONALE", size=10, bold=True, color=GOLD, align=WD_ALIGN_PARAGRAPH.CENTER, before=84, after=16)
    add_para(doc, "SMF Travel", size=31, bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    add_para(doc, "Descrizione completa delle funzionalità e base per i casi d’uso", size=15, color=TEAL, align=WD_ALIGN_PARAGRAPH.CENTER, after=28)
    add_para(doc, "Piattaforma SaaS multi-agenzia per trasformare un preventivo accettato in un’esperienza di viaggio digitale, operativa, personalizzata e coinvolgente.", size=11.5, italic=True, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=62)
    add_matrix(doc, ["Documento", "Versione", "Stato", "Data"], [["Catalogo funzionale", "1.0", "As-built per test", "30 agosto 2026"]], [2.35, 0.8, 1.55, 2.08])
    add_para(doc, "Destinatari", size=9, bold=True, color=GOLD, align=WD_ALIGN_PARAGRAPH.CENTER, before=22, after=4)
    add_para(doc, "Agenzie di viaggio · Analisti funzionali · Tester · Solution Architect · Partner commerciali", size=10.5, color=INK, align=WD_ALIGN_PARAGRAPH.CENTER, after=44)
    add_para(doc, "Il documento descrive il comportamento funzionale osservabile dell’applicazione e indica, per ogni capacità, i controlli minimi da trasformare in casi di test.", size=9.5, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_page_break()


def add_intro(doc):
    doc.add_heading("1. Scopo e modalità d’uso", level=1)
    add_para(doc, "Questo documento ha due obiettivi complementari: presentare in modo comprensibile ciò che SMF Travel offre e costituire il catalogo di riferimento dal quale derivare casi d’uso, scenari di collaudo e criteri di accettazione. La descrizione è organizzata per capacità funzionali numerate e riflette lo stato del software al 30 agosto 2026.")
    add_bullets(doc, [
        "Per la presentazione commerciale: usare le sezioni 2-4 per spiegare valore, utenti e ciclo di vita del viaggio.",
        "Per l’analisi funzionale: trasformare ogni scheda numerata in uno o più casi d’uso principali, alternativi e negativi.",
        "Per il collaudo: combinare i punti “Copertura test” con le matrici trasversali della sezione 8.",
        "Per il regression testing: mantenere invariati gli ID delle funzionalità e collegarli ai test automatici/manuali.",
    ])
    add_label_detail(doc, [
        ("Incluso", "Funzioni utente, automazioni AI, processi schedulati, regole di isolamento, comportamento offline e gestione degli errori osservabile."),
        ("Non incluso", "Dettaglio DDL, topologia infrastrutturale completa, procedure operative di disaster recovery e manuale tecnico API, già coperti dai documenti architetturali e dati."),
        ("Terminologia", "Nel prodotto si usa “gruppo”, non “famiglia”. Le vecchie denominazioni tecniche families/party possono permanere negli endpoint o nel database senza essere esposte all’utente."),
    ])

    doc.add_heading("2. Proposta di valore", level=1)
    add_para(doc, "SMF Travel accompagna l’intero passaggio dal preventivo già accettato al viaggio vissuto. L’agenzia carica il documento che già utilizza, l’intelligenza artificiale lo interpreta, l’agente corregge ciò che serve e pubblica un programma strutturato. Da quel momento ogni partecipante dispone sullo smartphone di programma, documenti, informazioni paese, contatti, spese, comunicazioni, ricordi e attività ludiche, anche con connettività debole.")
    add_matrix(doc, ["Beneficiario", "Valore offerto", "Risultato atteso"], [
        ["Agenzia", "Riduzione del lavoro manuale, programma unico e modificabile, distribuzione documenti e assistenza tracciata.", "Più viaggi digitalizzati con minori errori e meno messaggi dispersi."],
        ["Viaggiatore", "Companion mobile white-label, offline, informazioni contestuali, cassa di gruppo e gamification.", "Informazioni subito disponibili e maggiore coinvolgimento prima, durante e dopo il viaggio."],
        ["Superuser", "Gestione centralizzata di tenant, identità, branding e supporto tramite impersonazione controllata.", "Governance della piattaforma e assistenza più rapida."],
    ], [1.15, 3.37, 2.26])

    doc.add_heading("3. Attori e perimetri di accesso", level=1)
    add_matrix(doc, ["Attore", "Responsabilità", "Perimetro"], [
        ["Superuser", "Crea e governa agenzie, responsabili, branding, sospensioni, cancellazioni e assistenza.", "Intera piattaforma, con privilegi espliciti e audit."],
        ["Responsabile agenzia", "Amministra agenti e operatività dell’agenzia; crea, revisiona e mantiene viaggi.", "Una sola agenzia e i suoi dati."],
        ["Agente", "Gestisce viaggi, gruppi, viaggiatori, documenti e comunicazioni secondo il ruolo assegnato.", "Agenzia di appartenenza; impersonazione limitata ai viaggiatori dei propri viaggi."],
        ["Capogruppo", "Riferimento adulto; può trasferire il ruolo e supportare sblocchi/operazioni del gruppo.", "Proprio gruppo e partenza."],
        ["Viaggiatore adulto", "Consulta, registra dati, comunica e partecipa alle attività.", "Proprio viaggio e gruppo; eventuali classifiche di viaggio su opt-in."],
        ["Minore dipendente", "Partecipa nei limiti previsti; media subordinati al consenso.", "Proprio gruppo, senza autonomia di consenso o leadership."],
        ["Sistemi automatici", "OCR, AI, scheduler, push, sincronizzazione e cancellazioni differite.", "Scope tecnico derivato da agency_id, departure_id, party_id e identity."],
    ], [1.25, 3.36, 2.17])

    doc.add_heading("4. Ciclo di vita end-to-end", level=1)
    add_matrix(doc, ["Fase", "Evento", "Stato/risultato"], [
        ["1. Configurazione", "Il superuser crea l’agenzia, il responsabile e il branding.", "Tenant attivo e invito personale."],
        ["2. Import", "L’agente crea il viaggio e carica il preventivo.", "Viaggio in preparazione; job asincrono."],
        ["3. Interpretazione", "OCR e Bedrock estraggono la bozza strutturata.", "Bozza da revisionare oppure errore ritentabile."],
        ["4. Revisione", "L’agente controlla attività, città, servizi e pernottamenti.", "Bozza salvata e pronta alla conferma."],
        ["5. Pubblicazione", "La conferma materializza programma, partenza, cataloghi e contenuti.", "Viaggio validato/pubblicato."],
        ["6. Partecipanti", "L’agenzia crea gruppi, invita viaggiatori e assegna capogruppo.", "Accessi e privacy configurati."],
        ["7. Operatività", "Documenti, chat, modifiche programma, push e assistenza.", "Fonte unica e aggiornata per il gruppo."],
        ["8. Esperienza", "Programma, mappe, cassa, sfide, foto, ricordi e feedback.", "Utilizzo mobile online/offline."],
        ["9. Chiusura", "Contest e classifiche si consolidano; album finale disponibile.", "Memoria condivisibile del viaggio."],
    ], [1.15, 3.88, 1.75])


def add_catalog(doc):
    doc.add_page_break()
    doc.add_heading("5. Catalogo dettagliato delle funzionalità", level=1)
    add_para(doc, f"Il catalogo contiene {len(FUNCTIONS)} capacità. Ogni ID è stabile e può essere riportato nella matrice di tracciabilità dei casi d’uso e dei test.")
    domains = [
        ("IAM", "Identità, inviti e sessioni"), ("ADM", "Amministrazione della piattaforma"),
        ("AGY", "Operatività dell’agenzia"), ("TRP", "Viaggi, preventivi, AI e programma"),
        ("GRP", "Gruppi, viaggiatori e privacy"), ("DOC", "Documenti e biglietti"),
        ("EXP", "Esperienza del viaggiatore"), ("FIN", "Spese, cassa e pareggio"),
        ("GAM", "Sfide, quiz, giochi e contest"), ("MEM", "Ricordi, album e feedback"),
        ("PWA", "PWA, offline, push e accessibilità"),
    ]
    for prefix, title in domains:
        doc.add_heading(title, level=2)
        for f in [x for x in FUNCTIONS if x["id"].startswith(prefix + "-")]:
            add_function(doc, f)


def add_cross_cutting(doc):
    doc.add_page_break()
    doc.add_heading("6. Regole funzionali trasversali", level=1)
    rules = [
        ("RF-01", "Isolamento multi-tenant", "Ogni lettura e mutazione è limitata all’agenzia; gruppi, spese, ricordi, documenti e contest sono ulteriormente limitati a partenza e gruppo."),
        ("RF-02", "Unicità identità", "Lo username è globale e univoco; l’e-mail può essere condivisa e serve per inviti/recupero."),
        ("RF-03", "Human-in-the-loop", "L’AI prepara bozze e arricchimenti, ma il programma non viene pubblicato senza conferma dell’agente."),
        ("RF-04", "Idempotenza", "Import, pubblicazione, chat, finanza, media, scheduler e sync usano chiavi operative per rendere sicuri retry e doppio click."),
        ("RF-05", "Fonte unica", "Il programma pubblicato e le successive modifiche autorizzate costituiscono il riferimento comune per i viaggiatori."),
        ("RF-06", "Fusi e date", "Sblocchi, chiusure e notifiche usano il fuso della partenza; gli orologi mostrano Italia e destinazione."),
        ("RF-07", "Valute", "Importi originali e equivalente EUR sono distinti; ogni cambio applicato è visibile; movimenti di cassa non diventano automaticamente spese."),
        ("RF-08", "Media privati", "Documenti e foto non sono pubblici: metadati autorizzativi in Neon, oggetti in R2 e download temporanei firmati."),
        ("RF-09", "Minori", "Leadership solo adulti; i media dei minori richiedono consenso esplicito e revocabile."),
        ("RF-10", "Branding accessibile", "Logo e colore dipendono dall’agenzia, ma contrasto, testi scuri, focus e significato degli stati prevalgono."),
        ("RF-11", "Offline", "La consultazione usa cache; spese, cassa e feedback usano una coda locale con sincronizzazione e feedback all’utente."),
        ("RF-12", "Cancellazione sicura", "Le radici ad alto volume e i media con retention usano soft delete, processi asincroni e purge differito."),
    ]
    add_matrix(doc, ["ID", "Regola", "Descrizione verificabile"], rules, [0.75, 1.63, 4.40])

    doc.add_heading("7. Stati funzionali da coprire", level=1)
    add_matrix(doc, ["Entità", "Stati principali", "Transizioni critiche"], [
        ["Utente", "invited, active, disabled/removed", "invito→attivo; attivo→rimosso; sessione esistente dopo revoca"],
        ["Agenzia", "active, suspended, deleting/deleted", "sospensione reversibile; cancellazione asincrona a tappe"],
        ["Import", "uploaded, queued, processing, needs_review, failed, published", "retry idempotente; DLQ; publish dopo review"],
        ["Viaggio/partenza", "draft/preparing, to_review, validated/published, open/confirmed/in_progress, deleted", "materializzazione; aggiornamento programma; cancellazione"],
        ["Invito", "issued, accepted, expired/revoked", "uso singolo e indipendenza su e-mail condivisa"],
        ["Documento/media", "pending/uploading, ready, deleted", "upload diretto; autorizzazione; soft delete/retention"],
        ["Foto-prova", "queued/pending, approved, rejected, locked", "massimo due tentativi; punteggio solo approved"],
        ["Contest", "open, submitted, scoring, scored, closed/failed", "conferma blocca; chiusura 06:00; retry scoring"],
        ["Mutazione offline", "pending, complete, failed", "flush ordinato; retry; 409 idempotente"],
        ["Push", "active/revoked; run claimed/sent/failed", "revoca endpoint; job pianificato idempotente"],
    ], [1.23, 2.05, 3.50])


def add_test_matrices(doc):
    doc.add_page_break()
    doc.add_heading("8. Matrici per la progettazione dei casi d’uso", level=1)
    doc.add_heading("8.1 Matrice ruoli e autorizzazioni", level=2)
    add_matrix(doc, ["Operazione", "Superuser", "Agenzia", "Viaggiatore"], [
        ["Gestire agenzia/branding", "Sì", "No", "No"],
        ["Gestire agenti", "Lettura/supporto", "Responsabile/admin", "No"],
        ["Importare/revisionare/pubblicare", "Solo supporto impersonato", "Sì", "No"],
        ["Modificare programma pubblicato", "Solo supporto", "Sì", "No"],
        ["Gestire gruppi/viaggiatori/documenti", "Solo supporto", "Sì", "Solo profilo/leadership consentiti"],
        ["Spese/cassa/ricordi/sfide", "No salvo impersonazione", "Assistenza limitata", "Sì nel proprio gruppo"],
        ["Vedere classifiche viaggio", "Supporto", "Supporto", "Solo se regole/opt-in lo consentono"],
    ], [2.55, 1.42, 1.42, 1.39])

    doc.add_heading("8.2 Matrice isolamento dati", level=2)
    add_matrix(doc, ["Dato", "Scope minimo", "Test negativo obbligatorio"], [
        ["Anagrafiche/viaggi", "agency_id", "Utente agenzia A non legge/modifica agenzia B."],
        ["Gruppi/viaggiatori", "agency + departure + party", "Agente di un altro viaggio e viaggiatore di altro gruppo ricevono 403/404 coerente."],
        ["Documenti", "agency + departure + day + party", "URL/ID copiato in altro gruppo non permette metadata né contenuto."],
        ["Spese e cassa", "departure + party", "Totali e movimenti di gruppo A assenti in gruppo B."],
        ["Chat", "departure + party", "Nessuna lettura/invio su conversazione altrui."],
        ["Ricordi/foto", "departure + party + traveler/visibility", "Media altrui non firmabile; consenso minore applicato."],
        ["Classifiche", "party; departure solo opt-in", "Utente opt-out non compare nel ranking viaggio."],
    ], [1.8, 2.18, 2.80])

    doc.add_heading("8.3 Combinazioni tecniche da applicare a ogni caso d’uso", level=2)
    add_matrix(doc, ["Dimensione", "Valori/scenari"], [
        ["Connettività", "online stabile; lenta; offline prima dell’azione; perdita rete durante upload/invio; ritorno online"],
        ["Dispositivo", "desktop agenzia; tablet; Android Chrome; iPhone Safari/PWA; schermo piccolo e zoom"],
        ["Identità", "ruolo corretto; ruolo insufficiente; utente rimosso; agenzia sospesa; impersonazione"],
        ["Dati", "vuoto; minimo; massimo; testi lunghi; Unicode; duplicato; riferimento inesistente; record di altro tenant"],
        ["Concorrenza", "doppio click; due operatori; retry dopo timeout; scheduler ripetuto; stesso client_operation_id"],
        ["Tempo", "prima/esattamente/dopo unlock o close; timezone diversa; ora legale; date attraversano mezzanotte"],
        ["AI", "output valido; incompleto; schema errato; timeout; contenuto non pertinente; fallback; retry"],
        ["Media", "formato valido/non valido; grande; corrotto; upload interrotto; URL scaduta; consenso negato/revocato"],
        ["Accessibilità", "solo tastiera; screen reader; contrasto; focus; errori annunciati; reduced motion; target 48 px"],
    ], [1.40, 5.38])

    doc.add_heading("8.4 Tracciato consigliato del caso d’uso", level=2)
    add_matrix(doc, ["Campo", "Contenuto da compilare"], [
        ["ID e titolo", "UC-<dominio>-nnn e nome univoco; indicare gli ID funzionali coperti."],
        ["Obiettivo", "Risultato business osservabile."],
        ["Attori", "Attore primario, secondari e sistemi automatici coinvolti."],
        ["Precondizioni", "Stati, dati, permessi, configurazioni e connettività richiesti."],
        ["Trigger", "Evento che avvia il caso d’uso."],
        ["Flusso principale", "Passi numerati utente/sistema con risultato verificabile a ogni passaggio."],
        ["Flussi alternativi", "Varianti lecite: dati opzionali, stato già esistente, browser/device differente."],
        ["Eccezioni", "Validazione, autorizzazione, rete, storage, database, provider esterno e AI."],
        ["Postcondizioni", "Dati creati/modificati, stato, notifiche, audit e visibilità."],
        ["Criteri di accettazione", "Given/When/Then o elenco di asserzioni UI/API/database."],
        ["Dati di prova", "Tenant, viaggio, gruppo, utenti, valuta, fuso, file e date."],
        ["Priorità e tipo", "P0/P1/P2; funzionale, sicurezza, regressione, accessibilità, performance, offline."],
    ], [1.62, 5.16])

    doc.add_heading("8.5 Criterio di completezza della suite", level=2)
    add_bullets(doc, [
        "Ogni funzione IAM/ADM/AGY/TRP/GRP/DOC/EXP/FIN/GAM/MEM/PWA è collegata ad almeno un caso positivo e uno negativo.",
        "Ogni mutazione è provata con doppio invio, retry e accesso fuori scope.",
        "Ogni dato di gruppo è verificato con almeno due gruppi della stessa partenza e due agenzie differenti.",
        "Ogni funzione temporale è provata prima, al limite e dopo l’orario nel fuso della partenza.",
        "Ogni upload è provato per formato, dimensione, interruzione, autorizzazione e cancellazione/retention.",
        "Ogni flusso AI prevede output valido, non pertinente, incompleto, timeout e retry senza duplicati.",
        "L’esperienza viaggiatore è collaudata su smartphone reale, PWA installata, modalità aereo e ritorno online.",
        "Le pagine principali superano controlli WCAG 2.2 AA manuali e automatici pertinenti.",
    ])


def add_glossary(doc):
    doc.add_page_break()
    doc.add_heading("9. Glossario", level=1)
    add_matrix(doc, ["Termine", "Definizione funzionale"], [
        ["Agenzia / tenant", "Organizzazione cliente isolata che possiede utenti, viaggi, branding e dati."],
        ["Preventivo originale", "Documento PDF/DOC/DOCX caricato dall’agenzia e conservato come fonte."],
        ["Preventivo normalizzato", "DOCX ricostruito nel formato standard dopo la revisione."],
        ["Template/programma", "Contenuto editoriale del viaggio, versionato e riutilizzabile dalle partenze collegate."],
        ["Partenza", "Istanza datata del viaggio con giornate, gruppi e operatività."],
        ["Gruppo / party", "Nucleo di viaggiatori che condivide dati privati, spese, chat e classifiche."],
        ["Capogruppo / organizer", "Unico membro adulto che svolge funzioni di referente del gruppo."],
        ["AI con controllo umano", "Bedrock prepara o valuta contenuti, ma le decisioni editoriali critiche restano confermate dall’agente."],
        ["RLS", "Regole PostgreSQL che applicano l’isolamento dei dati anche a livello database."],
        ["PWA", "Applicazione web installabile e capace di funzionare parzialmente offline."],
        ["client_operation_id", "Identificativo univoco dell’operazione client usato per evitare duplicati nei retry."],
        ["Soft delete", "Rimozione logica dalla disponibilità senza cancellazione fisica immediata."],
    ], [1.62, 5.16])
    doc.add_heading("10. Conclusione", level=1)
    add_para(doc, "SMF Travel non è soltanto un programma di viaggio digitale: è una piattaforma operativa white-label che collega agenzia e partecipanti, trasforma documenti non strutturati in dati controllabili, protegge i perimetri di ogni gruppo e rende l’esperienza fruibile sul campo. Il catalogo funzionale e le matrici di questo documento consentono di costruire una suite di casi d’uso completa, tracciabile e ripetibile senza perdere le varianti più rischiose: multi-tenant, offline, fusi orari, valute, media privati, AI, processi schedulati e minori.")


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "SMF Travel - Catalogo Funzionale v1.0"
    doc.core_properties.subject = "Descrizione funzionale completa e base per casi d'uso"
    doc.core_properties.author = "SMF Travel"
    doc.core_properties.keywords = "SMF Travel, catalogo funzionale, casi d'uso, test, PWA, viaggi"
    add_cover(doc)
    add_intro(doc)
    add_catalog(doc)
    add_cross_cutting(doc)
    add_test_matrices(doc)
    add_glossary(doc)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
