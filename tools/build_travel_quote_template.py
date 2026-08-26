from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "templates" / "Modello_Preventivo_Viaggio_SMF_Travel.docx"

INK = "173C3C"
TEAL = "0B6462"
RUST = "C75B3B"
CREAM = "F7F1E7"
GOLD = "E9C77A"
LIGHT = "EEF5F3"
MUTED = "667170"
WHITE = "FFFFFF"
LINE = "CDD8D5"


def set_run_font(run, size=11, bold=False, color=INK, italic=False):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths_dxa, indent=120):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[min(index, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_text(cell, text, bold=False, color=INK, size=9, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    set_run_font(paragraph.add_run(text), size=size, bold=bold, color=color)


def add_table(doc, headers, rows, widths_dxa, header_fill=TEAL, font_size=9, indent=120):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        shade(table.rows[0].cells[index], header_fill)
        set_cell_text(table.rows[0].cells[index], header, bold=True, color=WHITE, size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_repeat_table_header(table.rows[0])
    for row_values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            set_cell_text(cells[index], str(value), size=font_size)
    set_table_geometry(table, widths_dxa, indent=indent)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.add_run(text)
    return paragraph


def add_callout(doc, title, body, fill=LIGHT):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_repeat_table_header(table.rows[0])
    shade(cell, fill)
    cell.text = ""
    p1 = cell.add_paragraph()
    p1.paragraph_format.space_after = Pt(3)
    set_run_font(p1.add_run(title), size=10, bold=True, color=TEAL)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(1)
    p2.paragraph_format.line_spacing = 1.15
    set_run_font(p2.add_run(body), size=9.5, color=INK)
    set_table_geometry(table, [9360])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    heading_tokens = {
        1: (16, TEAL, 18, 10),
        2: (13, TEAL, 14, 7),
        3: (12, INK, 10, 5),
    }
    for level, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def configure_section(section, landscape=False):
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Inches(11)
        section.page_height = Inches(8.5)
        section.top_margin = Inches(0.65)
        section.bottom_margin = Inches(0.65)
        section.left_margin = Inches(0.65)
        section.right_margin = Inches(0.65)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_1)
    run._r.append(instr_text)
    run._r.append(fld_char_2)


def configure_header_footer(section):
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    set_run_font(p.add_run("SMF TRAVEL  |  MODELLO PREVENTIVO STANDARD"), size=8, bold=True, color=MUTED)
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(0)
    set_run_font(fp.add_run("Modello v1.0  •  Pagina "), size=8, color=MUTED)
    add_page_number(fp)


def add_label_value_table(doc, rows):
    return add_table(doc, ["CAMPO", "VALORE"], rows, [2700, 6660], header_fill=TEAL, font_size=9.5)


def add_programme_page(doc, page_label):
    add_heading(doc, f"PROGRAMMA GIORNALIERO — {page_label}", 1)
    add_table(
        doc,
        ["CAMPO", "VALORE", "CAMPO", "VALORE"],
        [
            ["Numero giorno", "<<N>>", "Data (GG/MM/AAAA)", "<<DATA>>"],
            ["Titolo giornata", "<<TITOLO>>", "Paese", "<<PAESE>>"],
            ["Città principale", "<<CITTÀ>>", "Altre località", "<<LOCALITÀ, separate da virgola>>"],
        ],
        [1500, 3000, 1500, 3690],
        header_fill=RUST,
        font_size=8.5,
        indent=120,
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    set_run_font(p.add_run("Inserire una riga per ogni attività, nello stesso ordine in cui avverrà. Non accorpare visita, trasferimento e pasto nella stessa riga."), size=8.5, italic=True, color=MUTED)
    blank_rows = []
    activity_types = [
        "COLAZIONE", "TRASFERIMENTO", "VISITA", "PRANZO", "VISITA", "TRASFERIMENTO",
        "CENA", "TRASFERIMENTO", "PERNOTTAMENTO", "", "", "",
    ]
    for index, activity_type in enumerate(activity_types, start=1):
        blank_rows.append([index, activity_type, "", "", "<<TITOLO / SERVIZIO>>", "<<CITTÀ>>", "SÌ / NO / N.A.", "<<NOTE OPERATIVE / AUTISTA / TRENO / VOLO / INCLUSIONE>>"])
    add_table(
        doc,
        ["Ord.", "Tipo attività", "Inizio", "Fine", "Titolo / servizio", "Località", "Incluso", "Note"],
        blank_rows,
        [480, 1200, 650, 650, 2100, 1150, 760, 2700],
        header_fill=TEAL,
        font_size=8,
        indent=120,
    )
    add_table(
        doc,
        ["PERNOTTAMENTO DEL GIORNO", "CITTÀ", "CHECK-IN", "CHECK-OUT", "TRATTAMENTO", "NOTE"],
        [["<<NOME HOTEL oppure NESSUNO>>", "<<CITTÀ>>", "<<DATA>>", "<<DATA>>", "<<BB / HB / FB / RO>>", "<<TIPO CAMERA / LATE CHECK-OUT>>"]],
        [2700, 1300, 1100, 1100, 1100, 2380],
        header_fill=RUST,
        font_size=8.5,
        indent=120,
    )


def build():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    configure_section(doc.sections[0], landscape=False)
    configure_header_footer(doc.sections[0])

    # Cover / operating instructions — customer_pack pattern.
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(18)
    kicker.paragraph_format.space_after = Pt(4)
    set_run_font(kicker.add_run("MODELLO UFFICIALE PER AGENZIE"), size=10, bold=True, color=RUST)
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(5)
    set_run_font(title.add_run("Preventivo di viaggio"), size=30, bold=True, color=INK)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(18)
    set_run_font(subtitle.add_run("Struttura standard compatibile con l’importazione automatica in SMF Travel"), size=13.5, color=MUTED)

    add_table(
        doc,
        ["AGENZIA", "PREVENTIVO"],
        [["<<NOME AGENZIA>>", "<<CODICE PREVENTIVO>>"], ["<<REFERENTE E CONTATTI>>", "Versione <<N>> del <<GG/MM/AAAA>>"]],
        [4680, 4680],
        header_fill=TEAL,
        font_size=10,
    )
    add_callout(
        doc,
        "Perché usare questo modello",
        "Il preventivo Uzbekistan conteneva informazioni valide ma distribuite tra programma, tabelle hotel e testo descrittivo. Questo modello rende espliciti giorno, località, hotel, visite, pasti e trasferimenti, riducendo ambiguità e correzioni manuali dopo l’importazione.",
        fill=CREAM,
    )
    add_heading(doc, "Regole di compilazione", 1)
    rules = [
        ["1", "Non modificare i titoli delle sezioni e i nomi delle colonne."],
        ["2", "Usare una sola riga per attività e rispettare l’ordine cronologico reale."],
        ["3", "Scegliere il Tipo attività solo dall’elenco controllato riportato sotto."],
        ["4", "Usare date GG/MM/AAAA e orari locali HH:MM. Se l’orario non è noto, lasciare vuoto: non scrivere “mattina”, “pomeriggio” o “da confermare”."],
        ["5", "Scrivere sempre città e Paese con il nome completo; non usare abbreviazioni."],
        ["6", "Riportare ogni hotel sia nel riepilogo hotel sia nel pernottamento della giornata corrispondente."],
        ["7", "Per voli, treni e trasferimenti indicare partenza, arrivo, vettore/numero servizio e note operative in righe separate."],
        ["8", "Per colazione, pranzo e cena indicare sempre SÌ o NO nella colonna Incluso."],
    ]
    add_table(doc, ["#", "REGOLA"], rules, [600, 8760], header_fill=RUST, font_size=9)
    add_heading(doc, "Tipi attività ammessi", 2)
    type_text = "VOLO · TRENO · TRASFERIMENTO · COLAZIONE · PRANZO · CENA · VISITA · TEMPO_LIBERO · CHECK_IN · CHECK_OUT · PERNOTTAMENTO · ALTRO"
    add_callout(doc, "Valori controllati", type_text, fill=LIGHT)
    example = doc.add_paragraph()
    set_run_font(example.add_run("Esempio di una riga corretta: "), size=9.5, bold=True, color=TEAL)
    set_run_font(example.add_run("3 | TRASFERIMENTO | 18:30 | 19:15 | Hotel–ristorante | Samarcanda | SÌ | Autista: <<NOME>>, targa: <<TARGA>>"), size=9.5, color=INK)

    doc.add_page_break()
    add_heading(doc, "1. TESTATA DEL VIAGGIO", 1)
    add_label_value_table(doc, [
        ["Titolo del viaggio", "<<TITOLO COMMERCIALE>>"],
        ["Cliente / famiglia", "<<NOME CLIENTE O GRUPPO>>"],
        ["Paese o Paesi", "<<PAESE 1, PAESE 2>>"],
        ["Data inizio", "<<GG/MM/AAAA>>"],
        ["Data fine", "<<GG/MM/AAAA>>"],
        ["Numero giorni / notti", "<<GIORNI>> / <<NOTTI>>"],
        ["Numero viaggiatori", "<<ADULTI>> adulti · <<MINORI>> minori"],
        ["Fuso orario principale", "<<ESEMPIO: Asia/Tashkent>>"],
        ["Lingua della guida", "<<LINGUA oppure SENZA GUIDA>>"],
        ["Valuta del preventivo", "<<EUR / USD / ALTRO>>"],
    ])
    add_heading(doc, "2. QUOTAZIONE", 1)
    add_table(doc, ["VOCE", "IMPORTO", "VALUTA", "NOTE"], [
        ["Quota individuale", "<<IMPORTO>>", "<<EUR>>", "<<BASE CAMERA / SERVIZI>>"],
        ["Totale pratica", "<<IMPORTO>>", "<<EUR>>", "<<NUMERO PARTECIPANTI>>"],
        ["Acconto", "<<IMPORTO / PERCENTUALE>>", "<<EUR>>", "Scadenza <<GG/MM/AAAA>>"],
        ["Saldo", "<<IMPORTO>>", "<<EUR>>", "Scadenza <<GG/MM/AAAA>>"],
    ], [2600, 1600, 1100, 4060], header_fill=TEAL, font_size=9)
    add_heading(doc, "3. PARTECIPANTI", 1)
    add_table(doc, ["N.", "NOME E COGNOME", "DATA NASCITA", "TIPO DOCUMENTO", "NUMERO", "NOTE"], [
        [1, "<<VIAGGIATORE 1>>", "<<GG/MM/AAAA>>", "<<PASSAPORTO / CI>>", "<<NUMERO>>", ""],
        [2, "<<VIAGGIATORE 2>>", "<<GG/MM/AAAA>>", "<<PASSAPORTO / CI>>", "<<NUMERO>>", ""],
        [3, "<<VIAGGIATORE 3>>", "<<GG/MM/AAAA>>", "<<PASSAPORTO / CI>>", "<<NUMERO>>", ""],
    ], [500, 2200, 1300, 1700, 1600, 2060], header_fill=RUST, font_size=8.5)

    doc.add_page_break()
    add_heading(doc, "4. TRASPORTI PRINCIPALI", 1)
    p = doc.add_paragraph("Inserire una riga per ogni tratta. Non unire due voli in coincidenza nella stessa riga.")
    p.runs[0].italic = True
    p.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
    add_table(doc, ["TIPO", "DATA", "DA", "A", "PARTENZA", "ARRIVO", "VETTORE / N.", "INCLUSO", "NOTE"], [
        ["VOLO", "<<DATA>>", "<<AEROPORTO>>", "<<AEROPORTO>>", "<<HH:MM>>", "<<HH:MM +1>>", "<<COMPAGNIA / VOLO>>", "SÌ / NO", "<<SCALO / BAGAGLIO>>"],
        ["TRENO", "<<DATA>>", "<<STAZIONE>>", "<<STAZIONE>>", "<<HH:MM>>", "<<HH:MM>>", "<<TRENO / CLASSE>>", "SÌ / NO", "<<BIGLIETTO>>"],
        ["TRASFERIMENTO", "<<DATA>>", "<<LUOGO>>", "<<LUOGO>>", "<<HH:MM>>", "<<HH:MM>>", "<<AUTISTA / MEZZO>>", "SÌ / NO", "<<TARGA / CONTATTO>>"],
    ], [850, 900, 1050, 1050, 760, 760, 1250, 760, 1980], header_fill=TEAL, font_size=7.8)
    add_heading(doc, "5. RIEPILOGO HOTEL", 1)
    add_callout(doc, "Campo decisivo per l’importazione", "Inserire qui tutti gli hotel, anche quando nel programma giornaliero sono citati soltanto in una tabella esterna o nelle condizioni del preventivo.", fill=CREAM)
    add_table(doc, ["CHECK-IN", "CHECK-OUT", "HOTEL", "CITTÀ", "PAESE", "TRATTAMENTO", "CAMERA", "LINK / INDIRIZZO", "NOTE"], [
        ["<<DATA>>", "<<DATA>>", "<<NOME COMPLETO HOTEL>>", "<<CITTÀ>>", "<<PAESE>>", "<<BB/HB/FB/RO>>", "<<TIPO>>", "<<URL O INDIRIZZO>>", "<<LATE CHECK-OUT>>"],
        ["<<DATA>>", "<<DATA>>", "<<NOME COMPLETO HOTEL>>", "<<CITTÀ>>", "<<PAESE>>", "<<BB/HB/FB/RO>>", "<<TIPO>>", "<<URL O INDIRIZZO>>", ""],
        ["<<DATA>>", "<<DATA>>", "<<NOME COMPLETO HOTEL>>", "<<CITTÀ>>", "<<PAESE>>", "<<BB/HB/FB/RO>>", "<<TIPO>>", "<<URL O INDIRIZZO>>", ""],
        ["<<DATA>>", "<<DATA>>", "<<NOME COMPLETO HOTEL>>", "<<CITTÀ>>", "<<PAESE>>", "<<BB/HB/FB/RO>>", "<<TIPO>>", "<<URL O INDIRIZZO>>", ""],
    ], [900, 900, 1550, 900, 850, 900, 900, 1500, 960], header_fill=RUST, font_size=7.7)
    add_heading(doc, "6. DOCUMENTI ALLEGATI", 1)
    add_table(doc, ["DOCUMENTO", "PRESENTE", "NOME FILE / NOTE"], [
        ["Biglietti aerei", "SÌ / NO", "<<NOME FILE>>"],
        ["Biglietti ferroviari", "SÌ / NO", "<<NOME FILE>>"],
        ["Voucher hotel", "SÌ / NO", "<<NOME FILE>>"],
        ["Polizza assicurativa", "SÌ / NO", "<<NOME FILE>>"],
    ], [3000, 1400, 4960], header_fill=TEAL, font_size=9)

    # Landscape daily programme sheets.
    programme_section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(programme_section, landscape=True)
    configure_header_footer(programme_section)
    add_programme_page(doc, "COPIARE QUESTA PAGINA PER OGNI GIORNO")

    doc.add_page_break()
    add_programme_page(doc, "SECONDA PAGINA PRONTA ALL’USO")

    # Back to portrait for service conditions and acceptance.
    final_section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(final_section, landscape=False)
    configure_header_footer(final_section)
    add_heading(doc, "7. SERVIZI INCLUSI E NON INCLUSI", 1)
    add_table(doc, ["SERVIZIO", "INCLUSO", "DETTAGLIO"], [
        ["Voli internazionali", "SÌ / NO", "<<TRATTE / BAGAGLIO>>"],
        ["Voli interni", "SÌ / NO", "<<TRATTE / BAGAGLIO>>"],
        ["Treni", "SÌ / NO", "<<TRATTE / CLASSE>>"],
        ["Trasferimenti", "SÌ / NO", "<<PRIVATI / CONDIVISI / AUTISTA>>"],
        ["Guida", "SÌ / NO", "<<LINGUA / GIORNI>>"],
        ["Pernottamenti", "SÌ / NO", "<<NUMERO NOTTI / TRATTAMENTO>>"],
        ["Colazioni", "SÌ / NO", "<<NUMERO>>"],
        ["Pranzi", "SÌ / NO", "<<NUMERO>>"],
        ["Cene", "SÌ / NO", "<<NUMERO>>"],
        ["Ingressi ai siti", "SÌ / NO", "<<DETTAGLIO>>"],
        ["Assicurazione", "SÌ / NO", "<<COPERTURA / COMPAGNIA>>"],
        ["Mance / extra", "SÌ / NO", "<<DETTAGLIO>>"],
    ], [2800, 1300, 5260], header_fill=TEAL, font_size=9)
    add_heading(doc, "8. NOTE, CONDIZIONI E INFORMAZIONI UTILI", 1)
    add_label_value_table(doc, [
        ["Validità del preventivo", "<<DATA / NUMERO GIORNI>>"],
        ["Condizioni di annullamento", "<<TESTO>>"],
        ["Requisiti di ingresso", "<<PASSAPORTO / VISTO / ALTRO>>"],
        ["Farmaci / salute", "<<TESTO oppure VEDI ALLEGATO>>"],
        ["Connettività", "<<SIM / APP / COPERTURA>>"],
        ["Contatti di emergenza", "<<NUMERI / REFERENTI>>"],
        ["Altre note", "<<TESTO>>"],
    ])
    add_heading(doc, "9. REFERENTI OPERATIVI", 1)
    add_table(doc, ["RUOLO", "NOME", "TELEFONO", "EMAIL / APP", "DISPONIBILITÀ"], [
        ["Agente di viaggio", "<<NOME>>", "<<TELEFONO>>", "<<EMAIL>>", "<<ORARI>>"],
        ["Corrispondente locale", "<<NOME>>", "<<TELEFONO>>", "<<EMAIL / WHATSAPP>>", "<<ORARI>>"],
        ["Guida", "<<NOME>>", "<<TELEFONO>>", "<<LINGUA>>", "<<GIORNI>>"],
        ["Autista", "<<NOME>>", "<<TELEFONO>>", "<<TARGA / MEZZO>>", "<<GIORNI>>"],
    ], [1900, 1800, 1600, 2300, 1760], header_fill=RUST, font_size=8.5)
    add_heading(doc, "10. ACCETTAZIONE DEL PREVENTIVO", 1)
    add_callout(doc, "Conferma del cliente", "Il cliente dichiara di aver letto programma, servizi inclusi e non inclusi, condizioni e scadenze di pagamento.", fill=CREAM)
    add_table(doc, ["LUOGO E DATA", "NOME CLIENTE", "FIRMA CLIENTE", "FIRMA AGENZIA"], [["", "", "", ""]], [2200, 2400, 2380, 2380], header_fill=TEAL, font_size=9)

    # Accessibility and document properties.
    doc.core_properties.title = "Modello Preventivo di Viaggio - SMF Travel"
    doc.core_properties.subject = "Template standard per preventivi importabili in SMF Travel"
    doc.core_properties.author = "SMF Travel"
    doc.core_properties.keywords = "viaggio, preventivo, agenzia, itinerario, importazione"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
