from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))
from build_functional_catalog_doc import FUNCTIONS  # noqa: E402

OUTPUT = ROOT / "docs" / "testing" / "SMF_Travel_Casi_Uso_Completi_v1.0.docx"

NAVY = "12313B"
TEAL = "137F7B"
TEAL_LIGHT = "E8F4F2"
GOLD = "C46A3A"
INK = "162D35"
MUTED = "60747A"
PALE = "F8FAF9"
WHITE = "FFFFFF"
RED = "A33A32"

DOMAIN_NAMES = {
    "IAM": "Identità, inviti e sessioni",
    "ADM": "Amministrazione della piattaforma",
    "AGY": "Operatività dell’agenzia",
    "TRP": "Viaggi, preventivi, AI e programma",
    "GRP": "Gruppi, viaggiatori e privacy",
    "DOC": "Documenti e biglietti",
    "EXP": "Esperienza del viaggiatore",
    "FIN": "Spese, cassa e pareggio",
    "GAM": "Sfide, quiz, giochi e contest",
    "MEM": "Ricordi, album e feedback",
    "PWA": "PWA, offline, push e accessibilità",
}

PAGE_CASES = [
    ("VIS-01", "Pagina di login", "Tutti", "/login", "Username, password, visibilità password, errore inline, recupero password, reindirizzamento per ruolo.", "Desktop, smartphone, tastiera, zoom 200%, credenziali errate e agenzia sospesa."),
    ("VIS-02", "Attivazione e recupero account", "Invitato", "/attiva-account e flussi reset", "Token, username, nuova password, conferma, esito generico e collegamento scaduto.", "Token valido, scaduto, alterato, già usato; più username sulla stessa e-mail."),
    ("VIS-03", "Cruscotto superuser", "Superuser", "/admin", "Indicatori, navigazione Agenzie/Utenti, stato vuoto, errori e responsive.", "Zero/molti record, permesso negato, caricamento e indisponibilità dati."),
    ("VIS-04", "Gestione agenzie", "Superuser", "/admin/agenzie", "Creazione, espansione dettaglio, modifica dati, logo/colore, sostituzione responsabile, sospensione e cancellazione.", "Campi lunghi, file errato, colore chiaro, conferme distruttive, doppio click."),
    ("VIS-05", "Elenco utenti e login-come", "Superuser", "/admin/utenti", "Ricerca, ruoli, stato, target impersonazione, banner e ritorno alla sessione originaria.", "Target non autorizzato, sospeso, appartenente ad altro tenant e sessione scaduta."),
    ("VIS-06", "Dashboard agenzia ed elenco viaggi", "Responsabile/Agente", "/agenzia", "Branding, indicatori, filtri, ricerca, ordinamento, lista/schede, azioni coerenti con lo stato.", "Titoli lunghi, una sola riga, viewport desktop/tablet/mobile, zero/molti viaggi."),
    ("VIS-07", "Gestione agenti", "Responsabile", "/agenzia/agenti", "Elenco, nuovo agente, validazione username, invito ed eliminazione.", "Duplicato, e-mail condivisa, invio fallito, agente con attività e accesso fuori ruolo."),
    ("VIS-08", "Import e revisione preventivo", "Responsabile/Agente", "/agenzia/importazioni/[id]", "Stato elaborazione, giornate, attività, pernottamenti multipli, orari pertinenti, salvataggio e pubblicazione.", "Bozza incompleta, errore AI, ripresa, doppio pernottamento, testi lunghi e mobile."),
    ("VIS-09", "Programma pubblicato modificabile", "Responsabile/Agente", "/agenzia/viaggi/[id]/programma", "Tab Programma/Gruppi/Documenti/Preventivi, giornate, attività, cancellazione con motivo e salvataggio.", "ID non valido, concorrenza, attività cancellata, navigazione giorni e branding."),
    ("VIS-10", "Gruppi e viaggiatori", "Responsabile/Agente", "/agenzia/viaggi/[id] tab Gruppi", "Creazione/eliminazione gruppo, elenco viaggiatori, invito, rimozione, capogruppo e opt-in giochi.", "Minore capogruppo, gruppo non vuoto, username duplicato e isolamento tra gruppi."),
    ("VIS-11", "Documenti del viaggio", "Responsabile/Agente", "/agenzia/viaggi/[id]/documenti", "Selezione giornata e gruppo, descrizione, upload, elenco, download ed eliminazione.", "File non valido/grande, upload interrotto, URL scaduta, altro gruppo/tenant."),
    ("VIS-12", "Chat operativa", "Agenzia/Viaggiatore", "/agenzia/viaggi/[id]/chat e Chat viaggiatore", "Selettore gruppo, storico, invio, autore, timestamp, aggiornamento ed errore.", "Offline, doppio invio, messaggio lungo, altro gruppo e refresh concorrente."),
    ("VIS-13", "Home e programma viaggiatore", "Viaggiatore", "/viaggio", "Branding, viaggio corretto, giorno attivo, descrizione, scaletta senza vincolo orario, swipe e mappa.", "Più viaggi, programma vuoto, giornata su due città, geolocalizzazione negata e offline."),
    ("VIS-14", "Informazioni, frasi, documenti e SOS", "Viaggiatore", "/viaggio sezioni principali", "Orologi Italia/destinazione, cambio, nove sezioni paese, frasario locale, wallet documenti e contatti emergenza.", "Paese differente, dati assenti, traduzione non pertinente, documento altro gruppo e offline."),
    ("VIS-15", "Spese e cassa", "Viaggiatore", "/viaggio sezione Spese", "Totale EUR, valute locali, spesa, quote, cambio, prelievo, eliminazione, saldo e sync.", "Arrotondamenti, valuta errata, offline, doppio invio, altro gruppo e campi numerici estremi."),
    ("VIS-16", "Sfide e classifiche", "Viaggiatore", "/viaggio sezione Sfide", "Missioni, bingo, quiz, giochi, foto, classifiche, profilo, opt-in e cessione capogruppo.", "Prima/dopo sblocco, due tentativi, due foto, AI KO/timeout, opt-out e chiusura 06:00."),
    ("VIS-17", "PWA, installazione e accessibilità", "Viaggiatore", "shell globale e /accessibilita", "Banner installazione, istruzioni iOS, update disponibile, offline, push, safe-area, focus e target touch.", "iOS/Android, standalone, modalità aereo, ritorno online, contrasto colore chiaro, zoom e screen reader."),
]


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_run(run, size=11, bold=False, italic=False, color=INK, font="Calibri"):
    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), font)
    rpr.rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = rgb(color)


def shade(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths, indent_dxa=120):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    values = [round(w * 1440) for w in widths]
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(values)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for value in values:
        node = OxmlElement("w:gridCol")
        node.set(qn("w:w"), str(value))
        grid.append(node)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index])
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            tc_w.set(qn("w:w"), str(values[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def configure(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Heading 1", 16, TEAL, 18, 10),
        ("Heading 2", 13, TEAL, 14, 7),
        ("Heading 3", 12, NAVY, 10, 5),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    header = section.header.paragraphs[0]
    header.text = "SMF TRAVEL  |  CASI D'USO COMPLETI"
    set_run(header.runs[0], size=8.5, bold=True, color=TEAL)
    header.paragraph_format.space_after = Pt(3)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(footer.add_run("Versione 1.0 · 30 agosto 2026 · Catalogo di collaudo"), size=8, color=MUTED)


def add_para(doc, text, *, size=11, bold=False, italic=False, color=INK, align=None, before=0, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if align is not None:
        p.alignment = align
    set_run(p.add_run(text), size=size, bold=bold, italic=italic, color=color)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)


def add_matrix(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for index, header in enumerate(headers):
        shade(table.rows[0].cells[index], TEAL_LIGHT)
        set_run(table.rows[0].cells[index].paragraphs[0].add_run(header), size=9.2, bold=True, color=NAVY)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_run(cells[index].paragraphs[0].add_run(str(value)), size=8.8, color=INK)
    set_table_geometry(table, widths)
    return table


def add_label(doc, label, value, color=TEAL):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_together = True
    set_run(p.add_run(f"{label}: "), size=10, bold=True, color=color)
    set_run(p.add_run(value), size=10, color=INK)
    return p


def split_items(text):
    return [x.strip().rstrip(".") for x in re.split(r";|\.(?:\s+|$)", text) if x.strip()]


def domain_ui(prefix):
    values = {
        "IAM": "Campi con label persistenti, errori vicino al controllo, stato di caricamento, focus visibile e nessuna esposizione di token o password.",
        "ADM": "Layout desktop/tablet senza overflow, conferme chiare per azioni critiche, indicatori coerenti e controlli non disponibili nascosti o disabilitati.",
        "AGY": "Branding dell’agenzia, testi e icone scuri, tasti coerenti, elenco leggibile su una riga e alternativa a schede.",
        "TRP": "Stato di avanzamento riconoscibile, giornate ordinate, contenuti modificabili senza perdita dati e messaggi di errore operativi.",
        "GRP": "Gerarchia gruppo-viaggiatori immediata, capogruppo riconoscibile, azioni distruttive separate e nessun dato di altri gruppi.",
        "DOC": "Giornata, gruppo, descrizione, nome file e azioni sempre identificabili; progress upload e download accessibile.",
        "EXP": "Mobile-first, colore agenzia predominante con testo scuro, navigazione inferiore 48x48 px, safe-area e contenuto disponibile offline.",
        "FIN": "Importi, valuta e controvalore EUR non ambigui; tastiera numerica mobile; conferma, stato offline e sincronizzazione visibili.",
        "GAM": "Stato bloccato/aperto/completato evidente, punteggi leggibili, tentativi residui espliciti e nessuna risposta corretta anticipata.",
        "MEM": "Anteprime, visibilità, autore e stato upload chiari; media privati mai esposti con URL permanente.",
        "PWA": "Esperienza standalone, banner non invasivi, messaggi offline/update/push comprensibili, contrasto WCAG AA e reduced motion.",
    }
    return values[prefix]


def main_steps(f):
    steps = [
        f"L’attore accede alla funzione «{f['title']}» dal proprio percorso autorizzato.",
        f"Il sistema verifica identità, ruolo, tenant e prerequisiti: {f['pre']}",
    ]
    for item in split_items(f["flow"]):
        steps.append(item[0].upper() + item[1:] + ".")
    steps.extend([
        "Il sistema applica validazioni server-side, isolamento RLS e idempotenza dove la funzione modifica dati.",
        f"L’interfaccia aggiorna lo stato e rende osservabile l’esito: {f['outputs']}",
        "L’attore può ricaricare o riaprire la pagina senza duplicazioni né perdita delle informazioni confermate.",
    ])
    return steps


def negative_scenarios(f):
    base = split_items(f["tests"])
    return [
        *[f"Verificare: {item}." for item in base],
        "Ripetere l’azione con un’identità dello stesso ruolo ma appartenente a un’altra agenzia o a un altro gruppo: accesso negato senza fuga di informazioni.",
        "Inviare due volte la stessa richiesta o ritentare dopo un timeout: un solo effetto persistente e risposta coerente.",
        "Interrompere la rete durante l’azione: nessuno stato ambiguo; retry o messaggio recuperabile in base alla funzione.",
    ]


def add_case(doc, f):
    prefix = f["id"].split("-")[0]
    case_id = f"UC-{f['id']}"
    heading = doc.add_heading(f"{case_id} · {f['title']}", level=3)
    heading.paragraph_format.page_break_before = False
    add_label(doc, "Funzionalità coperta", f["id"])
    add_label(doc, "Obiettivo", f["purpose"])
    add_label(doc, "Attori", f["actors"])
    add_label(doc, "Precondizioni", f["pre"])
    add_label(doc, "Trigger", f"L’attore seleziona o avvia «{f['title']}»; per automazioni, scatta l’evento o l’orario previsto.")
    add_label(doc, "Priorità", "P0 se riguarda identità, isolamento, pubblicazione, denaro o sicurezza; altrimenti P1. Tipo: funzionale, sicurezza e regressione.")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    set_run(p.add_run("Flusso principale"), size=10.5, bold=True, color=NAVY)
    add_numbered(doc, main_steps(f))

    add_label(doc, "Regole di business", f["rules"])
    add_label(doc, "Requisiti della visualizzazione", domain_ui(prefix))

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    set_run(p.add_run("Varianti, errori e test negativi"), size=10.5, bold=True, color=RED)
    add_bullets(doc, negative_scenarios(f))

    add_label(doc, "Postcondizioni", f["outputs"])
    add_label(doc, "Audit e osservabilità", "Registrare attore reale, tenant, entità, operazione, esito e client_operation_id/job_id quando previsto; non registrare segreti, token o contenuti sensibili.")
    add_label(doc, "Criteri di accettazione", f"Il flusso principale produce «{f['outputs']}»; tutte le regole sono applicate lato server; i casi indicati in Copertura test danno un esito comprensibile e non modificano dati fuori scope.")
    add_label(doc, "Dati minimi di prova", "Due agenzie, due viaggi nella stessa agenzia, due gruppi nella stessa partenza, utenti di ruoli differenti, dato valido e non valido, stessa operazione ripetuta.")
    add_para(doc, "Esito test:  ☐ Non eseguito   ☐ Superato   ☐ Non superato   ☐ Bloccato     Evidenza/defect: ______________________________", size=9.5, color=MUTED, before=3, after=10)


def add_visual_case(doc, row):
    case_id, title, actor, route, content, variants = row
    doc.add_heading(f"UC-{case_id} · {title}", level=3)
    add_label(doc, "Attori", actor)
    add_label(doc, "Percorso", route)
    add_label(doc, "Obiettivo", "Verificare che ogni contenuto, stato e comando della pagina sia visibile, coerente, autorizzato e utilizzabile sui dispositivi previsti.")
    add_label(doc, "Elementi obbligatori", content)
    add_numbered(doc, [
        "Aprire la pagina con dati rappresentativi e verificare titolo, navigazione, branding, ordine di lettura e stato attivo.",
        "Eseguire ogni comando interattivo almeno una volta e verificare focus, caricamento, conferma, errore e aggiornamento dei dati.",
        "Ridimensionare a 1440, 1024, 768, 390 e 320 px; nessun contenuto essenziale deve sovrapporsi, uscire dal viewport o richiedere scroll orizzontale globale.",
        "Ripetere con tastiera, zoom 200%, contrasto elevato e preferenza reduced-motion; focus e messaggi devono essere percepibili.",
        "Verificare skeleton/stato vuoto/errore/retry, quindi ricaricare la pagina e controllare la persistenza corretta.",
    ])
    add_label(doc, "Varianti obbligatorie", variants)
    add_label(doc, "Criteri di accettazione", "Tutti gli elementi elencati sono presenti e leggibili; le azioni rispettano ruolo e scope; nessun overflow o testo bianco sul colore agenzia; target touch almeno 48x48 px nel percorso mobile; errori senza dettagli tecnici.")
    add_para(doc, "Esito test:  ☐ Non eseguito   ☐ Superato   ☐ Non superato   ☐ Bloccato     Evidenza/defect: ______________________________", size=9.5, color=MUTED, after=10)


def add_cover(doc):
    add_para(doc, "CATALOGO DI COLLAUDO", size=10, bold=True, color=GOLD, align=WD_ALIGN_PARAGRAPH.CENTER, before=74, after=16)
    add_para(doc, "SMF Travel", size=30, bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    add_para(doc, "Casi d’uso completi e copertura di ogni funzionalità e visualizzazione", size=15, color=TEAL, align=WD_ALIGN_PARAGRAPH.CENTER, after=26)
    add_para(doc, "Suite funzionale di riferimento per collaudo manuale, automazione, sicurezza multi-tenant, mobile, offline, AI e accessibilità.", size=11.5, italic=True, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=54)
    add_matrix(doc, ["Documento", "Versione", "Copertura", "Data"], [["Casi d’uso completi", "1.0", f"{len(FUNCTIONS)} funzioni + {len(PAGE_CASES)} viste", "30 agosto 2026"]], [2.15, 0.8, 1.85, 1.70])
    add_para(doc, "Destinatari", size=9, bold=True, color=GOLD, align=WD_ALIGN_PARAGRAPH.CENTER, before=24, after=4)
    add_para(doc, "Product Owner · Agenzie · QA/Test Manager · Sviluppatori · Solution Architect · Security e Compliance", size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_page_break()


def add_intro(doc):
    doc.add_heading("1. Scopo, perimetro e modalità d’uso", level=1)
    add_para(doc, "Questo documento converte l’intero catalogo funzionale as-built di SMF Travel in casi d’uso eseguibili. Copre tutte le 72 capacità note e aggiunge 17 casi di verifica visuale per le pagine e le aree di esperienza. I casi non sostituiscono i test tecnici di unità o integrazione: definiscono il comportamento osservabile che tali test devono dimostrare.")
    add_bullets(doc, [
        "Eseguire ogni caso principale almeno nel flusso positivo e nelle varianti negative indicate.",
        "Applicare sempre due agenzie e due gruppi per dimostrare isolamento e assenza di data leak.",
        "Per ogni mutazione provare doppio click, timeout, retry e identificatore idempotente.",
        "Per AI, storage, push ed e-mail raccogliere evidenze sia applicative sia del servizio esterno.",
        "Per il viaggiatore ripetere il percorso su smartphone reale, PWA installata, offline e ritorno online.",
    ])
    doc.add_heading("2. Convenzioni e ambiente minimo", level=1)
    add_matrix(doc, ["Area", "Fixture minima", "Evidenza richiesta"], [
        ["Tenant", "Agenzia A attiva, Agenzia B attiva, Agenzia sospesa", "Screenshot/UI, risposta API, query di controllo RLS"],
        ["Identità", "Superuser, responsabile, agente, 2 adulti, 1 minore, utenti con e-mail condivisa", "Sessione, ruolo, audit e accesso negato"],
        ["Viaggi", "Bozza, in elaborazione, da revisionare, pubblicato, in corso", "Stato UI e record materializzati"],
        ["Gruppi", "Due gruppi nella stessa partenza e capogruppo adulto", "Visibilità e isolamento di documenti/spese/chat"],
        ["Servizi", "Cognito/e-mail, R2, SQS/Lambda, Bedrock, Neon, push", "Correlation ID, log, stato job e assenza duplicati"],
        ["Dispositivi", "Desktop, Android Chrome, iPhone Safari/PWA", "Responsive, offline, safe-area e accessibilità"],
    ], [1.05, 3.12, 2.33])
    doc.add_heading("3. Regole di esecuzione trasversali", level=1)
    add_bullets(doc, [
        "Sicurezza: ogni caso di lettura o scrittura comprende un tentativo con ID appartenente a un altro tenant/gruppo.",
        "Privacy: token, password, chiavi, URL firmate scadute e dati dei minori non devono comparire nei log o nei messaggi.",
        "Tempo: casi con sblocco/chiusura vanno eseguiti prima, al limite e dopo l’orario nel fuso della partenza.",
        "Accessibilità: ordine focus, label, errore annunciato, contrasto 4,5:1, zoom e target touch sono criteri di accettazione, non controlli opzionali.",
        "Dati: verificare sempre persistenza, audit, stato dopo refresh e assenza di effetti duplicati.",
    ])


def add_traceability(doc):
    doc.add_page_break()
    doc.add_heading("6. Matrice di tracciabilità", level=1)
    rows = []
    for prefix, name in DOMAIN_NAMES.items():
        ids = [f["UC-{x['id']}"] if False else f"UC-{x['id']}" for x in FUNCTIONS if x["id"].startswith(prefix + "-")]
        rows.append([prefix, name, len(ids), ", ".join(ids)])
    rows.append(["VIS", "Pagine e visualizzazioni", len(PAGE_CASES), ", ".join(f"UC-{x[0]}" for x in PAGE_CASES)])
    add_matrix(doc, ["Dominio", "Area", "Casi", "Identificativi"], rows, [0.7, 1.75, 0.55, 3.5])
    doc.add_heading("7. Criterio di completamento", level=1)
    add_bullets(doc, [
        "Tutti gli 89 casi hanno esito e collegamento a evidenza o defect.",
        "Nessun P0 è non eseguito, bloccato o non superato prima del go-live.",
        "Le anomalie corrette sono rieseguite nel caso originario e nella suite di regressione del dominio.",
        "I casi cross-tenant, offline, AI, media e temporali sono eseguiti con fixture ripristinabili.",
        "Il report finale distingue difetto applicativo, problema dati, problema ambiente e limite del servizio esterno.",
    ])


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure(doc)
    doc.core_properties.title = "SMF Travel - Casi d'uso completi v1.0"
    doc.core_properties.subject = "Copertura funzionale e visuale completa per collaudo"
    doc.core_properties.author = "SMF Travel"
    doc.core_properties.keywords = "SMF Travel, casi d'uso, test, collaudo, PWA, AI, multi-tenant"
    add_cover(doc)
    add_intro(doc)

    doc.add_page_break()
    doc.add_heading("4. Casi d’uso funzionali", level=1)
    add_para(doc, f"Sono presenti {len(FUNCTIONS)} casi, uno per ogni capacità del catalogo funzionale. Gli identificativi mantengono il legame diretto tra requisito, test ed eventuale defect.")
    for prefix, name in DOMAIN_NAMES.items():
        doc.add_heading(name, level=2)
        for f in [x for x in FUNCTIONS if x["id"].startswith(prefix + "-")]:
            add_case(doc, f)

    doc.add_page_break()
    doc.add_heading("5. Casi d’uso delle pagine e visualizzazioni", level=1)
    add_para(doc, "Questi casi verificano l’esperienza completa delle pagine, includendo contenuti, stati, responsive, branding, accessibilità e azioni. Integrano i casi funzionali evitando che una funzione corretta lato API resti inutilizzabile nell’interfaccia.")
    doc.add_heading("Pagine e aree applicative", level=2)
    for row in PAGE_CASES:
        add_visual_case(doc, row)

    add_traceability(doc)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
