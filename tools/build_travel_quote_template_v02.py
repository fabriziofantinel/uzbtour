import argparse
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


INK = "173C3C"
TEAL = "0B6462"
RUST = "C75B3B"
CREAM = "F7F1E7"
MUTED = "667170"
WHITE = "FFFFFF"


def set_run_font(run, size=9, bold=False, color=INK, italic=False):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = tr_pr.find(qn("w:tblHeader"))
    if marker is None:
        marker = OxmlElement("w:tblHeader")
        tr_pr.append(marker)
    marker.set(qn("w:val"), "true")


def set_table_geometry(table, widths_dxa, indent=0):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[min(index, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_text(cell, value, *, size=8, bold=False, color=INK, center=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    set_run_font(paragraph.add_run(str(value)), size=size, bold=bold, color=color)


def make_table(doc, headers, rows, widths, *, header_fill=TEAL, font_size=8):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        shade(table.rows[0].cells[index], header_fill)
        set_cell_text(table.rows[0].cells[index], header, size=8, bold=True, color=WHITE, center=True)
    repeat_header(table.rows[0])
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            set_cell_text(cells[index], value, size=font_size, center=index in (0, 1, 4))
    set_table_geometry(table, widths)
    return table


def replace_paragraph_text(paragraph, old, new):
    if old not in paragraph.text:
        return False
    for run in paragraph.runs:
        if old in run.text:
            run.text = run.text.replace(old, new)
            return True
    paragraph.text = paragraph.text.replace(old, new)
    return True


def insert_after(anchor, element):
    anchor.addnext(element)


def build(reference: Path, output: Path):
    document = Document(reference)

    replacements = {
        "PROGRAMMA GIORNALIERO — COPIARE QUESTA PAGINA PER OGNI GIORNO": "3. PROGRAMMA GIORNALIERO - COPIARE QUESTA PAGINA PER OGNI GIORNO",
        "9. REFERENTI OPERATIVI": "6. REFERENTI OPERATIVI",
        "10. ACCETTAZIONE DEL PREVENTIVO": "7. ACCETTAZIONE DEL PREVENTIVO",
    }
    for paragraph in document.paragraphs:
        for old, new in replacements.items():
            replace_paragraph_text(paragraph, old, new)

    # Testata: il riepilogo breve alimenta la descrizione generale del viaggio.
    trip_header = document.tables[1]
    row = trip_header.add_row().cells
    set_cell_text(row[0], "Descrizione sintetica del viaggio", size=9)
    set_cell_text(row[1], "<<SINTESI COMMERCIALE DEL VIAGGIO>>", size=9)
    source_widths = []
    for cell in trip_header.rows[0].cells:
        tc_w = cell._tc.tcPr.find(qn("w:tcW"))
        source_widths.append(int(tc_w.get(qn("w:w"))) if tc_w is not None else 4680)
    set_table_geometry(trip_header, source_widths)

    # Scheda giornaliera: campi geografici espliciti.
    day_header = document.tables[3]
    set_cell_text(day_header.rows[1].cells[0], "Numero giorno", size=8)
    set_cell_text(day_header.rows[1].cells[1], "<<N>>", size=8)
    set_cell_text(day_header.rows[1].cells[2], "Data (GG/MM/AAAA)", size=8)
    set_cell_text(day_header.rows[1].cells[3], "<<DATA>>", size=8)
    set_cell_text(day_header.rows[2].cells[0], "Titolo giornata", size=8)
    title_cell = day_header.rows[2].cells[1].merge(day_header.rows[2].cells[3])
    set_cell_text(title_cell, "<<TITOLO DELLA GIORNATA>>", size=8)
    set_cell_text(day_header.rows[3].cells[0], "Città principale", size=8)
    set_cell_text(day_header.rows[3].cells[1], "<<CITTÀ>>", size=8)
    set_cell_text(day_header.rows[3].cells[2], "Paese", size=8)
    set_cell_text(day_header.rows[3].cells[3], "<<PAESE>>", size=8)
    set_table_geometry(day_header, [1600, 3800, 1500, 7068])

    description_paragraph = next(
        paragraph for paragraph in document.paragraphs
        if paragraph.text.strip().startswith("Descrizione estesa della giornata")
    )
    description_paragraph.clear()
    description_paragraph.paragraph_format.space_before = Pt(8)
    description_paragraph.paragraph_format.space_after = Pt(4)
    set_run_font(description_paragraph.add_run("Descrizione estesa della giornata"), size=9, bold=True, color=TEAL)

    old_activity = document.tables[4]

    description_table = make_table(
        document,
        ["DESCRIZIONE ESTESA"],
        [["<<TESTO COMPLETO DEL PROGRAMMA DELLA GIORNATA>>"]],
        [13968],
        header_fill=RUST,
        font_size=8.5,
    )
    insert_after(description_paragraph._p, description_table._tbl)

    allowed = document.add_paragraph()
    allowed.paragraph_format.space_before = Pt(5)
    allowed.paragraph_format.space_after = Pt(5)
    set_run_font(allowed.add_run("Tipi ammessi: "), size=7.5, bold=True, color=TEAL)
    set_run_font(
        allowed.add_run("COLAZIONE, TRASFERIMENTO, VOLO, TRENO, VISITA, PRANZO, CENA, TEMPO LIBERO, INCONTRO, ALTRO. Una sola attività per riga. Ogni attività inserita è inclusa nel preventivo; ciò che non è presente non è incluso."),
        size=7.5,
        color=MUTED,
    )
    insert_after(description_table._tbl, allowed._p)

    activity_rows = [
        [str(index), "<<TIPO>>", "<<TITOLO / SITO / SERVIZIO>>", "<<LOCALITÀ>>", "<<NOTE / TRATTA / N. VOLO O TRENO>>"]
        for index in range(1, 10)
    ]
    activity_table = make_table(
        document,
        ["Ord.", "Tipo attività", "Titolo / sito / servizio", "Località", "Note"],
        activity_rows,
        [500, 1700, 3400, 2200, 6168],
        header_fill=TEAL,
        font_size=7.2,
    )
    insert_after(allowed._p, activity_table._tbl)

    hotel_label = document.add_paragraph()
    hotel_label.paragraph_format.space_before = Pt(7)
    hotel_label.paragraph_format.space_after = Pt(4)
    set_run_font(hotel_label.add_run("Hotel e pernottamento della giornata"), size=9, bold=True, color=TEAL)
    insert_after(activity_table._tbl, hotel_label._p)

    hotel_table = make_table(
        document,
        ["NOME HOTEL", "CITTÀ", "PAESE", "TRATTAMENTO", "NOTE"],
        [["<<NOME UFFICIALE HOTEL oppure NESSUNO>>", "<<CITTÀ>>", "<<PAESE>>", "<<BB / HB / FB / RO>>", "<<TIPO CAMERA / NOTE>>"]],
        [3200, 1750, 1600, 1800, 5618],
        header_fill=RUST,
        font_size=8,
    )
    insert_after(hotel_label._p, hotel_table._tbl)
    old_activity._element.getparent().remove(old_activity._element)

    # Elimina i paragrafi vuoti ereditati dalla vecchia tabella: se restano
    # prima del cambio sezione Word crea una pagina orizzontale bianca.
    sibling = hotel_table._tbl.getnext()
    while sibling is not None and sibling.tag == qn("w:p"):
        properties = sibling.find(qn("w:pPr"))
        if properties is not None and properties.find(qn("w:sectPr")) is not None:
            break
        text_nodes = sibling.findall(".//" + qn("w:t"))
        if any((node.text or "").strip() for node in text_nodes):
            break
        next_sibling = sibling.getnext()
        sibling.getparent().remove(sibling)
        sibling = next_sibling

    for section in document.sections:
        for paragraph in section.footer.paragraphs:
            for run in paragraph.runs:
                run.text = run.text.replace("v1.0", "v2.0")

    document.core_properties.title = "Modello Preventivo di Viaggio SMF Travel v02"
    document.core_properties.subject = "Template per preventivi importabili in SMF Travel"
    document.core_properties.author = "SMF Travel"
    document.core_properties.last_modified_by = "SMF Travel"
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("reference", type=Path)
    parser.add_argument("output", type=Path)
    arguments = parser.parse_args()
    build(arguments.reference, arguments.output)
