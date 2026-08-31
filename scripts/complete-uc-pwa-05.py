from collections import Counter
from pathlib import Path
from docx import Document

path = Path("docs/testing/SMF_Travel_Rapporto_Completo_Test_v1.1_2026-08-30.docx")
document = Document(path)
cases = document.tables[4]
target = None
for row in cases.rows[1:]:
    if row.cells[0].text.strip() == "UC-PWA-05":
        target = row
        break
if target is None:
    raise RuntimeError("UC-PWA-05 non trovato")

target.cells[2].text = "SUPERATO"
target.cells[3].text = (
    "Verificato su Android/Chrome e PWA installata: consenso notifiche, registrazione e rinnovo "
    "della sottoscrizione FCM, revoca automatica degli endpoint invalidi dopo rotazione VAPID, "
    "consegna con app chiusa/in background, scheduler AWS ogni cinque minuti timezone-aware, "
    "idempotenza e audit. Notifica finale ricevuta con logo Golden Terra Travel corretto."
)

statuses = Counter(row.cells[2].text.strip().upper() for row in cases.rows[1:])
labels = {
    "superati": statuses["SUPERATO"],
    "parziali": statuses["PARZIALE"],
    "bloccati": statuses["BLOCCATO"],
    "non superati": statuses["NON SUPERATO"],
}
for row in document.tables[1].rows:
    label = row.cells[0].text.strip().lower()
    for key, value in labels.items():
        if key in label:
            row.cells[1].text = str(value)
            break

for paragraph in document.paragraphs:
    if paragraph.text.strip().startswith("Sono stati censiti 107 casi d'uso."):
        paragraph.text = (
            f"Sono stati censiti 107 casi d'uso. L'esecuzione ha prodotto {labels['superati']} casi "
            f"superati, {labels['parziali']} parzialmente coperti, {labels['bloccati']} bloccati da "
            f"prerequisiti esterni e {labels['non superati']} casi completamente non superati."
        )
        break

document.save(path)
print(path.resolve())
