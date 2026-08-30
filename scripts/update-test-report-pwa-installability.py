from pathlib import Path
from docx import Document

path = Path("docs/testing/SMF_Travel_Rapporto_Completo_Test_v1.1_2026-08-30.docx")
document = Document(path)
document.paragraphs[4].text = (
    "Sono stati censiti 107 casi d'uso. L'esecuzione ha prodotto 42 casi superati, "
    "57 parzialmente coperti, 8 bloccati da prerequisiti esterni e 0 casi completamente "
    "non superati. Le suite strutturali, la build, il modello dati, le migrazioni fino alla "
    "112, Analytics, chat, gestione spese, governance contenuti, album PDF e PWA risultano conformi."
)
summary = document.tables[1]
summary.rows[2].cells[1].text = "57"
summary.rows[3].cells[1].text = "8"
commands = document.tables[3]
row = commands.add_row().cells
row[0].text = "acceptance:pwa-installability"
row[1].text = "SUPERATO"
row[2].text = "9 immagini, 3 shortcut"
row[3].text = "Manifest, dimensioni asset, prompt Android, guida iOS, registrazione SW e aggiornamenti verificati."
cases = document.tables[4]
for row in cases.rows[1:]:
    if row.cells[0].text.strip() == "UC-PWA-01":
        row.cells[2].text = "PARZIALE"
        row.cells[3].text = (
            "Installabilita verificata automaticamente: manifest standalone portrait, icone maskable, Apple touch icon, "
            "splash iOS, tre shortcut, prompt Android e istruzioni Safari. Resta la prova fisica iOS/Android."
        )
        break
else:
    raise RuntimeError("UC-PWA-01 non trovato")
document.save(path)
print(path.resolve())
