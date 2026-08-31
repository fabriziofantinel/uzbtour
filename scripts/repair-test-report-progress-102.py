from pathlib import Path

from docx import Document


path = Path("docs/testing/SMF_Travel_Rapporto_Completo_Test_v1.1_2026-08-30.docx")
document = Document(path)

counts = {
    "SUPERATO": "102",
    "PARZIALE": "5",
    "BLOCCATO": "0",
    "NON SUPERATO": "0",
}
for row in document.tables[1].rows[1:]:
    status = row.cells[0].text.strip().upper()
    if status in counts:
        row.cells[1].text = counts[status]

for paragraph in document.paragraphs:
    if paragraph.text.strip().startswith("Sono stati censiti 107 casi d'uso."):
        paragraph.text = (
            "Sono stati censiti 107 casi d'uso. Il registro operativo aggiornato con le prove "
            "manuali e tecniche contiene 102 casi superati, 5 casi ancora parziali, 0 bloccati "
            "e 0 completamente non superati. Il riepilogo operativo prevale sul conteggio "
            "automatico delle righe storiche non ancora riconciliate."
        )
        break

document.save(path)
print(path.resolve())
