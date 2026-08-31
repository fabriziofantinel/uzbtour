from collections import Counter
from pathlib import Path

from docx import Document


path = Path("docs/testing/SMF_Travel_Rapporto_Completo_Test_v1.1_2026-08-30.docx")
document = Document(path)
cases = document.tables[4]
target = None
for row in cases.rows[1:]:
    if row.cells[0].text.strip() == "UC-IAM-01":
        target = row
        break
if target is None:
    raise RuntimeError("UC-IAM-01 non trovato")

target.cells[2].text = "SUPERATO"
target.cells[3].text = (
    "Verificato in sessione reale: accesso con credenziali valide, messaggio controllato per "
    "password errata e username inesistente, persistenza dopo refresh/riapertura, logout con "
    "invalidazione della sessione, due username distinti sulla stessa email e assenza di token "
    "nell'URL. Il viaggiatore accede esclusivamente al proprio viaggio e gruppo. Verifica tecnica: "
    "cookie HttpOnly, Secure in produzione e SameSite=Lax; revoca del refresh token e pulizia di "
    "cache/storage al logout; mapping Cognito-Neon solo per utente, membership e agenzia attivi; "
    "protezione Cognito PreventUserExistenceErrors abilitata. Lo smoke DB del ruolo applicativo "
    "non e' applicabile alla connessione locale owner e non viene conteggiato come difetto."
)

statuses = Counter(row.cells[2].text.strip().upper() for row in cases.rows[1:])
labels = {
    "superati": statuses["SUPERATO"],
    "parziali": statuses["PARZIALE"],
    "bloccati": statuses["BLOCCATO"],
    "non superati": statuses["NON SUPERATO"],
}
for row in document.tables[1].rows:
    label = row.cells[0].text.strip().lower()
    for key, value in labels.items():
        if key in label:
            row.cells[1].text = str(value)
            break

for paragraph in document.paragraphs:
    if paragraph.text.strip().startswith("Sono stati censiti 107 casi d'uso."):
        paragraph.text = (
            f"Sono stati censiti 107 casi d'uso. L'esecuzione ha prodotto {labels['superati']} casi "
            f"superati, {labels['parziali']} parzialmente coperti, {labels['bloccati']} bloccati da "
            f"prerequisiti esterni e {labels['non superati']} casi completamente non superati."
        )
        break

document.save(path)
print(path.resolve())
