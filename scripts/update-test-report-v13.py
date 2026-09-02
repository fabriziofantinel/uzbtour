from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


SOURCE = Path("docs/testing/SMF_Travel_Rapporto_Completo_Test_v1.2_2026-09-01.docx")
OUTPUT = Path("docs/testing/SMF_Travel_Rapporto_Completo_Test_v1.3_2026-09-02.docx")
BLUE = "167A84"
LIGHT_BLUE = "E8F3F4"
LIGHT_GREEN = "E7F4EA"
LIGHT_AMBER = "FFF4D6"
INK = "102A35"
MUTED = "53666D"


def set_font(run, size=9, bold=False, color=INK):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    element = properties.find(qn("w:shd"))
    if element is None:
        element = OxmlElement("w:shd")
        properties.append(element)
    element.set(qn("w:fill"), fill)


def set_cell_width(cell, width):
    properties = cell._tc.get_or_add_tcPr()
    element = properties.find(qn("w:tcW"))
    if element is None:
        element = OxmlElement("w:tcW")
        properties.append(element)
    element.set(qn("w:w"), str(width))
    element.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    properties = table._tbl.tblPr
    table_width = properties.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        properties.append(table_width)
    table_width.set(qn("w:w"), str(sum(widths)))
    table_width.set(qn("w:type"), "dxa")
    layout = properties.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        properties.append(layout)
    layout.set(qn("w:type"), "fixed")
    indent = properties.find(qn("w:tblInd"))
    if indent is None:
        indent = OxmlElement("w:tblInd")
        properties.append(indent)
    indent.set(qn("w:w"), "120")
    indent.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_cell_text(cell, text, *, bold=False, color=INK, size=8.4):
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.05
    set_font(paragraph.add_run(text), size=size, bold=bold, color=color)


document = Document(SOURCE)

for section in document.sections:
    footer = section.footer.paragraphs[0]
    footer.clear()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(footer.add_run("Versione 1.3 | 2 settembre 2026"), size=8, color=MUTED)

document.add_page_break()
document.add_heading("Addendum di riesecuzione - 2 settembre 2026", level=1)
lead = document.add_paragraph()
lead.paragraph_format.space_after = Pt(10)
set_font(
    lead.add_run(
        "La presente sezione registra una nuova riesecuzione tecnica del collaudo. "
        "Non sostituisce le evidenze manuali già raccolte: distingue i controlli nuovamente "
        "superati da quelli che richiedono il ripristino dell'ambiente locale."
    ),
    size=10,
)

rows = [
    ("Quality guard e TypeScript", "SUPERATO", "Baseline, 114 migrazioni e 24 confini runtime verificati."),
    ("Build Next.js 16.3.2", "SUPERATO", "Compilazione, TypeScript e generazione di 43 pagine completate."),
    ("Schema Neon V3", "SUPERATO", "78 tabelle, 61 con RLS; nessun vincolo o indice invalido."),
    ("Flussi Neon transazionali", "SUPERATO", "Profili Paese, gruppi, inviti, analytics, engagement, foto e write cutover."),
    ("PWA, offline e sync", "SUPERATO", "Installabilità, cache privata, documenti offline e idempotenza finanziaria."),
    ("Push e pianificazione", "SUPERATO", "Contratti client/service worker e invio pianificato verificati."),
    ("Qualità deterministica import", "SUPERATO", "Rilevate tutte le otto classi di anomalia attese."),
    ("Invocazioni Bedrock reali", "DA RIESEGUIRE", "Sessione AWS locale scaduta; nessuna chiamata AI completata in questa sessione."),
    ("Autorizzazione ruolo DB runtime", "DA RIESEGUIRE", "Manca una DATABASE_URL locale valida per il ruolo applicativo ridotto."),
    ("Browser autenticato", "DA RIESEGUIRE", "Le sessioni disponibili sono ferme al login; richiesta una sessione autenticata per i ruoli principali."),
]

table = document.add_table(rows=1, cols=3)
table.style = "Table Grid"
table.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_geometry(table, [2700, 1500, 5160])
for cell, label in zip(table.rows[0].cells, ("Area", "Esito", "Evidenza")):
    shade(cell, BLUE)
    add_cell_text(cell, label, bold=True, color="FFFFFF", size=8.5)

for area, status, evidence in rows:
    row = table.add_row()
    add_cell_text(row.cells[0], area, bold=True, size=8.2)
    if status == "SUPERATO":
        shade(row.cells[1], LIGHT_GREEN)
        add_cell_text(row.cells[1], status, bold=True, color="23623A", size=8)
    else:
        shade(row.cells[1], LIGHT_AMBER)
        add_cell_text(row.cells[1], status, bold=True, color="8A5A00", size=8)
    add_cell_text(row.cells[2], evidence, size=8)

document.add_heading("Anomalie corrette durante la riesecuzione", level=2)
for text in (
    "Il validatore dello schema ammette ora il totale corrente di 78 tabelle, mantenendo invariati gli altri gate.",
    "La fixture del quiz usa date relative al giorno di esecuzione e non produce più falsi errori dopo il 1 settembre 2026.",
):
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.add_run(text)

document.add_heading("Valutazione corrente", level=2)
document.add_paragraph(
    "La build e i controlli automatizzati disponibili sono verdi. Il punto 1 del piano resta aperto "
    "finché non vengono rieseguiti Bedrock, il ruolo database runtime e i percorsi browser autenticati. "
    "Questi elementi sono classificati come verifiche pendenti di ambiente, non come difetti applicativi confermati."
)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
document.save(OUTPUT)
print(OUTPUT.resolve())
