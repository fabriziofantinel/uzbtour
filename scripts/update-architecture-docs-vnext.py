from copy import deepcopy
from pathlib import Path
from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]


def add_bullets(doc, entries):
    for entry in entries:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(entry)


def add_release_page(doc, title, subtitle, sections):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    heading = doc.add_heading(title, level=1)
    heading.paragraph_format.keep_with_next = True
    p = doc.add_paragraph(subtitle)
    p.runs[0].italic = True
    p.runs[0].font.size = Pt(9)
    for name, intro, bullets in sections:
        doc.add_heading(name, level=2)
        if intro:
            doc.add_paragraph(intro)
        add_bullets(doc, bullets)


def clone_and_update(source, target, title, subtitle, sections):
    doc = Document(source)
    add_release_page(doc, title, subtitle, sections)
    props = doc.core_properties
    props.title = title
    props.subject = "Aggiornamento architetturale SMF Travel del 3 settembre 2026"
    props.comments = "Aggiornato con Tour Leader temporaneo, governance AI e osservabilità."
    doc.save(target)


clone_and_update(
    ROOT / "docs/architecture/SMF_Travel_Architettura_Soluzione_v1.2.docx",
    ROOT / "docs/architecture/SMF_Travel_Architettura_Soluzione_v1.3.docx",
    "Addendum architetturale - versione 1.3",
    "Stato consolidato al 3 settembre 2026. Le sezioni precedenti restano valide salvo quanto integrato qui.",
    [
        ("Tour Leader limitato alla partenza", "Il ruolo operativo non coincide con un membro permanente dell'agenzia.", [
            "Il responsabile dell'agenzia può invitare un professionista esterno con username univoco, e-mail di attivazione e intervallo valid_from/valid_until.",
            "L'autorizzazione è verificata sulla singola partenza a ogni accesso; la revoca è immediata, motivata e registrata nell'audit.",
            "Dopo il login il Tour Leader accede a un'area dedicata con sole partenze attive e autorizzate, presenze, chat e segnalazioni operative essenziali.",
            "Il flusso usa identificativi IAM UUID nativi; non crea iam.agency_memberships permanenti.",
        ]),
        ("Governance Bedrock e costi", "La generazione resta asincrona e osservabile.", [
            "Le citazioni di contenuti sensibili sono filtrate tramite allowlist di fonti istituzionali; una risposta priva di almeno una fonte attendibile viene rifiutata.",
            "Il test di accettazione Bedrock verifica dal vivo il grounding; il replay deterministico resta nella CI e non consuma token.",
            "ops.ai_model_prices conserva prezzi Nova 2 Lite versionati per modello, regione, validità e fonte ufficiale AWS.",
            "La telemetria AI è associata al job originario e consente riconciliazione di token, latenza, modello e costo stimato.",
        ]),
        ("Osservabilità e gestione errori", "L'impianto è provider-neutral e non richiede un nuovo servizio a pagamento.", [
            "instrumentation.ts intercetta gli errori server catturati da Next.js e produce log JSON strutturati nei Runtime Logs Vercel.",
            "Ogni errore API inatteso riceve un errorId UUID, restituito nel payload e nell'header x-smf-error-id per la correlazione con assistenza e log.",
            "I record includono ambiente, deployment, route e tipo di esecuzione; i valori di contesto sono limitati e gli oggetti non ammessi vengono oscurati.",
            "I worker AWS continuano a esporre log e allarmi CloudWatch; il deployment SAM è verificato con stack UPDATE_COMPLETE.",
        ]),
    ],
)

clone_and_update(
    ROOT / "docs/data-model/SMF_Travel_Modello_Logico_Dati_v1.4.docx",
    ROOT / "docs/data-model/SMF_Travel_Modello_Logico_Dati_v1.5.docx",
    "Estensione del modello logico - versione 1.5",
    "Entità e vincoli logici introdotti o consolidati al 3 settembre 2026.",
    [
        ("Assegnazione temporanea del Tour Leader", "Relazione tra utente, agenzia e partenza priva di membership permanente.", [
            "DepartureStaffAssignment identifica utente, partenza, ruolo tour_leader, stato, autore e intervallo temporale di efficacia.",
            "Una persona può essere invitata e attivata come identità IAM senza appartenere stabilmente all'organico dell'agenzia.",
            "La validità richiede valid_until maggiore di valid_from; la revoca conserva autore, istante e motivazione.",
            "L'accesso operativo è valido soltanto quando lo stato è active e l'istante corrente ricade nell'intervallo autorizzato.",
        ]),
        ("Osservabilità delle elaborazioni AI", "Il costo è derivato da utilizzo misurato e tariffa temporalmente valida.", [
            "AiModelPrice identifica modello, regione, unità di prezzo per milione di token, fonte e periodo di efficacia.",
            "AiInvocationTelemetry collega job, operazione, modello, regione, token input/output, latenza, risultato e costo stimato.",
            "Le citazioni ufficiali appartengono al dossier generato e sono sottoposte a policy prima della persistenza applicativa.",
        ]),
        ("Correlazione degli errori", "L'errorId non è un'entità di dominio persistente obbligatoria.", [
            "L'identificatore di errore viene prodotto dal runtime e correlato ai log di deployment, evitando la duplicazione di stack trace nel database transazionale.",
            "I dati di contesto sono minimizzati: nessuna credenziale, payload documentale o dato sanitario viene incluso automaticamente.",
        ]),
    ],
)

clone_and_update(
    ROOT / "docs/data-model/SMF_Travel_Modello_Fisico_Dati_v1.4.docx",
    ROOT / "docs/data-model/SMF_Travel_Modello_Fisico_Dati_v1.5.docx",
    "Estensione del modello fisico - versione 1.5",
    "Oggetti PostgreSQL e contratti runtime aggiunti fino alla migrazione 143.",
    [
        ("travel.departure_staff_assignments", "Tabella esistente estesa dalle migrazioni 140, 142 e 143.", [
            "Colonne temporali: valid_from TIMESTAMPTZ NOT NULL e valid_until TIMESTAMPTZ NOT NULL con CHECK valid_until > valid_from.",
            "Revoca: revoked_at, revoked_by UUID e revocation_reason; il vincolo mantiene coerenza tra status active/revoked e metadati.",
            "Indice parziale departure_staff_assignments_active_period_idx su departure_id, user_id, valid_from e valid_until per assegnazioni attive.",
        ]),
        ("Funzioni app native", "Tutti i nuovi attori sono UUID IAM, senza aumento del debito di firme legacy.", [
            "app.provision_departure_tour_leader_v3(UUID, ...): crea identità invitata, mapping compatibile, invito e assegnazione temporanea, senza agency_membership.",
            "app.assign_tour_leader_period_v3(UUID, UUID, UUID, TIMESTAMPTZ, TIMESTAMPTZ): assegna o riattiva un incarico con periodo esplicito.",
            "app.list_departure_tour_leaders_v3(UUID, UUID) e app.list_my_tour_leader_departures_v3(UUID): letture autorizzate per gestione e dashboard.",
            "app.revoke_tour_leader_v3(UUID, UUID, UUID, TEXT): revoca atomica e auditata.",
            "app.resolve_cognito_authenticated_user(TEXT) espone anche native_user_id per i contratti runtime UUID-native.",
        ]),
        ("Tariffe e telemetria AI", "Migrazione 141 e strutture operative precedenti.", [
            "ops.ai_model_prices contiene le tariffe ufficiali Nova 2 Lite per eu-central-1 e us-east-1, espresse in microUSD per milione di token e versionate per validità.",
            "Le invocazioni Bedrock persistono metrica di utilizzo e riferimento al job; la stima del costo seleziona la tariffa coerente con modello, regione e istante.",
        ]),
        ("Contratti di esercizio", "Controlli applicativi aggiunti senza nuove tabelle per gli errori.", [
            "Le risposte HTTP 500 includono errorId e header x-smf-error-id; instrumentation.ts emette application_error in JSON.",
            "Migrazioni 142 e 143 sono riproducibili dalla catena di bootstrap; la validazione target rileva 88 tabelle, 68 con RLS, zero vincoli o indici invalidi.",
        ]),
    ],
)
