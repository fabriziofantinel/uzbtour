from pathlib import Path
from docx import Document

path = Path("docs/testing/SMF_Travel_Rapporto_Completo_Test_v1.1_2026-08-30.docx")
document = Document(path)
document.paragraphs[4].text = (
    "Sono stati censiti 107 casi d'uso. L'esecuzione ha prodotto 42 casi superati, "
    "56 parzialmente coperti, 9 bloccati da prerequisiti esterni e 0 casi completamente "
    "non superati. Le suite strutturali, la build, il modello dati, le migrazioni fino alla "
    "112, Analytics, chat, gestione spese, governance contenuti, album finale PDF e pacchetto offline risultano conformi."
)
summary = document.tables[1]
summary.rows[2].cells[1].text = "56"
summary.rows[3].cells[1].text = "9"

findings = document.tables[2]
row = findings.add_row().cells
row[0].text = "DEF-002"
row[1].text = "Chiusa"
row[2].text = "PWA offline"
row[3].text = "URL relativi non normalizzati e documenti salvati in una cache diversa da quella consultata offline."
row[4].text = "Normalizzazione rispetto all'origine, cache media coerente e acceptance offline superata."

commands = document.tables[3]
row = commands.add_row().cells
row[0].text = "acceptance:pwa-offline"
row[1].text = "SUPERATO"
row[2].text = "3 risorse + logout"
row[3].text = "Programma, JSON viaggio e documento consultati senza rete; cache private cancellate al logout."

cases = document.tables[4]
for row in cases.rows[1:]:
    if row.cells[0].text.strip() == "UC-PWA-02":
        row.cells[2].text = "PARZIALE"
        row.cells[3].text = (
            "Service worker verificato con acceptance ripetibile: download esplicito, programma, dati e documenti "
            "disponibili offline, cache private eliminate al logout. Resta il collaudo fisico su iOS e Android."
        )
        break
else:
    raise RuntimeError("UC-PWA-02 non trovato")

document.paragraphs[25].text = (
    "La baseline tecnica e stabile e le nuove funzionalita risultano integrate senza regressioni di build o schema. "
    "Sono dimostrati album PDF e pacchetto offline, incluso l'isolamento e la cancellazione delle cache private. "
    "Restano da completare i flussi infrastrutturali end-to-end e il collaudo PWA su dispositivi mobili reali."
)
document.save(path)
print(path.resolve())
