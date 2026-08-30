from pathlib import Path
from docx import Document
path=Path("docs/testing/SMF_Travel_Rapporto_Completo_Test_v1.1_2026-08-30.docx");document=Document(path)
document.paragraphs[4].text=("Sono stati censiti 107 casi d'uso. L'esecuzione ha prodotto 42 casi superati, 59 parzialmente coperti, 6 bloccati da prerequisiti esterni e 0 casi completamente non superati. Le suite strutturali, la build, il modello dati, le migrazioni fino alla 115, Analytics, chat, gestione spese, governance contenuti, album PDF e PWA risultano conformi.")
document.paragraphs[11].text=("Migrazioni fino alla 115 verificate: KPI, governance fonti, contratti runtime, pgcrypto, contest a due foto, bingo globale, Web Push multi-agenzia e pianificazione timezone-aware.")
summary=document.tables[1];summary.rows[2].cells[1].text="59";summary.rows[3].cells[1].text="6"
findings=document.tables[2]
for values in [
 ("DEF-004","Chiusa","Push schedulato","Il claim falliva per ambiguita PL/pgSQL sulla colonna departure_id.","Migrazione 114 con conflict constraint esplicito; acceptance superata."),
 ("DEF-005","Chiusa","Timezone","Reminder e quiz usavano date calcolate con timezone differenti vicino alla mezzanotte.","Migrazione 115: entrambi calcolati nel timezone della partenza."),
]:
 row=findings.add_row().cells
 for cell,value in zip(row,values):cell.text=value
commands=document.tables[3]
for values in [
 ("db:migrate:v3:scheduled-push-claim","SUPERATO","Migrazione 114","Claim PL/pgSQL corretto e applicato."),
 ("db:migrate:v3:scheduled-push-local-date","SUPERATO","Migrazione 115","Data locale della partenza applicata a reminder e quiz."),
 ("acceptance:v3:scheduled-push","SUPERATO","Rollback","Reminder, quiz, idempotenza, retry e audit consegna verificati."),
]:
 row=commands.add_row().cells
 for cell,value in zip(row,values):cell.text=value
for row in document.tables[4].rows[1:]:
 if row.cells[0].text.strip()=="UC-PWA-05":
  row.cells[2].text="PARZIALE";row.cells[3].text=("Pianificazione verificata su Neon: reminder partenza, quiz, data locale, idempotenza, retry e audit. Il piano Vercel Hobby consente un solo cron giornaliero con precisione oraria, quindi non garantisce le 20:00 locali per tutti i fusi; resta inoltre il test su dispositivo reale.");break
else:raise RuntimeError("UC-PWA-05 non trovato")
document.save(path);print(path.resolve())
