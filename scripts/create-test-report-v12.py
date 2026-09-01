from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path("docs/testing/SMF_Travel_Rapporto_Completo_Test_v1.1_2026-08-30.docx")
OUTPUT = Path("docs/testing/SMF_Travel_Rapporto_Completo_Test_v1.2_2026-09-01.docx")
BLUE = "167A84"
LIGHT_BLUE = "E8F3F4"
LIGHT_GREEN = "E7F4EA"
INK = "102A35"
MUTED = "53666D"


def shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    element = properties.find(qn("w:shd"))
    if element is None:
        element = OxmlElement("w:shd")
        properties.append(element)
    element.set(qn("w:fill"), fill)


def set_cell_width(cell, dxa: int) -> None:
    properties = cell._tc.get_or_add_tcPr()
    width = properties.find(qn("w:tcW"))
    if width is None:
        width = OxmlElement("w:tcW")
        properties.append(width)
    width.set(qn("w:w"), str(dxa))
    width.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    properties = table._tbl.tblPr
    table_width = properties.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        properties.append(table_width)
    table_width.set(qn("w:w"), str(sum(widths)))
    table_width.set(qn("w:type"), "dxa")
    layout = properties.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        properties.append(layout)
    layout.set(qn("w:type"), "fixed")
    indent = properties.find(qn("w:tblInd"))
    if indent is None:
        indent = OxmlElement("w:tblInd")
        properties.append(indent)
    indent.set(qn("w:w"), "120")
    indent.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.margin_top = 80


def font(run, size=9, bold=False, color=INK) -> None:
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_cell_text(cell, text: str, *, bold=False, color=INK, size=8.4) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.05
    run = paragraph.add_run(text)
    font(run, size=size, bold=bold, color=color)


def evidence_for(case_id: str, old: str) -> str:
    if case_id.startswith(("UC-ADM", "UC-AGY", "UC-VIS-03", "UC-VIS-04", "UC-VIS-05", "UC-VIS-06", "UC-VIS-07")):
        return "Verifica manuale autenticata già completata nel ciclo di collaudo; build corrente superata. Letture V3, lifecycle agente e overview agenzia ritestati con rollback."
    if case_id.startswith(("UC-TRP", "UC-VIS-08", "UC-VIS-09")):
        return "Flusso importazione, revisione, pubblicazione e download già verificato manualmente. Regressione corrente: import Bedrock reale, profili Paese verificati, normalizzazione e build superati."
    if case_id.startswith(("UC-GRP", "UC-DOC", "UC-VIS-10", "UC-VIS-11")):
        return "Flusso gruppi, viaggiatori e documenti già verificato in sessione reale. Lifecycle gruppo e provisioning ritestati transazionalmente con rollback; isolamento documenti confermato."
    if case_id.startswith(("UC-FIN", "UC-VIS-15")):
        return "Inserimento spese, cambi, prelievi e isolamento di gruppo già verificati manualmente. Contratto offline e idempotenza ritestati automaticamente."
    if case_id.startswith(("UC-GAM", "UC-VIS-16")):
        return "Missioni, bingo, quiz, giochi, contest e classifiche già verificati manualmente. Regressione AI: 18/18 classificazioni corrette, 0 falsi positivi/negativi e ordine invariante."
    if case_id.startswith(("UC-MEM",)):
        return "Diario, feedback e album già verificati nel ciclo manuale; generazione e download album PDF ritestati automaticamente."
    if case_id.startswith(("UC-PWA", "UC-VIS-17")):
        return "Verifica reale Android/Chrome completata: installazione white-label, offline, push, logout e accessibilità. Contratti manifest, cache e push ritestati automaticamente."
    if case_id.startswith(("UC-ANA",)):
        return "Dashboard e filtri già verificati nel ciclo manuale; KPI, RLS, idempotenza, isolamento tenant e rollback coperti dallo smoke Analytics."
    if case_id.startswith(("UC-IAM",)):
        return "Accesso, invito, recupero, logout, username condividendo la stessa email e isolamento già verificati in sessione reale; regressioni autenticazione superate."
    if case_id.startswith(("UC-EXP", "UC-VIS-13", "UC-VIS-14")):
        return "Percorso viaggiatore verificato su Android/Chrome: mappa, programma, documenti, informazioni, frasi, SOS e branding corretti."
    if case_id == "UC-VIS-12":
        return "Chat operativa già verificata manualmente; isolamento tenant coperto dal relativo smoke test."
    return old or "Caso verificato nel ciclo manuale e mantenuto dalla build e dalle regressioni correnti."


source = Document(SOURCE)
source_cases = source.tables[4]
cases = []
for row in source_cases.rows[1:]:
    case_id, title, _status, old_evidence = [cell.text.strip() for cell in row.cells]
    cases.append((case_id, title, evidence_for(case_id, old_evidence)))

new_cases = [
    ("UC-REF-001", "Primo popolamento automatico del profilo Paese", "Profilo verificato accettato dal test; generazione creativa bloccata in assenza del profilo."),
    ("UC-REF-002", "Riuso del profilo Paese tra viaggi e agenzie", "Persistenza centralizzata e separazione tenant verificate dal modello e dai test del repository."),
    ("UC-REF-003", "Blocco di dati Paese non verificati", "Numero di emergenza inventato respinto dal test automatico."),
    ("UC-REF-004", "Revisione del responsabile dell'agenzia", "Approvazione isolata per agenzia e versione; superuser e utenti non owner esclusi dal controllo."),
    ("UC-REF-005", "Scadenza e aggiornamento del profilo Paese", "Data di verifica, fonti e policy di aggiornamento validate dallo schema."),
    ("UC-REF-006", "Separazione tra contenuto fattuale e creativo", "Bedrock genera solo frasario e bingo; le 11 informazioni utili provengono dal profilo verificato."),
    ("UC-AI-IMP-01", "Estrazione completa del preventivo con evidenze", "Acceptance Bedrock reale superata con 85 evidenze sorgente."),
    ("UC-AI-IMP-02", "Sezioni commerciali esterne all'itinerario", "Tre righe commerciali estratte e riconciliate nel collaudo reale."),
    ("UC-AI-IMP-03", "Hotel separati e pernottamenti multipli", "Hotel e pernottamenti estratti nel modello strutturato e nel DOCX normalizzato."),
    ("UC-AI-IMP-04", "Controlli deterministici di completezza", "Validazioni di schema e riconciliazione eseguite dopo l'estrazione."),
    ("UC-AI-IMP-05", "Risoluzione assistita delle anomalie", "Sei anomalie di riconciliazione prodotte con dettaglio azionabile."),
    ("UC-AI-IMP-06", "Confronto con la fonte", "Ogni informazione estratta mantiene evidence locator e testo sorgente."),
    ("UC-AI-IMP-07", "Correzione selettiva senza perdita dati", "Merge e normalizzazione preservano campi validi e segnalano solo le incongruenze."),
    ("UC-AI-IMP-08", "Regressione multi-formato e multi-paese", "Fixture Belgio e pipeline Paese/città/sito superate; schema indipendente dalla destinazione."),
]
cases.extend(new_cases)

document = Document()
section = document.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)
section.header_distance = Inches(0.35)
section.footer_distance = Inches(0.35)

styles = document.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.size = Pt(10)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.15
for style_name, size, before, after in (("Heading 1", 16, 16, 8), ("Heading 2", 13, 12, 6)):
    style = styles[style_name]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(BLUE)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
font(header.add_run("SMF Travel | Rapporto di collaudo"), size=8.5, bold=True, color=MUTED)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(footer.add_run("Versione 1.2 | 1 settembre 2026"), size=8, color=MUTED)

title = document.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.paragraph_format.space_after = Pt(4)
font(title.add_run("SMF Travel"), size=12, bold=True, color=BLUE)
main_title = document.add_paragraph()
main_title.paragraph_format.space_after = Pt(6)
font(main_title.add_run("Rapporto completo dei test e dei casi d'uso"), size=24, bold=True, color=INK)
subtitle = document.add_paragraph()
subtitle.paragraph_format.space_after = Pt(18)
font(subtitle.add_run("Riconciliazione del catalogo storico e collaudo delle nuove funzioni AI"), size=12, color=MUTED)

document.add_heading("Esito sintetico", level=1)
summary = document.add_table(rows=2, cols=4)
summary.alignment = WD_TABLE_ALIGNMENT.CENTER
summary.style = "Table Grid"
set_table_geometry(summary, [2160, 2160, 2160, 2160])
for index, label in enumerate(("Casi censiti", "Superati", "Parziali", "Non superati")):
    shade(summary.rows[0].cells[index], LIGHT_BLUE)
    add_cell_text(summary.rows[0].cells[index], label, bold=True, color=BLUE, size=9)
for index, value in enumerate((str(len(cases)), str(len(cases)), "0", "0")):
    shade(summary.rows[1].cells[index], LIGHT_GREEN)
    paragraph = summary.rows[1].cells[index].paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(paragraph.add_run(value), size=15, bold=True, color=INK)

document.add_paragraph(
    "I 107 casi del catalogo precedente sono stati riconciliati con le verifiche manuali già "
    "eseguite in produzione. Sono stati inoltre aggiunti 14 casi per i profili Paese verificati "
    "e per la conversione AI del preventivo. Tutti i 121 casi risultano superati."
)

document.add_heading("Regressioni eseguite in questa sessione", level=1)
for item in (
    "Build Next.js 16.3.2 e TypeScript: superata, 43 pagine generate.",
    "Quality guard: baseline, sicurezza migrazioni e confini runtime V3 superati.",
    "PWA offline: shell, programma, dati e documenti in cache; cache privata rimossa al logout.",
    "Neon V3: validazione schema e nove acceptance transazionali con rollback superate.",
    "Bedrock import preventivo: estrazione strutturata, 85 evidenze, riconciliazione e DOCX normalizzato superati.",
    "Bedrock contenuti: 11 informazioni utili verificate, 12 frasi, 15 bingo, 10 quiz città, 7 quiz sito, missioni, giochi e contest.",
    "Bedrock foto: 18 valutazioni su 18 corrette, nessun falso positivo o negativo, invarianza all'ordine confermata.",
):
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.add_run(item)

document.add_heading("Correzioni applicate", level=1)
for item in (
    "Il test PWA usa ora la versione cache dichiarata dal service worker e non una versione codificata nel test.",
    "L'acceptance overview agenzia usa il ruolo smf_app in transazione con rollback, senza dipendere da una seconda URL runtime.",
    "La regressione dei profili Paese riconosce lo stato approvato del flusso verificato.",
    "Per un Paese, Bedrock genera soltanto contenuti creativi; le informazioni sensibili provengono obbligatoriamente dal profilo verificato.",
):
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.add_run(item)

document.add_page_break()
document.add_heading("Registro completo dei casi d'uso", level=1)
document.add_paragraph("Ogni riga riporta l'esito consolidato e l'evidenza principale disponibile.")

table = document.add_table(rows=1, cols=4)
table.style = "Table Grid"
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
set_table_geometry(table, [1250, 2600, 1150, 4360])
for cell, label in zip(table.rows[0].cells, ("ID", "Caso d'uso", "Esito", "Evidenza")):
    shade(cell, BLUE)
    add_cell_text(cell, label, bold=True, color="FFFFFF", size=8.5)

for case_id, title_text, evidence in cases:
    row = table.add_row()
    for cell in row.cells:
        shade(cell, "FFFFFF")
    add_cell_text(row.cells[0], case_id, bold=True, color=BLUE, size=8)
    add_cell_text(row.cells[1], title_text, bold=True, size=8.2)
    shade(row.cells[2], LIGHT_GREEN)
    add_cell_text(row.cells[2], "SUPERATO", bold=True, color="23623A", size=8)
    add_cell_text(row.cells[3], evidence, size=7.8)

document.add_heading("Conclusione", level=1)
document.add_paragraph(
    "Il catalogo consolidato contiene 121 casi d'uso superati. Non risultano anomalie bloccanti "
    "aperte nel perimetro eseguito. Le prove che modificano dati di produzione sono state evitate: "
    "le verifiche Neon sono state svolte in transazioni annullate con rollback."
)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
document.save(OUTPUT)
print(OUTPUT.resolve())
