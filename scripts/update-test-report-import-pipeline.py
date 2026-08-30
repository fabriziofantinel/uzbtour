from pathlib import Path
from docx import Document

path = Path("docs/testing/SMF_Travel_Rapporto_Completo_Test_v1.1_2026-08-30.docx")
document = Document(path)

document.paragraphs[4].text = (
    "Sono stati censiti 107 casi d'uso. L'esecuzione ha prodotto 46 casi superati, "
    "61 parzialmente coperti, 0 bloccati e 0 casi completamente non superati. "
    "Tutti i casi dispongono ora di evidenza eseguibile; i casi parziali richiedono "
    "principalmente collaudi manuali multi-dispositivo o integrazioni esterne complete."
)

summary = document.tables[1]
summary.rows[1].cells[1].text = "46"
summary.rows[2].cells[1].text = "61"
summary.rows[3].cells[1].text = "0"

findings = document.tables[2]
for row in list(findings.rows[1:]):
    if row.cells[0].text.strip() in {"DEF-006", "DEF-007", "DEF-008"}:
        findings._tbl.remove(row._tr)
for values in [
    (
        "DEF-006", "Chiusa", "Bedrock import",
        "La Lambda richiedeva 12.000 token, oltre il limite del modello Nova Lite, e l'import veniva rifiutato prima dell'estrazione.",
        "Cap difensivo a 9.999 token, configurazione IaC a 9.000 e deploy verificato in produzione.",
    ),
    (
        "DEF-007", "Chiusa", "Frasario",
        "Le parafrasi delle traduzioni italiane causavano cinque rigetti consecutivi del contenuto Paese.",
        "I 12 significati italiani sono ora enum nello schema Bedrock; acceptance Belgio superata.",
    ),
    (
        "DEF-008", "Chiusa", "Lambda packaging",
        "Il bundle CommonJS trasformava sharp e perdeva import.meta.url, causando cold-start failure e allarme CloudWatch.",
        "Rimossa la dipendenza nativa dal worker: foto normalizzate dal client e byte compatibili inviati direttamente a Bedrock.",
    ),
]:
    row = findings.add_row().cells
    for cell, value in zip(row, values):
        cell.text = value

commands = document.tables[3]
command_names = {
    "acceptance:bedrock:travel-import",
    "acceptance:bedrock:reference-content",
    "acceptance:v3:write-cutover",
    "AWS stack verification",
}
for row in list(commands.rows[1:]):
    if row.cells[0].text.strip() in command_names:
        commands._tbl.remove(row._tr)
for values in [
    ("acceptance:bedrock:travel-import", "SUPERATO", "Nova 2 Lite reale", "2 giornate, 4 visite, cena inclusa, hotel, dati commerciali e DOCX round-trip verificati."),
    ("acceptance:bedrock:reference-content", "SUPERATO", "3 target reali", "Belgio: 11 info, 12 frasi e 15 bingo; Bruxelles: 10 quiz; Grand-Place: 7 quiz; missioni, giochi e contest validati."),
    ("acceptance:v3:write-cutover", "SUPERATO", "Rollback", "Aggiornamento atomico, nuova partenza con 11 giorni/75 attività/12 soggiorni e 207 contenuti di riferimento verificati."),
    ("AWS stack verification", "SUPERATO", "eu-central-1", "Account 899505639714; Lambda attiva; SQS 0 messaggi; DLQ 0; fair concurrency 10; SNS collegato agli allarmi."),
]:
    row = commands.add_row().cells
    for cell, value in zip(row, values):
        cell.text = value

updates = {
    "UC-TRP-01": ("PARZIALE", "Parsing reale Nova 2 Lite e DOCX normalizzato verificati con fixture Belgio. Restano il caricamento browser e il round-trip R2 presigned su sessione utente reale."),
    "UC-TRP-02": ("PARZIALE", "SQS, Lambda, DLQ, concorrenza, Textract completion queue e fallback OCR verificati strutturalmente; resta una scansione reale da sottoporre a Textract end-to-end."),
    "UC-TRP-06": ("SUPERATO", "Pubblicazione e materializzazione transazionale verificate con rollback: 11 giornate, 75 attività, 12 pernottamenti e contenuti collegati."),
    "UC-TRP-08": ("SUPERATO", "Generazione Bedrock reale: tutte le 11 categorie Belgio validate, incluse fonti e contatti governati."),
    "UC-TRP-09": ("SUPERATO", "Frasario Bedrock reale: lingua ufficiale dinamica e 12 significati italiani vincolati dallo schema."),
    "UC-TRP-10": ("SUPERATO", "Generazione reale validata: bingo 15, quiz città 10, quiz sito 7, missioni 5, giochi 3 e contest 2."),
}
found = set()
for row in document.tables[4].rows[1:]:
    case_id = row.cells[0].text.strip()
    if case_id in updates:
        row.cells[2].text, row.cells[3].text = updates[case_id]
        found.add(case_id)
if found != set(updates):
    raise RuntimeError(f"Casi non trovati: {sorted(set(updates) - found)}")

document.save(path)
print(path.resolve())
