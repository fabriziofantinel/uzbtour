from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
BLACK = RGBColor(0, 0, 0)


def heading(doc, text, level):
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True
    for run in paragraph.runs:
        run.font.color.rgb = BLACK
    return paragraph


def bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.add_run(text)
    return paragraph


def update_running_version(doc, version):
    pattern = re.compile(r"\b(?:v|Versione\s+)\d+(?:[.,]\d+)+", re.IGNORECASE)

    def paragraphs(container):
        yield from container.paragraphs
        for table in container.tables:
            for row in table.rows:
                for cell in row.cells:
                    yield from cell.paragraphs

    for section in doc.sections:
        for container in (section.header, section.footer):
            for paragraph in paragraphs(container):
                for run in paragraph.runs:
                    if pattern.search(run.text):
                        run.text = pattern.sub(
                            lambda match: (
                                f"v{version}"
                                if match.group(0).lower().startswith("v")
                                and not match.group(0).lower().startswith("versione")
                                else f"Versione {version}"
                            ),
                            run.text,
                        )


def add_addendum(source, target, title, subtitle, sections, subject, version):
    doc = Document(source)
    update_running_version(doc, version)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    heading(doc, title, 1)
    intro = doc.add_paragraph(subtitle)
    intro.runs[0].italic = True
    intro.runs[0].font.size = Pt(9)
    for section_title, lead, entries in sections:
        heading(doc, section_title, 2)
        if lead:
            doc.add_paragraph(lead)
        for entry in entries:
            bullet(doc, entry)
    props = doc.core_properties
    props.title = title
    props.subject = subject
    props.comments = "Allineato alla migrazione 200, alla rooming list e al ciclo post viaggio."
    doc.save(target)


add_addendum(
    ROOT / "docs/architecture/SMF_Travel_Architettura_Soluzione_v1.5.docx",
    ROOT / "docs/architecture/SMF_Travel_Architettura_Soluzione_v1.6.docx",
    "Addendum architetturale versione 1.6",
    "Stato as built allineato alla migrazione 200 e verificato l'8 settembre 2026.",
    [
        (
            "Rooming list",
            "La rooming list è un sottodominio operativo dei pernottamenti già presenti nel programma.",
            [
                "Responsabile, agente e accompagnatore assegnato operano tramite route Next.js autorizzate e stored API PostgreSQL; la guida è esclusa.",
                "Le camere e gli occupanti sono isolati per agenzia, partenza, pernottamento e gruppo mediante chiavi composte e RLS forzata.",
                "La validazione server segnala viaggiatori non assegnati, capienza ecceduta e minori privi di un adulto dello stesso gruppo.",
                "L'esportazione DOCX è prodotta lato server e contiene nomi, composizione camere ed esigenze operative, senza numeri di documento o dati sanitari.",
            ],
        ),
        (
            "Valutazione e passaparola post viaggio",
            "Il flusso è deterministico e non invoca servizi AI.",
            [
                "Due giorni dopo la data di rientro il sistema rende disponibile una valutazione da 0 a 10 con commento facoltativo e pianifica una notifica push idempotente.",
                "I promotori con voto 9 o 10 ricevono le azioni recensione pubblica e segnalazione amico; i detrattori con voto da 0 a 6 vengono indirizzati alla chat privata; i voti 7 e 8 seguono un percorso neutro.",
                "Il codice passaparola è deterministico e riferito all'agenzia, alla partenza e al viaggiatore senza esporre identificativi sensibili.",
                "Analytics aggrega media e commenti per partenza, gruppo e viaggiatore e li correla ai feedback di tappa nel medesimo tenant.",
            ],
        ),
        (
            "Inventario e ricostruibilità",
            "Il catalogo non è più validato soltanto tramite un conteggio complessivo.",
            [
                "Il manifesto database/v3-table-inventory.json elenca nominalmente 99 tabelle applicative attese.",
                "La ricostruzione da vuoto aggiunge esclusivamente ops.repository_migrations come registro tecnico dei checksum, portando a 100 il conteggio del branch temporaneo.",
                "Il gate confronta nomi mancanti e inattesi, 73 tabelle RLS, vincoli, indici, indici tenant-leading e riconciliazioni shadow.",
                "La CI ricostruisce tutte le migrazioni 001-200 su un branch Neon effimero e lo elimina al termine.",
            ],
        ),
    ],
    "Architettura SMF Travel allineata allo schema 200",
    "1.6",
)

add_addendum(
    ROOT / "docs/data-model/SMF_Travel_Modello_Logico_Dati_v1.7.docx",
    ROOT / "docs/data-model/SMF_Travel_Modello_Logico_Dati_v1.8.docx",
    "Estensione del modello logico versione 1.8",
    "Entità, relazioni e invarianti introdotte dalla migrazione 200.",
    [
        (
            "Assegnazione camere",
            "RoomingRoom rappresenta una camera relativa a un pernottamento e a un gruppo.",
            [
                "Un pernottamento può avere molte camere e ogni camera appartiene a una sola partenza, un solo soggiorno e un solo gruppo.",
                "RoomingRoomOccupant associa una camera a un viaggiatore già membro attivo del gruppo della stessa partenza.",
                "Ogni viaggiatore può comparire una sola volta nella rooming list dello stesso pernottamento.",
                "La tipologia determina la capienza: singola 1, doppia o matrimoniale 2, tripla 3.",
                "Un minore può essere assegnato soltanto se nella stessa camera è presente almeno un adulto del medesimo gruppo.",
            ],
        ),
        (
            "Valutazione post viaggio",
            "PostTripReview rappresenta una valutazione personale della partenza.",
            [
                "La cardinalità è una valutazione per viaggiatore e partenza, aggiornabile in modo idempotente.",
                "Il voto è un intero compreso fra 0 e 10; il commento è facoltativo e limitato.",
                "La valutazione è ammessa dalla data locale pari al secondo giorno successivo al rientro.",
                "Il collegamento a gruppo e agenzia consente aggregazioni tenant-safe senza duplicare i feedback di tappa.",
                "L'URL di recensione appartiene alla configurazione dell'agenzia; il codice referral deriva da agenzia, partenza e viaggiatore.",
            ],
        ),
    ],
    "Modello logico SMF Travel allineato allo schema 200",
    "1.8",
)

add_addendum(
    ROOT / "docs/data-model/SMF_Travel_Modello_Fisico_Dati_v1.7.docx",
    ROOT / "docs/data-model/SMF_Travel_Modello_Fisico_Dati_v1.8.docx",
    "Estensione del modello fisico versione 1.8",
    "Dizionario fisico degli oggetti introdotti dalla migrazione 200 e inventario canonico.",
    [
        (
            "Tabelle rooming",
            "Le tabelle sono travel.rooming_rooms e travel.rooming_room_occupants.",
            [
                "rooming_rooms conserva agency_id, departure_id, stay_id, party_id, numero camera, tipologia, capacità, esigenze e note; indici e chiavi seguono il prefisso tenant.",
                "rooming_room_occupants conserva la relazione camera-viaggiatore, con vincoli composti verso camera e profilo della partenza.",
                "Entrambe le tabelle hanno RLS forzata e policy basate sul contesto app.agency_id.",
                "app.can_manage_rooming_list_v3 e app.save_rooming_list_v3 applicano ruolo, assegnazione, capienza, appartenenza e tutela dei minori in transazione.",
            ],
        ),
        (
            "Tabella valutazioni",
            "journey.post_trip_reviews conserva le risposte post viaggio.",
            [
                "Chiave logica univoca su agenzia, partenza e viaggiatore; voto SMALLINT con CHECK 0-10 e commento limitato.",
                "RLS forzata e indici tenant-leading supportano scrittura personale e aggregazioni di agenzia.",
                "Le stored API leggono e salvano la valutazione, gestiscono l'URL pubblico dell'agenzia e alimentano Analytics.",
                "Il worker push usa il tipo post_trip_review con data locale rientro più due giorni e deduplicazione applicativa.",
            ],
        ),
        (
            "Inventario fisico verificato",
            "Lo schema applicativo contiene 99 tabelle nominalmente censite e 73 protette da RLS.",
            [
                "Il manifesto v3-table-inventory.json è la fonte versionata del gate nominale.",
                "ops.repository_migrations è una tabella tecnica opzionale del bootstrap da vuoto e non appartiene al catalogo applicativo.",
                "Validazione: zero tabelle mancanti o inattese, zero vincoli non validati, zero indici invalidi e zero tabelle tenant senza indice leading.",
                "La catena riproducibile comprende le migrazioni 001-200; CURRENT_SCHEMA_VERSION coincide con 200_v3_rooming_and_post_trip_reviews.",
            ],
        ),
    ],
    "Modello fisico SMF Travel allineato allo schema 200",
    "1.8",
)

functional_sections = [
    (
        "Rooming list per pernottamento",
        "L'operatore prepara e consegna all'hotel l'assegnazione camere.",
        [
            "Attori: responsabile, agente e accompagnatore assegnato; la guida non accede alla funzione.",
            "L'utente seleziona pernottamento e gruppo, crea camere, sceglie tipologia e assegna i viaggiatori disponibili.",
            "La pagina evidenzia persone non assegnate, sovraccarichi, duplicati e minori senza adulto dello stesso gruppo.",
            "Il salvataggio è atomico: un errore impedisce la persistenza parziale.",
            "Il documento per l'hotel riporta viaggio, struttura, date, camere, ospiti ed esigenze operative; esclude numeri di documento e informazioni sanitarie.",
        ],
    ),
    (
        "Valutazione e passaparola",
        "Il viaggiatore mantiene un punto di contatto con l'agenzia dopo il rientro.",
        [
            "La richiesta compare dal secondo giorno successivo al rientro e può essere richiamata tramite push.",
            "Il viaggiatore assegna un voto intero 0-10 e può aggiungere o modificare un commento.",
            "Con voto 9-10 l'app propone recensione pubblica e codice passaparola; con voto 0-6 propone la chat privata con l'agenzia; con voto 7-8 mostra un ringraziamento neutro.",
            "Responsabile e agente configurano il collegamento alla recensione pubblica e consultano medie e commenti in Analytics.",
            "Le analisi distinguono partenza, gruppo e viaggiatore e affiancano la valutazione complessiva ai feedback per città e tappa.",
        ],
    ),
]

add_addendum(
    ROOT / "docs/functional/SMF_Travel_Analisi_Funzionale_Completa_v2.1.docx",
    ROOT / "docs/functional/SMF_Travel_Analisi_Funzionale_Completa_v2.2.docx",
    "Estensione funzionale versione 2.2",
    "Copertura completa delle funzioni prima e dopo il viaggio introdotte nella release con schema 200.",
    functional_sections,
    "Analisi funzionale SMF Travel aggiornata alla rooming list e al post viaggio",
    "2.2",
)

add_addendum(
    ROOT / "docs/functional/SMF_Travel_Catalogo_Funzionale_v1.3.docx",
    ROOT / "docs/functional/SMF_Travel_Catalogo_Funzionale_v1.4.docx",
    "Catalogo funzionale versione 1.4",
    "Nuove capacità disponibili nella release con schema 200.",
    [
        (
            "ROOM 01 Gestione assegnazione camere",
            "Valore: sostituire il foglio Excel manuale e preparare la rooming list per ogni hotel.",
            functional_sections[0][2][0:4],
        ),
        (
            "ROOM 02 Documento per hotel",
            "Valore: produrre un documento operativo condivisibile senza dati sensibili.",
            [functional_sections[0][2][4]],
        ),
        (
            "POST 01 Valutazione complessiva",
            "Valore: raccogliere il gradimento nel momento successivo al rientro.",
            functional_sections[1][2][0:3],
        ),
        (
            "POST 02 Recensione passaparola e recupero",
            "Valore: trasformare i promotori in opportunità commerciali e gestire privatamente l'insoddisfazione.",
            [functional_sections[1][2][2], functional_sections[1][2][3]],
        ),
        (
            "ANA 03 Analisi post viaggio",
            "Valore: confrontare valutazione complessiva e percezione delle singole tappe.",
            [functional_sections[1][2][4]],
        ),
    ],
    "Catalogo funzionale SMF Travel aggiornato alla rooming list e al post viaggio",
    "1.4",
)

use_cases = [
    ("UC ROOM 01 Apertura per ruolo", "Responsabile, agente e accompagnatore assegnato aprono Rooming list; guida e utenti estranei ricevono accesso negato."),
    ("UC ROOM 02 Creazione camere", "Per ciascun pernottamento e gruppo l'operatore crea camere singole, doppie, matrimoniali e triple e assegna gli occupanti."),
    ("UC ROOM 03 Persone non assegnate", "La pagina elenca tutti i viaggiatori attivi del gruppo ancora privi di camera nello specifico pernottamento."),
    ("UC ROOM 04 Controllo capienza", "Il sistema rifiuta una camera con più occupanti della capacità dichiarata e non persiste modifiche parziali."),
    ("UC ROOM 05 Tutela minori", "Il sistema rifiuta un minore in camera senza almeno un adulto appartenente allo stesso gruppo."),
    ("UC ROOM 06 Esportazione hotel", "Il DOCX contiene struttura, date, camere, nomi ed esigenze e non contiene numeri di documento o dati sanitari."),
    ("UC POST 01 Finestra temporale", "Prima del secondo giorno dopo il rientro la valutazione non è disponibile; dalla data prevista viene mostrata e notificata una sola volta."),
    ("UC POST 02 Salvataggio idempotente", "Il viaggiatore salva voto 0-10 e commento, ripete l'operazione e ottiene un solo record aggiornato."),
    ("UC POST 03 Percorsi per voto", "Voti 9-10 propongono recensione e referral, 7-8 ringraziamento neutro, 0-6 chat privata con l'agenzia."),
    ("UC POST 04 Isolamento e Analytics", "L'agenzia vede soltanto medie e commenti del proprio tenant, aggregati per partenza e correlati ai feedback di tappa."),
]

add_addendum(
    ROOT / "docs/testing/SMF_Travel_Casi_Uso_Completi_v1.4.docx",
    ROOT / "docs/testing/SMF_Travel_Casi_Uso_Completi_v1.5.docx",
    "Casi d'uso aggiuntivi versione 1.5",
    "Casi manuali da integrare nella suite operatore per la migrazione 200.",
    [(title, description, ["Precondizioni, dati di prova, esito atteso ed evidenza devono essere registrati nella scheda di collaudo."]) for title, description in use_cases],
    "Casi d'uso SMF Travel aggiornati alla rooming list e al post viaggio",
    "1.5",
)

add_addendum(
    ROOT / "docs/testing/SMF_Travel_Rapporto_Completo_Test_v1.6_2026-09-07.docx",
    ROOT / "docs/testing/SMF_Travel_Rapporto_Completo_Test_v1.7_2026-09-08.docx",
    "Aggiornamento tecnico del rapporto test versione 1.7",
    "Evidenze automatiche completate l'8 settembre 2026; collaudo operatore intenzionalmente sospeso fino alla nuova suite.",
    [
        (
            "Evidenze tecniche superate",
            "La release con schema 200 ha superato i gate gratuiti e deterministici.",
            [
                "Quality release, build Next.js, unit test e percorsi browser pubblici superati.",
                "Ricostruzione Neon da vuoto delle migrazioni 001-200 superata su branch effimero.",
                "Inventario nominale: 99 tabelle applicative, 73 RLS, zero nomi mancanti o inattesi; il bootstrap aggiunge il solo registro checksum ops.repository_migrations.",
                "Smoke test di lettura operativa, catalogo e gamification e isolamento cross-tenant superati.",
                "Nessun test Bedrock live eseguito e nessun credito AI consumato.",
            ],
        ),
        (
            "Casi ancora da eseguire",
            "Gli esiti funzionali non vengono anticipati senza prova dell'operatore.",
            [
                "Eseguire UC ROOM 01-06 con responsabile, agente, accompagnatore e guida negativa.",
                "Eseguire UC POST 01-04 simulando le date in un tenant di collaudo e verificando push, referral, chat privata e Analytics.",
                "Registrare dispositivo, ruolo, dati usati, risultato, evidenza e identificativo di eventuale anomalia.",
            ],
        ),
    ],
    "Rapporto test SMF Travel aggiornato allo schema 200",
    "1.7",
)
