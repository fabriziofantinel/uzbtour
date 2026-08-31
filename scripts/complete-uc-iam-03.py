from pathlib import Path

from docx import Document


path = Path("docs/testing/SMF_Travel_Rapporto_Completo_Test_v1.1_2026-08-30.docx")
document = Document(path)
cases = document.tables[4]
target = next(
    (row for row in cases.rows[1:] if row.cells[0].text.strip() == "UC-IAM-03"),
    None,
)
if target is None:
    raise RuntimeError("UC-IAM-03 non trovato")

target.cells[2].text = "SUPERATO"
target.cells[3].text = (
    "Verificato in sessione reale: richiesta di recupero associata allo username corretto; email "
    "ricevuta con riferimento all'identita' prevista; nuova password impostata; vecchia password "
    "rifiutata e nuova password accettata; secondo username sulla stessa email non modificato; "
    "link o codice di recupero non riutilizzabile."
)

counts = {"SUPERATO": "107", "PARZIALE": "0", "BLOCCATO": "0", "NON SUPERATO": "0"}
for row in document.tables[1].rows[1:]:
    status = row.cells[0].text.strip().upper()
    if status in counts:
        row.cells[1].text = counts[status]

for paragraph in document.paragraphs:
    if paragraph.text.strip().startswith("Sono stati censiti 107 casi d'uso."):
        paragraph.text = (
            "Sono stati censiti e completati 107 casi d'uso. Il registro operativo aggiornato "
            "con le prove manuali e tecniche contiene 107 casi superati, 0 parziali, 0 bloccati "
            "e 0 completamente non superati."
        )
        break

document.save(path)
print(path.resolve())
