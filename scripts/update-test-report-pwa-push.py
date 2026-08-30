from pathlib import Path
from docx import Document

path=Path("docs/testing/SMF_Travel_Rapporto_Completo_Test_v1.1_2026-08-30.docx")
document=Document(path)
document.paragraphs[4].text=(
    "Sono stati censiti 107 casi d'uso. L'esecuzione ha prodotto 42 casi superati, "
    "58 parzialmente coperti, 7 bloccati da prerequisiti esterni e 0 casi completamente "
    "non superati. Le suite strutturali, la build, il modello dati, le migrazioni fino alla "
    "113, Analytics, chat, gestione spese, governance contenuti, album PDF e PWA risultano conformi."
)
document.paragraphs[11].text=(
    "Migrazioni fino alla 113 verificate: KPI, governance fonti, contratti runtime, pgcrypto, "
    "contest a due foto, scrittura bingo globale e identita Web Push multi-agenzia."
)
summary=document.tables[1]
summary.rows[2].cells[1].text="58"
summary.rows[3].cells[1].text="7"
findings=document.tables[2]
row=findings.add_row().cells
row[0].text="DEF-003";row[1].text="Chiusa";row[2].text="Web Push"
row[3].text="La funzione di sottoscrizione assegnava il record del resolver Cognito a un UUID e falliva in produzione."
row[4].text="Migrazione 113 applicata; identita Cognito diretta, upsert multi-agenzia e acceptance con rollback superata."
commands=document.tables[3]
for values in [
    ("db:migrate:v3:web-push-identity","SUPERATO","Migrazione 113","Dry-run, applicazione e gate runtime/multi-agenzia superati."),
    ("acceptance:v3:web-push","SUPERATO","Rollback","Upsert idempotente, scope partenza/gruppo e revoca verificati."),
    ("acceptance:pwa-push","SUPERATO","Service worker","Notifica background, deep-link, focus finestra e richiesta permesso verificati."),
]:
    row=commands.add_row().cells
    for cell,value in zip(row,values):cell.text=value
cases=document.tables[4]
for row in cases.rows[1:]:
    if row.cells[0].text.strip()=="UC-PWA-04":
        row.cells[2].text="PARZIALE"
        row.cells[3].text=(
            "Contratto completo verificato: permesso, subscription Cognito multi-agenzia, idempotenza, scope partenza/gruppo, "
            "revoca 404/410, notifica background e deep-link. Resta la ricezione fisica su iOS/Android."
        )
        break
else:raise RuntimeError("UC-PWA-04 non trovato")
document.save(path)
print(path.resolve())
