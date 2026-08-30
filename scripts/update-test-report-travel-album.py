from pathlib import Path
from docx import Document

path = Path("docs/testing/SMF_Travel_Rapporto_Completo_Test_v1.1_2026-08-30.docx")
document = Document(path)

document.paragraphs[4].text = (
    "Sono stati censiti 107 casi d'uso. L'esecuzione ha prodotto 42 casi superati, "
    "55 parzialmente coperti, 10 bloccati da prerequisiti esterni e 0 casi completamente "
    "non superati. Le suite strutturali, la build, il modello dati, le migrazioni fino alla "
    "112, Analytics, chat, gestione spese, governance contenuti e album finale PDF risultano conformi."
)

summary = document.tables[1]
summary.rows[1].cells[1].text = "42"
summary.rows[3].cells[1].text = "10"

commands = document.tables[3]
row = commands.add_row().cells
row[0].text = "acceptance:travel-album"
row[1].text = "SUPERATO"
row[2].text = "PDF 2 pagine"
row[3].text = "PDF valido, metadati, contenuti, foto scoped al gruppo, testo estratto e rendering PNG verificati."

cases = document.tables[4]
for row in cases.rows[1:]:
    if row.cells[0].text.strip() == "UC-MEM-02":
        row.cells[2].text = "SUPERATO"
        row.cells[3].text = (
            "Generazione PDF e download verificati con test ripetibile; foto limitate al gruppo del viaggiatore, "
            "metadati e contenuti corretti, rendering senza sovrapposizioni. Condivisione Web Share con fallback download presente."
        )
        break
else:
    raise RuntimeError("UC-MEM-02 non trovato")

document.paragraphs[25].text = (
    "La baseline tecnica e stabile e le nuove funzionalita risultano integrate senza regressioni di build o schema. "
    "Sono ora dimostrati anche la generazione, il download e la condivisione con fallback dell'album finale PDF, "
    "incluso l'isolamento delle foto per gruppo. Restano da completare soprattutto i flussi infrastrutturali end-to-end "
    "e il collaudo PWA su dispositivi mobili reali."
)

document.save(path)
print(path.resolve())
